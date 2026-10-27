"""تصدير يومية الأفراد بصيغة أرشيف 28-9-2026 الثابتة."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .afraad import BASIC_SERVICE_NAMES, emergency_horizontal_pairs
from .board_export import WEEKDAY_BY_INDEX, order_xml_children
from .export_fit import (PageFitError, TextMeter, estimate_table_height, fit_widths,
                         percentile_need)
from .labels import officer_label, personnel_label


TEMPLATE = Path(__file__).with_name("docx_templates") / "afraad_roster.docx"
LAST_FIT = {}
EMPTY_BASIC = "ــــــــــــ"
EMPTY_VALUE = "ــــــــ"
EMERGENCY_TITLE = "الخدمـــــــــــات الطارئــــــــــــــــــــــــــة"
PAIRS_PER_ROW = 2          # الملفات الحقيقية (14-10 حتى 21-10): زوجان في الصف = 4 خلايا
IMPOSSIBLE_RATIO = 1.15    # التقدير تقريبي؛ فوق هذا فقط نرفض بدل أن نفشل التصدير
FIT_TARGET = 1.0           # التقدير متشائم (ملف 14-10 الحقيقي بـ42 صفًا يتسع في Word)
PAGE_PT = 841.9
MARGINS_PT = (284 + 142) / 20
ROW_BORDER_PT = .5
CELL_HMARGIN_PT = 10.8     # 108+108 twips من نمط TableGrid
# مقاسات الخط بنصف نقطة (مدى الأرشيف الحقيقي): الأكبر أولاً، والتقدير يختار الأول الذي يتسع.
# em=None => بلا sz (يرث 11pt من النمط كما في ملف 21-10).
SIZE_TIERS = (
    dict(title=24, head=20, svc=20, names=18, small=14, em_title=18, em=None),
    dict(title=24, head=18, svc=16, names=16, small=14, em_title=18, em=18),
    dict(title=24, head=16, svc=16, names=14, small=12, em_title=16, em=16),
    dict(title=20, head=14, svc=14, names=14, small=12, em_title=14, em=14),
)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _title(day):
    try:
        d = date.fromisoformat(day)
        return (f"يومية الخدمات عن يوم {WEEKDAY_BY_INDEX[d.weekday()]} "
                f"{d.day}/{d.month}/{d.year}م خدمات الاساسية والطارئة")
    except ValueError:
        return f"يومية الأفراد — {day}"


def _person_text(person):
    person = person or {}
    if person.get("person_id"):
        who = {**person, "name": person.get("full_name") or person.get("name", "")}
        name = personnel_label(who) if person.get("role") else person.get("name", "")
    else:
        name = person.get("name", "")
    bits = [name]
    if person.get("phone"):
        bits.append(str(person["phone"]))
    return "\n".join(bits) if any(bits) else EMPTY_BASIC


def _emergency_person(person):
    person = person or {}
    if person.get("category") == "officers":
        return officer_label(person)
    if person.get("person_id") or person.get("role"):
        value = personnel_label(person)
        if person.get("phone"):
            value += "\n" + str(person["phone"])
        return value
    return str(person.get("name") or "")


def _occasional_who(row):
    people = [officer_label(o) for o in row.get("officers") or []]
    people.extend(_emergency_person(p) for p in row.get("personnel") or [])
    value = "، ".join(item for item in people if item)
    return value or EMPTY_BASIC


def _literal(value, empty=EMPTY_VALUE):
    return str(value) if value not in (None, "") else empty


def _strength(row):
    return _literal(row.get("count"))


def _occasional_count(row):
    count = row.get("conscript_count")
    if not count:
        count = sum((entry.get("count") or 0) for entry in row.get("conscripts") or [])
    return count or EMPTY_VALUE


def _occasional_weapon(row):
    if row.get("weapon"):
        return row["weapon"]
    classes = [entry.get("class", "") for entry in row.get("conscripts") or []]
    return " + ".join(dict.fromkeys(value for value in classes if value)) or EMPTY_VALUE


def _occasional_time(row):
    value = str(row.get("time") or "")
    party = str(row.get("party") or "")
    if party and party not in value:
        value = f'{value} "{party}"'.strip()
    return value or EMPTY_VALUE


def _basic_values(row):
    name = str(row.get("name") or "")
    split = name.startswith("كمين 109 -")
    if split:
        parent, _, child = name.partition(" - ")
        return ["", parent, child or EMPTY_VALUE, _person_text(row.get("morning")),
                _person_text(row.get("night")), _strength(row),
                _literal(row.get("weapon")), _literal(row.get("schedule"))]
    return ["", name, _person_text(row.get("morning")), _person_text(row.get("night")),
            _strength(row), _literal(row.get("weapon")), _literal(row.get("schedule"))]


def _export_basic_rows(rows):
    if rows:
        normalized = []
        for row in rows:
            value = dict(row)
            name = " ".join(str(value.get("name") or "").split())
            if name == "كمين 109":
                value["name"] = "كمين 109 - بيك اب"
            elif name == "مدرعة" and normalized:
                previous = str(normalized[-1].get("name") or "")
                if previous.startswith("كمين 109 -"):
                    value["name"] = "كمين 109 - مدرعة"
            normalized.append(value)
        return normalized
    return [{"name": name, "morning": {}, "night": {}, "count": "", "weapon": "",
             "schedule": ""} for name in BASIC_SERVICE_NAMES]



def _basic_rows_values(rows):
    """[(split, values)] — values بلا عمود «م» (يُرقَّم تلقائيًا في Word)."""
    result = []
    for row in _export_basic_rows(rows):
        values = _basic_values(row)[1:]
        result.append((str(row.get("name") or "").startswith("كمين 109 -"), values))
    return result


def _horizontal_values(pair_group):
    values = []
    for pair in pair_group:
        values.extend([str(pair.get("label") or pair.get("name") or ""),
                       str(pair.get("person") or pair.get("who") or "")])
    return values


def _emergency_rows(afraad):
    """[('pair'|'row', values)] بنفس المنطق السابق، لكن زوجان في الصف كالأرشيف."""
    rows = []
    officer_pairs = afraad.get("officer_pairs") or []
    for start in range(0, len(officer_pairs), PAIRS_PER_ROW):
        rows.append(("pair", _horizontal_values(officer_pairs[start:start + PAIRS_PER_ROW])))
    for source in afraad.get("occasional") or []:
        pairs = emergency_horizontal_pairs(source)
        if len(pairs) > 1:
            for start in range(0, len(pairs), PAIRS_PER_ROW):
                rows.append(("pair", _horizontal_values(pairs[start:start + PAIRS_PER_ROW])))
            continue
        rows.append(("row", [source.get("label") or source.get("name") or EMPTY_VALUE,
                             _occasional_who(source), _occasional_count(source),
                             _occasional_weapon(source), _occasional_time(source)]))
    return rows


# ------------------------------------------------------------------ XML

def _q(name):
    return "{%s}%s" % (W_NS, name)


def _set_size(rpr, half):
    """half=None يحذف sz فيرث الحجم من النمط."""
    if rpr is None:
        return
    for tag in ("sz", "szCs"):
        node = rpr.find(_q(tag))
        if half is None:
            if node is not None:
                rpr.remove(node)
            continue
        if node is None:
            node = rpr.makeelement(_q(tag), {})
            rpr.append(node)
        node.set(_q("val"), str(half))


def _first_paragraph_protos(tc):
    paragraph = tc.find(_q("p"))
    ppr = paragraph.find(_q("pPr"))
    run = paragraph.find(_q("r"))
    rpr = run.find(_q("rPr")) if run is not None else None
    if rpr is None and ppr is not None:
        rpr = ppr.find(_q("rPr"))
    return ppr, rpr


def _fill(tc, text, half, protos=None):
    """يكتب النص (سطر لكل فقرة) في tc بنسخ pPr/rPr من خلية حقيقية؛ النص والحجم فقط يتغيران."""
    ppr_proto, rpr_proto = protos or _first_paragraph_protos(tc)
    for paragraph in tc.findall(_q("p")):
        tc.remove(paragraph)
    lines = str(text).split("\n") if text not in (None, "") else [""]
    for line in lines:
        paragraph = tc.makeelement(_q("p"), {})
        if ppr_proto is not None:
            ppr = deepcopy(ppr_proto)
            _set_size(ppr.find(_q("rPr")), half)
            paragraph.append(ppr)
        if line:
            run = paragraph.makeelement(_q("r"), {})
            rpr = deepcopy(rpr_proto)
            _set_size(rpr, half)
            run.append(rpr)
            node = run.makeelement(_q("t"), {})
            node.text = line
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            run.append(node)
            paragraph.append(run)
        tc.append(paragraph)


def _fit_tier(rows, widths, usable, widths_for=None):
    """يختار أكبر مقاس حقيقي يتسع تقديريًا؛ التقدير لا يفرض ارتفاعات أبدًا.

    `widths_for(tier)`: يعيد عروض الأعمدة المحسوبة من محتوى اليوم عند هذا المقاس."""
    meter = TextMeter()
    total = 0.0
    for tier in SIZE_TIERS:
        if widths_for is not None:
            widths = widths_for(tier)
        total = _estimate(rows, widths, tier, meter)
        overhead = tier["title"] / 2 * 1.2 + 10 + 24
        total += overhead
        if total <= usable * FIT_TARGET:
            return tier, total
    if total <= usable * IMPOSSIBLE_RATIO:
        return SIZE_TIERS[-1], total
    raise PageFitError("تعذر احتواء اليومية في صفحة واحدة (نسبة تقديرية %.2f)" % (total / usable))


def _estimate(rows, widths, tier, meter):
    total = 0.0
    for kind, values, minimum in rows:
        row_height = minimum / 20
        for index, value in enumerate(values):
            role, width, font = _role(kind, index, widths)
            half = tier.get(role)
            size = (half or 22) / 2
            lines = str(value or "").split("\n")
            spec = {"text": "\n".join(lines), "size_pt": size, "bold": True, "font": font,
                    "hmargins": CELL_HMARGIN_PT, "margins": (0.0, 2.0 * len(lines))}
            height = estimate_table_height([[spec]], [width / 20])[0]
            row_height = max(row_height, height)
        total += row_height + ROW_BORDER_PT
    return total


BASE_FONT, EM_FONT = "Times New Roman", "Bahnschrift SemiBold"


def _role(kind, index, widths):
    if kind in ("regular", "split"):
        offset = 0 if kind == "regular" else 1
        # values بلا «م»: regular [svc, m, n, qawam, taslih, intizam]
        if kind == "regular":
            roles = ["svc", "names", "names", "small", "small", "small"]
            cols = widths["regular"][1:]
        else:
            roles = ["svc", "svc", "names", "names", "small", "small", "small"]
            cols = widths["split"][1:]
        return roles[index], cols[index], BASE_FONT
    if kind == "heading":
        return "em_title", widths["heading"][0], BASE_FONT
    if kind == "pair":
        return "em", widths["pair"][index], EM_FONT
    return "em", widths["row"][index], EM_FONT


def _tc_widths(tr):
    return [int(tc.find(_q("tcPr")).find(_q("tcW")).get(_q("w"))) for tc in tr.findall(_q("tc"))]


def _set_title(document, day, half):
    paragraph = next(p for p in document.paragraphs if "يومية الخدمات" in p.text)
    runs = paragraph._p.findall(_q("r"))
    first = runs[0]
    first.find(_q("t")).text = _title(day)
    first.find(_q("t")).set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for extra in runs[1:]:
        paragraph._p.remove(extra)
    _set_size(first.find(_q("rPr")), half)
    _set_size(paragraph._p.find(_q("pPr")).find(_q("rPr")), half)



# الأعمدة المنطقية (م، الخدمة، الصباحية، الليلية، القوام، التسليح، الانتظام) وعدد
# أعمدة الشبكة الفعلية لكل منها في القالب (g0 و g13 شريحتان صغيرتان ثابتتان).
LOGICAL_GROUPS = ((1,), (2, 3), (4, 5), (6, 7, 8), (9, 10), (11,), (12,))
LOGICAL_RANGES = ((574, 609), (1937, 2005), (1950, 2319), (2070, 2246), (2046, 2358),
                  (940, 1119), (1073, 1123))
LOGICAL_ROLES = ("svc", "svc", "names", "names", "small", "small", "small")   # أول عمود (م) يتجاهل
MEASURE_FONT = "Times New Roman"


def _template_grid(tbl):
    return [int(col.get(_q("w"))) for col in tbl.find(_q("tblGrid")).findall(_q("gridCol"))]


def _logical_needs(basic, tier):
    """احتياج كل عمود منطقي (twips) من نصوص الخدمات الأساسية عند مقاس الخط المختار."""
    meter = TextMeter()
    # values بلا «م»: regular [svc, m, n, qawam, taslih, intizam]؛ المقسوم [parent, child, m, n, ...]
    columns = {1: [], 2: [], 3: [], 4: [], 5: [], 6: []}
    for is_split, values in basic:
        mapping = ([None, None, 2, 3, 4, 5, 6] if is_split else [1, 2, 3, 4, 5, 6])
        for index, text in enumerate(values):
            logical = mapping[index] if index < len(mapping) else None
            if logical is None:
                continue
            role = LOGICAL_ROLES[logical]
            size = (tier.get(role) or 22) / 2
            lines = str(text or "").split("\n")
            width = max(meter.text_width(line, size, True, MEASURE_FONT) for line in lines)
            columns[logical].append((width + CELL_HMARGIN_PT + 3) * 20)
    return [0] + [percentile_need(columns[i]) for i in range(1, 7)]


def _apply_grid(tbl, trs, logical):
    """يعيد توزيع شبكة الجدول على الأعمدة المنطقية الجديدة (كل مجموعة تحافظ على
    تناسب أعمدتها الداخلية)، ثم يضبط tcW لكل خلية = مجموع الأعمدة التي تمتد عليها."""
    base = _template_grid(tbl)
    grid = list(base)
    for group, target in zip(LOGICAL_GROUPS, logical):
        old = sum(base[i] for i in group)
        if group == (12,):
            target -= base[13]            # شريحة الشبكة الأخيرة (gridAfter) ثابتة
            old -= 0
        for i in group:
            grid[i] = round(base[i] * target / old) if old else base[i]
        grid[group[-1]] += target - sum(grid[i] for i in group)
    for col, width in zip(tbl.find(_q("tblGrid")).findall(_q("gridCol")), grid):
        col.set(_q("w"), str(width))
    for tr in trs:
        pr = tr.find(_q("trPr"))
        before = pr.find(_q("gridBefore")) if pr is not None else None
        start = int(before.get(_q("val"))) if before is not None else 0
        for tc in tr.findall(_q("tc")):
            span_node = tc.find(_q("tcPr")).find(_q("gridSpan"))
            span = int(span_node.get(_q("val"))) if span_node is not None else 1
            tc.find(_q("tcPr")).find(_q("tcW")).set(_q("w"), str(sum(grid[start:start + span])))
            start += span
    return grid


def build_docx(afraad):
    document = Document(TEMPLATE)
    table = document.tables[0]
    tbl = table._tbl
    trs = tbl.findall(_q("tr"))
    header, regular, split_first, split_next, heading, pair, emergency = trs
    for tr in trs:
        tbl.remove(tr)
    def row_widths():
        return {"regular": _tc_widths(regular),
                "split": _tc_widths(split_first)[:3] + _tc_widths(regular)[2:],
                "heading": _tc_widths(heading), "pair": _tc_widths(pair),
                "row": _tc_widths(emergency)}
    base_grid = _template_grid(tbl)
    logical_total = sum(base_grid[1:14])
    state = {}

    def widths_for(tier):
        logical = fit_widths(_logical_needs(basic, tier), LOGICAL_RANGES, logical_total)
        state["grid"] = _apply_grid(tbl, trs, logical)
        state["logical"] = logical
        return row_widths()

    widths = row_widths()
    # صف 109 المقسوم = صف عادي بعد استبدال خليتي الخدمة بـ (م، الاسم، الفرع) من الملف الحقيقي
    def split_tr(source):
        tr = deepcopy(regular)
        tr.remove(tr.find(_q("trPr")))
        tr.insert(0, deepcopy(source.find(_q("trPr"))))
        cells = tr.findall(_q("tc"))
        new = [deepcopy(c) for c in source.findall(_q("tc"))[:3]]
        position = list(tr).index(cells[0])
        tr.remove(cells[0])
        tr.remove(cells[1])
        for offset, cell in enumerate(new):
            tr.insert(position + offset, cell)
        return tr

    basic = _basic_rows_values(afraad.get("basic"))
    emergency_rows = _emergency_rows(afraad)

    spec = []
    for is_split, values in basic:
        spec.append(("split" if is_split else "regular", values, 297 if is_split else 0))
    spec.append(("heading", [EMERGENCY_TITLE], 30))
    for kind, values in emergency_rows:
        spec.append((kind, values, 261 if kind == "pair" else 108))
    usable = PAGE_PT - MARGINS_PT
    tier, total = _fit_tier(spec, widths, usable, widths_for)
    widths_for(tier)                     # الشبكة النهائية عند المقاس المختار
    split_protos = (split_tr(split_first), split_tr(split_next))

    # الرأس
    head_cells = header.findall(_q("tc"))
    for tc in head_cells:
        _fill(tc, "".join(t.text or "" for t in tc.iter(_q("t"))), tier["head"])
    tbl.append(header)

    group = 0
    for is_split, values in basic:
        if is_split:
            tr = deepcopy(split_protos[1 if group else 0])
            group += 1
        else:
            group = 0
            tr = deepcopy(regular)
        cells = tr.findall(_q("tc"))
        if is_split:
            roles = [None, "svc", "svc", "names", "names", "small", "small", "small"]
            texts = [None] + list(values)
            continued = group > 1
        else:
            roles = [None, "svc", "names", "names", "small", "small", "small"]
            texts = [None] + list(values)
            continued = False
        for index, (tc, role, text) in enumerate(zip(cells, roles, texts)):
            if role is None:
                continue
            if continued and index == 1:
                continue            # خلية الاسم المدموجة رأسيًا (استكمال): فارغة كالملف الحقيقي
            _fill(tc, text, tier[role])
        tbl.append(tr)

    tr = deepcopy(heading)
    _fill(tr.find(_q("tc")), EMERGENCY_TITLE, tier["em_title"])
    tbl.append(tr)

    for kind, values in emergency_rows:
        if kind == "pair":
            tr = deepcopy(pair)
            cells = tr.findall(_q("tc"))
            values = list(values) + [""] * (len(cells) - len(values))
        else:
            tr = deepcopy(emergency)
            cells = tr.findall(_q("tc"))
        for tc, text in zip(cells, values):
            _fill(tc, text, tier["em"])
        tbl.append(tr)

    _set_title(document, afraad["date"], tier["title"])
    order_xml_children(document)
    LAST_FIT.clear()
    LAST_FIT.update(scale=round(tier["svc"] / 20, 2), ratio=total / usable, total_pt=total,
                    usable_pt=usable, tier=SIZE_TIERS.index(tier),
                    grid=list(state["grid"]), logical=list(state["logical"]))

    buf = BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf
