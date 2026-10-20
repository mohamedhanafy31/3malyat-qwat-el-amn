"""Deterministic, renderer-independent fitting helpers for daily exports."""
from dataclasses import dataclass
import logging
from pathlib import Path
import subprocess
from typing import Iterable

from PIL import ImageFont
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

log = logging.getLogger(__name__)
WIDTH_FACTOR = 1.07
LINE_HEIGHT_FACTOR = 1.1499


class PageFitError(RuntimeError):
    pass


class TextMeter:
    _fonts = {}
    _all_paths = None
    _widths = {}

    def __init__(self):
        self._paths = None

    def _font_paths(self):
        if TextMeter._all_paths is None:
            paths = []
            try:
                out = subprocess.check_output(["fc-list", ":", "file"], text=True,
                                              stderr=subprocess.DEVNULL)
                paths.extend(line.split(":", 1)[0] for line in out.splitlines())
            except (OSError, subprocess.SubprocessError):
                pass
            for root in (Path("/usr/share/fonts"), Path.home() / ".fonts"):
                if root.exists():
                    paths.extend(str(p) for p in root.rglob("*.ttf"))
            TextMeter._all_paths = list(dict.fromkeys(paths))
        self._paths = TextMeter._all_paths
        return TextMeter._all_paths

    def _font_path(self, name, bold=False):
        key = (name or "Arial").lower()
        wanted = [key]
        if "sultan" in key or "bahnschrift" in key or "simplified" in key:
            wanted += ["arial", "times new roman", "noto sans arabic", "dejavu sans"]
        elif "times" in key:
            wanted += ["times new roman", "arial", "noto sans arabic"]
        else:
            wanted += ["arial", "noto sans arabic", "dejavu sans"]
        candidates = self._font_paths()
        for wanted_name in wanted:
            matches = [p for p in candidates if wanted_name in Path(p).stem.lower()]
            if bold:
                matches.sort(key=lambda p: ("bold" not in p.lower(), p))
            if matches:
                return matches[0]
        return candidates[0] if candidates else None

    def _font(self, name, size_pt, bold):
        path = self._font_path(name, bold)
        if not path:
            return None
        key = (path, round(size_pt, 2), bold)
        if key not in self._fonts:
            try:
                self._fonts[key] = ImageFont.truetype(
                    path, max(1, round(size_pt * 96 / 72)),
                    layout_engine=ImageFont.Layout.RAQM)
                log.debug("TextMeter using %s for %s", path, name)
            except OSError:
                self._fonts[key] = None
        return self._fonts[key]

    def text_width(self, text, size_pt, bold=False, font_name="Arial"):
        text = str(text or "")
        key = (text, round(size_pt, 2), bool(bold), font_name)
        if key in self._widths:
            return self._widths[key]
        font = self._font(font_name, size_pt, bold)
        if font is None:
            value = len(text) * size_pt * 0.55 * WIDTH_FACTOR
        else:
            width = font.getlength(text) * 72 / 96
            value = width * WIDTH_FACTOR
        self._widths[key] = value
        return value

    def wrap_lines(self, text, width_pt, size_pt, bold=False, font_name="Arial"):
        lines = []
        for paragraph in str(text or "").splitlines() or [""]:
            words = paragraph.split()
            if not words:
                lines.append("")
                continue
            line = ""
            for word in words:
                if not line and self.text_width(word, size_pt, bold, font_name) > width_pt:
                    chunk = ""
                    for char in word:
                        if chunk and self.text_width(chunk + char, size_pt, bold, font_name) > width_pt:
                            lines.append(chunk)
                            chunk = char
                        else:
                            chunk += char
                    line = chunk
                    continue
                candidate = f"{line} {word}".strip()
                if line and self.text_width(candidate, size_pt, bold, font_name) > width_pt:
                    lines.append(line)
                    line = word
                else:
                    line = candidate
            lines.append(line)
        return lines or [""]


def _cell_value(cell):
    if isinstance(cell, dict):
        return cell
    return {"text": str(cell or "")}


def estimate_table_height(rows, grid_widths_pt, meter=None, slack=True):
    meter = meter or TextMeter()
    total = 0.0
    heights = []
    vertical = []
    for row in rows:
        row_height = 0.0
        for index, raw in enumerate(row):
            cell = _cell_value(raw)
            colspan = max(1, int(cell.get("colspan", 1)))
            width = sum(grid_widths_pt[index:index + colspan])
            margins = cell.get("margins", (0.0, 0.0))
            if isinstance(margins, (int, float)):
                margins = (float(margins), float(margins))
            inner = max(1.0, width - sum(margins))
            size = float(cell.get("size_pt", 10))
            lines = meter.wrap_lines(cell.get("text", ""), inner, size,
                                     bool(cell.get("bold")), cell.get("font", "Arial"))
            need = len(lines) * LINE_HEIGHT_FACTOR * size + sum(margins) + 1.0
            if slack and lines and len(lines) > 1:
                last = meter.text_width(lines[-1], size, bool(cell.get("bold")),
                                        cell.get("font", "Arial"))
                if last >= inner * .9:
                    need += LINE_HEIGHT_FACTOR * size
            row_height = max(row_height, need)
            rowspan = max(1, int(cell.get("rowspan", 1)))
            if rowspan > 1:
                vertical.append((len(heights), rowspan, need))
        heights.append(row_height)
        total += row_height
    for start, rowspan, need in vertical:
        current = sum(heights[start:start + rowspan])
        if current < need:
            heights[min(len(heights) - 1, start + rowspan - 1)] += need - current
    return sum(heights), heights


@dataclass
class FitResult:
    scale: float
    heights_pt: list
    total_pt: float
    usable_pt: float
    ratio: float


def fit_table(table_spec, usable_height_pt, target=.94, min_scale=.55, step=.02):
    rows = table_spec["rows"] if isinstance(table_spec, dict) else table_spec.rows
    widths = table_spec["grid_widths_pt"] if isinstance(table_spec, dict) else table_spec.grid_widths_pt
    original = table_spec.get("base_scale", 1.0) if isinstance(table_spec, dict) else 1.0
    scale = 1.0
    while scale >= min_scale - 1e-9:
        scaled = []
        for row in rows:
            scaled.append([{**_cell_value(c), "size_pt": max(4.0, float(_cell_value(c).get("size_pt", 10)) * scale)}
                           for c in row])
        total, heights = estimate_table_height(scaled, widths)
        if total <= usable_height_pt * target:
            return FitResult(round(scale, 2), heights, total, usable_height_pt,
                             total / usable_height_pt)
        scale = round(scale - step, 2)
    raise PageFitError("تعذر احتواء اليومية في صفحة واحدة (نسبة تقديرية %.2f)" %
                       (total / usable_height_pt if usable_height_pt else float("inf")))


def _set_run_size(run, size):
    rpr = run._r.get_or_add_rPr()
    for tag in ("w:sz", "w:szCs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
        node.set(qn("w:val"), str(max(1, round(size * 2))))
        if node.getparent() is None:
            rpr.append(node)


def apply_fit_to_docx_table(table, fit_result):
    for row, height in zip(table.rows, fit_result.heights_pt):
        row.height = Pt(round(height + 1, 1))
        tr_pr = row._tr.get_or_add_trPr()
        node = tr_pr.find(qn("w:trHeight"))
        if node is None:
            node = OxmlElement("w:trHeight")
        node.set(qn("w:val"), str(round((height + 1) * 20)))
        node.set(qn("w:hRule"), "exact")
        if node.getparent() is None:
            tr_pr.append(node)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                ppr = paragraph._p.get_or_add_pPr()
                spacing = ppr.find(qn("w:spacing"))
                if spacing is None:
                    spacing = OxmlElement("w:spacing")
                    ppr.append(spacing)
                spacing.set(qn("w:before"), "0")
                spacing.set(qn("w:after"), "0")
                spacing.set(qn("w:lineRule"), "exact")
                for run in paragraph.runs:
                    current = run.font.size.pt if run.font.size else 10
                    final_size = max(4, round(current * fit_result.scale * 2) / 2)
                    spacing.set(qn("w:line"), str(round(final_size * LINE_HEIGHT_FACTOR * 20)))
                    _set_run_size(run, final_size)


def usable_height_pt(section, final_paragraph_allowance=2):
    value = section.page_height.pt - section.top_margin.pt - section.bottom_margin.pt
    return max(1, value - final_paragraph_allowance)


def docx_table_spec(table, default_size=10):
    """Extract rows without python-docx's repeated views of merged cells."""
    widths = [c.width.pt if c.width else 40 for c in table.columns]
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        physical = row.cells
        for index, cell in enumerate(physical):
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            span = 1
            while index + span < len(physical) and physical[index + span]._tc is cell._tc:
                span += 1
            runs = [r for p in cell.paragraphs for r in p.runs]
            size = next((r.font.size.pt for r in runs if r.font.size), default_size)
            cells.append({"text": cell.text, "size_pt": size,
                          "bold": any(r.bold for r in runs), "colspan": span})
        rows.append(cells)
    return rows, widths


def configure_xlsx_one_page(ws):
    from openpyxl.utils import get_column_letter
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.print_options.horizontalCentered = True
    if ws.max_row and ws.max_column:
        ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
