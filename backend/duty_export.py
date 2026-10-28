"""تصدير يومية تشغيل الضباط بصيغة ملف Word الحقيقي (أكتوبر 2026).

القالب `docx_templates/duty_roster.docx` نسخة من ملف أرشيف حقيقي: الشبكة وعروض
الأعمدة وارتفاعات الصفوف وترتيب الجسم والشعارات والجدول الملخّص (17 عمودًا)
كلها من الملف الحقيقي، والمُصدِّر لا يغيّر إلا النص. خصائص الخط (الخطوط
النسقية majorBidi، عريض، 9pt للجدول الأول و8pt للملخّص) تُنسخ من خصائص
أول run في الخلية في القالب ولا تُكتب من جديد؛ والتصغير يتم لخلية واحدة
طويلة فقط (8 ثم 7pt) كما يفعل من يكتب الملف يدويًا، وليس على الجدول كله.
"""
from copy import deepcopy
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell

from .board_export import WEEKDAY_BY_INDEX, order_xml_children
from .constants import (OFFICER_SECTIONS, SECTION_FORCE, SECTION_GUARDS,
                        SECTION_OUTSIDE)
from .export_fit import PageFitError, TextMeter, usable_height_pt
from .labels import RANK_SHORT, officer_label

TEMPLATE = Path(__file__).with_name("docx_templates") / "duty_roster.docx"
LAST_FIT = {}
ROSTER_SLOTS = {SECTION_FORCE: 24, SECTION_GUARDS: 8, SECTION_OUTSIDE: 2}

ROSTER_SIZE_PT = 9          # sz=18 في كل خلايا الجدول الأول بالملف الحقيقي
SUMMARY_SIZE_PT = 8         # sz=16 في الجدول الملخّص
SHRINK_STEPS_PT = (9, 8, 7)  # تصغير خلية واحدة طويلة (سابقة يوم 21: 8 و7)
MAX_CELL_LINES = 2          # أكثر من سطرين بـ9pt => الكاتب اليدوي يصغّر الخلية
CELL_MARGIN_PT = 10.8       # هامش الخلية الأيمن + الأيسر (108+108 twips) من نمط الجدول
REAL_LINE_FACTOR = 1.1499 * 1.15   # ارتفاع سطر Times New Roman × line=276
FIT_TARGET = 1.0
FIT_LIMIT = 1.08            # المقدّر متشائم؛ فوق هذا نرفض بدل أن نكسر الصفحة الواحدة
FONT_NAME = "Times New Roman"


def _physical_cells(row):
    """خلايا XML الفعلية للصف بلا الأسماء المستعارة للدمج الرأسي في python-docx."""
    return [_Cell(tc, row) for tc in row._tr.tc_lst]


def _rest_text(row):
    system = str(row.get("rest_system") or "").strip()
    if not system or system == "—" or (system == "أسبوعية" and not row.get("rest_day")):
        return "------------"
    if system == "أسبوعية":
        return row["rest_day"]
    return system


def _status_period_text(row):
    """صياغة الحالة الممتدة في Word؛ المفتوحة تذكر يوم البداية فقط."""
    period = row.get("persistent_status") or {}
    if row.get("status_source") != "period" or not period:
        return ""
    status = period.get("status") or row.get("status") or "خارج"
    if period.get("title") and period.get("label"):
        return period["label"]      # «مأمورية: <الاسم> — <التفاصيل> — من … حتى …»
    try:
        start = date.fromisoformat(period.get("start") or "").strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        start = period.get("start") or ""
    if not period.get("end"):
        return f"{status} من يوم {start}" if start else status
    try:
        end = date.fromisoformat(period["end"]).strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        end = period["end"]
    return f"{status} من {start} حتى {end}" if start else f"{status} حتى {end}"


def _leave_text(leave):
    """صياغة الراحة داخل «التشغيل اليومي» في ملف Word بس."""
    raw_type = str((leave or {}).get("raw_type") or "").strip()
    if raw_type == "أسبوعية":
        return "راحة أسبوعية"

    label = str((leave or {}).get("type") or raw_type or "راحة").strip()
    label = {
        "نصف شهرية": "راحة نصف شهرية",
        "شهرية": "راحة شهرية",
        "مجمعة": "راحة مجمعة",
    }.get(raw_type, label)

    def short_day(value):
        try:
            parsed = date.fromisoformat(str(value or ""))
        except ValueError:
            return str(value or "")
        return f"{parsed.day}/{parsed.month}"

    return_date = (leave or {}).get("return_date")
    if not return_date:
        try:
            return_date = (date.fromisoformat(str((leave or {}).get("end") or ""))
                           + timedelta(days=1)).isoformat()
        except ValueError:
            return_date = ""
    return f"{label} (من {short_day((leave or {}).get('start'))} والعودة {short_day(return_date)})"


def _daily_text(row):
    """«التشغيل اليومي»: الخدمة التلقائية من اللوحة + التشغيل الإضافي اليدوي."""
    # الفرقة المحسوبة من سجل الالتحاق ليس لها حالة يومية مكتوبة؛ من دون
    # هذا الفرع كان ملف Word يطبع «خوارج» فقط رغم أن الشاشة تعرض الاسم
    # والمدة والمكان. التكليف الفعلي يظل أقوى من مدى الفرقة، لذلك نطبقها
    # فقط عندما كانت الخانة النهائية نفسها «خوارج/فرقة».
    if (row.get("group"), row.get("bucket")) == ("خوارج", "فرقة") and row.get("course"):
        from .duty import course_with_note
        return course_with_note(row["course"], row.get("text"))
    period_text = _status_period_text(row)
    if period_text:
        return period_text
    if row.get("text"):
        return row["text"]
    if row.get("leave"):
        return _leave_text(row["leave"])
    return row.get("status") or row.get("group") or "عمل"


def _grouped_rows(rows):
    by_section = {}
    for row in rows:
        by_section.setdefault(row.get("section") or SECTION_FORCE, []).append(row)
    return [(section, by_section.get(section, [])) for section in OFFICER_SECTIONS]


def _title_pieces(day):
    """نصوص أجزاء العنوان بنفس تقطيع الملف الحقيقي (10 runs).

    أسماء الأيام تُكتب بالهمزة (الأحد/الأربعاء) بقرار المالك، على عكس الملف الحقيقي."""
    d = date.fromisoformat(day)
    return ["يومية تشغيل الضبـــاط بإدارة قـوات أمن السويس عـن يوم", " ",
            WEEKDAY_BY_INDEX[d.weekday()], " الموافق", " ", str(d.day), " /",
            str(d.month), "/", f"{d.year}م"]


def _set_title(document, day):
    """يعيد كتابة نص runs العنوان في الفقرة الأصلية؛ خصائصها (14pt، تسطير،
    cs=Sultan bold) تبقى من القالب."""
    pieces = _title_pieces(day)
    for paragraph in document.paragraphs:
        if "يومية تشغيل الضب" not in paragraph.text:
            continue
        runs = [r for r in paragraph._p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
        if len(runs) < len(pieces):
            # قالب غير معتاد: أبقِ أول run واكتب العنوان كله فيه.
            for extra in runs[1:]:
                paragraph._p.remove(extra)
            runs[0].find(qn("w:t")).text = "".join(pieces)
            return
        for run, text in zip(runs, pieces):
            node = run.find(qn("w:t"))
            node.text = text
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for extra in runs[len(pieces):]:
            paragraph._p.remove(extra)
        return


def _base_rpr(cell, size_pt):
    """خصائص run لخلية: نسخة من أول run في القالب (أو علامة الفقرة) وتغيير الحجم فقط."""
    paragraph = cell._tc.find(qn("w:p"))
    run = next((r for p in cell._tc.findall(qn("w:p")) for r in p.findall(qn("w:r"))
                if r.find(qn("w:rPr")) is not None), None)
    source = run.find(qn("w:rPr")) if run is not None else None
    if source is None:
        ppr = paragraph.find(qn("w:pPr")) if paragraph is not None else None
        source = ppr.find(qn("w:rPr")) if ppr is not None else None
    rpr = deepcopy(source) if source is not None else OxmlElement("w:rPr")
    for tag in ("w:b", "w:bCs", "w:rtl"):
        if rpr.find(qn(tag)) is None:
            rpr.append(OxmlElement(tag))
    _set_rpr_size(rpr, size_pt)
    return rpr


def _set_rpr_size(rpr, size_pt):
    half = str(round(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        node.set(qn("w:val"), half)


def _write_cell(cell, text, *, size_pt, paragraphs=None):
    """يستبدل نص الخلية فقط مع إبقاء pPr/rPr القالب (تباعد 276، توسيط، الخطوط).

    `paragraphs`: قائمة نصوص لفقرة لكل عنصر (خلية الأسماء)؛ تُنسخ فقرة القالب الأولى."""
    tc = cell._tc
    rpr = _base_rpr(cell, size_pt)
    template = tc.find(qn("w:p"))
    for extra in tc.findall(qn("w:p"))[1:]:
        tc.remove(extra)
    texts = paragraphs if paragraphs is not None else [text]
    blank = deepcopy(template)
    for child in list(blank):
        if child.tag != qn("w:pPr"):
            blank.remove(child)
    pprs = []
    tc.remove(template)
    for value in texts:
        paragraph = deepcopy(blank)
        ppr = paragraph.find(qn("w:pPr"))
        if ppr is not None:
            mark = ppr.find(qn("w:rPr"))
            if mark is not None:
                _set_rpr_size(mark, size_pt)
        if value not in (None, ""):
            run = OxmlElement("w:r")
            run.append(deepcopy(rpr))
            node = OxmlElement("w:t")
            node.text = str(value)
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            run.append(node)
            paragraph.append(run)
        tc.append(paragraph)
        pprs.append(paragraph)
    return pprs


def _summary_values(summary):
    """قيم صف الإجمالي بترتيب أعمدة الملف الحقيقي (17 عمودًا):
    أصل القوة، خارجية(2)، داخلية(2)، طبية(2)، خوارج: تقصيرة/راحة/دورية/غياب/
    مرضي/فرقة/مأمورية/انتداب، حراسات، (الصافي = الأسماء)."""
    external, internal = summary["خارجية"], summary["داخلية"]
    medical, outside = summary["طبية"], summary["خوارج"]
    morning = str(external["صباحية"])
    if external.get("بحث"):
        morning = (f"{morning} + {external['بحث']} بحث" if external["صباحية"]
                   else f"{external['بحث']} بحث")
    # «N معسكر فرعي» بيتكتب في الصف اللي تحت الداخلية (`_sparse_summary`)،
    # مش جنب الرقم — غير كده كان بيتطبع مرتين.
    internal_morning = str(internal["صباحية"])
    fixed_outside = {"تقصيرة", "راحة", "طارئة", "غياب", "مرضي", "فرقة", "انتداب", "مأمورية"}
    other_emergency = sum(value for key, value in outside.items() if key not in fixed_outside)
    return [summary["أصل القوة"], morning, external["ليلية"], internal_morning,
            internal["ليلية"], medical["موجود"], medical["راحة"], outside["تقصيرة"],
            outside["راحة"], outside["طارئة"] + other_emergency, outside["غياب"], outside["مرضي"],
            outside["فرقة"], outside.get("مأمورية", 0), outside["انتداب"], summary["حراسات"]]


def _net_names(rows):
    """«رتبة / اسم» لكل ضابط صافي — فقرة لكل ضابط (الترقيم من قائمة Word)."""
    names = []
    for row in rows:
        if row.get("group") == "صافي":
            label = officer_label(row, full=False)
            names.append(label.replace("/ ", " / ", 1) if "/ " in label and " / " not in label
                         else label)
    return names


def _sparse_summary(rows):
    subcamp = sum(1 for row in rows if row.get("group") == "معسكر فرعي")
    deputed = False
    for row in rows:
        deputed = deputed or row.get("status") == "انتداب"
    return (f"{subcamp} معسكر فرعي" if subcamp else ""), ("منتدب" if deputed else "")




def _fill_summary(table, summary, rows, names_size=SUMMARY_SIZE_PT):
    header, names_row, total, sparse = (_physical_cells(table.rows[index]) for index in range(4))
    values = _summary_values(summary)
    # صف الإجمالي: 16 قيمة ثم خلية الصافي (استمرار دمج رأسي لخلية الأسماء).
    for index, value in enumerate(values):
        _write_cell(total[index], value, size_pt=SUMMARY_SIZE_PT)
    names = _net_names(rows)
    pprs = _write_cell(names_row[16], None, size_pt=names_size, paragraphs=names or [""])
    if not names:
        for paragraph in pprs:
            ppr = paragraph.find(qn("w:pPr"))
            numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
            if numpr is not None:
                ppr.remove(numpr)
    subcamp, deputed = _sparse_summary(rows)
    for index in range(len(sparse)):
        if sparse[index]._tc.tcPr.find(qn("w:vMerge")) is not None and index not in (3, 4):
            continue                      # خلايا استمرار الدمج تبقى كما في القالب
        _write_cell(sparse[index], subcamp if index == 3 else deputed if index == 4 else "",
                    size_pt=SUMMARY_SIZE_PT)
    return names


def _grid_widths_pt(table):
    return [int(col.get(qn("w:w"))) / 20 for col in table._tbl.tblGrid.findall(qn("w:gridCol"))]


def _row_min_pt(tr):
    height = tr.find(qn("w:trPr") + "/" + qn("w:trHeight"))
    return int(height.get(qn("w:val"))) / 20 if height is not None else 0.0


def _cell_span(tc):
    node = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
    return int(node.get(qn("w:val"))) if node is not None else 1


def _lines(meter, text, width_pt, size_pt):
    inner = max(1.0, width_pt - CELL_MARGIN_PT)
    return len(meter.wrap_lines(text, inner, size_pt, True, FONT_NAME))


def _pick_size(meter, text, width_pt):
    """9pt كالأصل؛ لو احتاج أكثر من سطرين صغّر هذه الخلية وحدها إلى 8 ثم 7pt."""
    for size in SHRINK_STEPS_PT:
        if _lines(meter, text, width_pt, size) <= MAX_CELL_LINES:
            return size
    return SHRINK_STEPS_PT[-1]


def _roster_plan(table, values):
    """values: {(row_index, tc_index): text}. يرجع {(row, tc): size_pt}."""
    meter = TextMeter()
    grid = _grid_widths_pt(table)
    plan = {}
    for (row_index, tc_index), text in values.items():
        tcs = table.rows[row_index]._tr.tc_lst
        start = sum(_cell_span(tc) for tc in tcs[:tc_index])
        width = sum(grid[start:start + _cell_span(tcs[tc_index])])
        plan[(row_index, tc_index)] = _pick_size(meter, text, width)
    return plan, meter, grid


def _estimate_total(document, values, plan, names, meter, grid, names_size=SUMMARY_SIZE_PT):
    """تقدير ارتفاع الجدولين (نقطة) بالمقاسات المخطّطة — للقرار فقط، لا يُكتب في الملف."""
    roster, summary = document.tables[:2]
    total = 0.0
    for row_index, row in enumerate(roster.rows):
        need = _row_min_pt(row._tr)
        tcs = row._tr.tc_lst
        for tc_index, tc in enumerate(tcs):
            key = (row_index, tc_index)
            if key not in values:
                continue
            start = sum(_cell_span(t) for t in tcs[:tc_index])
            width = sum(grid[start:start + _cell_span(tc)])
            size = plan[key]
            need = max(need, _lines(meter, values[key], width, size) * size * REAL_LINE_FACTOR + 1)
        total += need
    sgrid = _grid_widths_pt(summary)
    base = sum(_row_min_pt(row._tr) for row in summary.rows)
    names_width = sgrid[-1]
    names_need = sum(_lines(meter, name, names_width - 9, names_size) for name in names) \
        * names_size * REAL_LINE_FACTOR + 1
    return total + max(base, names_need)


def _plan_fit(document, values, names):
    roster = document.tables[0]
    plan, meter, grid = _roster_plan(roster, values)
    usable = usable_height_pt(document.sections[0])
    names_size = SUMMARY_SIZE_PT
    total = _estimate_total(document, values, plan, names, meter, grid, names_size)
    # قائمة صافي طويلة جدًا (لم تحدث في الملفات الحقيقية): صغّر خلية الأسماء وحدها
    # أولًا لأنها الأكثر تأثيرًا على الارتفاع، ثم أطول خلايا الجدول الأول.
    while total / usable > FIT_TARGET and names_size > 6:
        names_size -= 1
        total = _estimate_total(document, values, plan, names, meter, grid, names_size)
    # لو الصفحة ما كفتش: صغّر أطول الخلايا واحدة واحدة (دائمًا خلية لا الجدول).
    while total / usable > FIT_TARGET:
        candidates = []
        for key, size in plan.items():
            if size <= SHRINK_STEPS_PT[-1]:
                continue
            tcs = roster.rows[key[0]]._tr.tc_lst
            start = sum(_cell_span(t) for t in tcs[:key[1]])
            width = sum(grid[start:start + _cell_span(tcs[key[1]])])
            lines = _lines(meter, values[key], width, size)
            if lines > 1:
                candidates.append((lines * size, key))
        if not candidates:
            break
        plan[max(candidates)[1]] -= 1
        total = _estimate_total(document, values, plan, names, meter, grid, names_size)
    ratio = total / usable
    if ratio > FIT_LIMIT:
        raise PageFitError("تعذر احتواء يومية الضباط في صفحة واحدة")
    LAST_FIT.clear()
    LAST_FIT.update(scale=1.0, ratio=ratio, total_pt=total, usable_pt=usable,
                    shrunk=sorted((k, v) for k, v in plan.items() if v < ROSTER_SIZE_PT),
                    names_size=names_size)
    return plan, names_size


def build_docx(duty):
    document = Document(TEMPLATE)
    _set_title(document, duty["date"])

    table = document.tables[0]
    grouped = dict(_grouped_rows(duty["rows"]))
    # نجمع نصوص كل الخلايا أولًا ثم نقرر مقاس كل خلية على حدة ثم نكتب.
    values = {}
    number = 0
    for section in OFFICER_SECTIONS:
        for offset in range(ROSTER_SLOTS[section]):
            item = grouped.get(section, [])[offset] if offset < len(grouped.get(section, [])) else None
            row_index = offset + (1 if section == SECTION_FORCE else
                                  26 if section == SECTION_GUARDS else 35)
            number += 1
            texts = [f"{number}-", "", "", "", "", ""]
            if item:
                texts = [f"{number}-", RANK_SHORT.get(item.get("role"), item.get("role") or ""),
                         item.get("name") or "", item.get("post") or "", _daily_text(item),
                         _rest_text(item)]
            for tc_index, text in enumerate(texts):
                values[(row_index, tc_index)] = text
    names = _net_names(duty["rows"])
    plan, names_size = _plan_fit(document, values, names)
    for (row_index, tc_index), text in values.items():
        cells = _physical_cells(table.rows[row_index])
        _write_cell(cells[tc_index], text, size_pt=plan[(row_index, tc_index)])

    _fill_summary(document.tables[1], duty["summary"], duty["rows"], names_size)
    order_xml_children(document)
    buf = BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf
