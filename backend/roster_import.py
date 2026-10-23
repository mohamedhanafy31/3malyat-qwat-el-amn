"""قراءة كشوف Word واستبدال القوة النشطة مع إبقاء التاريخ صالحًا."""
from datetime import date
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
import re

from docx import Document

from .store import reserve_id


_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
_ALEF = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه"})
MAX_FILE_BYTES = 4 * 1024 * 1024


def clean(value):
    return " ".join(str(value or "").translate(_DIGITS).replace("\xa0", " ").split())


def name_key(value):
    value = clean(value).translate(_ALEF).lower()
    return re.sub(r"[^\w\u0600-\u06ff]", "", value)


def source_hash(blob):
    return sha256(blob).hexdigest()


def effective_date(filename, supplied=""):
    value = clean(supplied)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        try:
            date.fromisoformat(value)
            return value
        except ValueError:
            pass
    match = re.search(r"(20\d{2})[-_](\d{1,2})[-_](\d{1,2})", clean(filename))
    if match:
        try:
            return date(int(match[1]), int(match[2]), int(match[3])).isoformat()
        except ValueError:
            pass
    return date.today().isoformat()


def _ensure_size(blob):
    if not blob:
        raise ValueError("الملف فارغ.")
    if len(blob) > MAX_FILE_BYTES:
        raise ValueError("حجم ملف الكشف أكبر من 4 ميجابايت.")


def parse_conscripts(blob):
    _ensure_size(blob)
    try:
        document = Document(BytesIO(blob))
    except Exception as exc:
        raise ValueError("تعذر قراءة ملف المجندين. اختر ملف DOCX صحيحًا.") from exc
    if not document.tables:
        raise ValueError("لا يوجد جدول أسماء في ملف المجندين.")
    names = []
    for table in document.tables:
        for row in table.rows:
            cells = [clean(cell.text) for cell in row.cells]
            name = next((cell for cell in reversed(cells) if cell), "")
            if not name or name_key(name) in {"الاسم", "اسم"}:
                continue
            names.append(name)
    duplicates = sorted(name for name in set(names) if names.count(name) > 1)
    if duplicates:
        raise ValueError("يوجد أسماء مجندين مكررة في الملف: " + "، ".join(duplicates[:5]))
    if not names:
        raise ValueError("لم يتم العثور على أسماء مجندين في الملف.")
    return [{"row": index, "name": name, "phone": "", "is_driver": True}
            for index, name in enumerate(names, 1)]


class _WordHtmlTable(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell, self.in_cell = [], [], [], False

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "tr":
            self.row = []
        elif tag in ("td", "th"):
            self.cell, self.in_cell = [], True

    def handle_data(self, data):
        if self.in_cell:
            self.cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ("td", "th") and self.in_cell:
            self.row.append(clean(" ".join(self.cell)))
            self.in_cell = False
        elif tag == "tr" and self.row:
            self.rows.append(self.row)


def _canonical_role(value):
    words = clean(value).split()
    if not words:
        return ""
    words[0] = {"امين": "أمين", "جندي": "جندي"}.get(words[0], words[0])
    words = [{"اول": "أول", "ثاني": "ثان"}.get(word, word) for word in words]
    return " ".join(words)


def parse_personnel(blob):
    _ensure_size(blob)
    text = None
    for encoding in ("utf-8-sig", "windows-1256"):
        try:
            text = blob.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None or "<table" not in text.lower():
        raise ValueError("تعذر قراءة ملف الأفراد. يجب أن يكون ملف DOC صادرًا من Word بنفس بنية الكشف.")
    parser = _WordHtmlTable()
    parser.feed(text)
    header_at = next((i for i, row in enumerate(parser.rows)
                      if len(row) >= 7 and name_key(row[2]) == name_key("الاسم")), None)
    if header_at is None:
        raise ValueError("لم يتم العثور على أعمدة كشف الأفراد المطلوبة.")
    out = []
    for source_row, values in enumerate(parser.rows[header_at + 1:], 1):
        if len(values) < 7:
            continue
        number, role, name, post, address, phone, code = map(clean, values[:7])
        if not name:
            continue
        out.append({"row": int(number) if number.isdigit() else source_row,
                    "role": _canonical_role(role), "name": name, "post": post,
                    "address": address, "phone": phone, "code": code})
    if not out:
        raise ValueError("كشف الأفراد لا يحتوي على سجلات قابلة للاستيراد.")
    duplicate_names = [name for name, count in _counts(name_key(r["name"]) for r in out).items()
                       if count > 1]
    if duplicate_names:
        raise ValueError("يوجد أسماء أفراد مكررة في الملف.")
    return out


def _counts(values):
    out = {}
    for value in values:
        if value:
            out[value] = out.get(value, 0) + 1
    return out


def _phone_keys(row):
    values = [row.get("phone", ""), *(row.get("other_phones") or [])]
    return {re.sub(r"\D", "", clean(value)) for value in values
            if re.sub(r"\D", "", clean(value))}


def _personnel_matches(current, incoming_rows):
    """Match a new official roster to existing identities.

    Names in the October file were expanded from short names to full names,
    so identity cannot depend on exact name alone. Official code wins when it
    exists in both sources, then any current/alternate phone, exact normalized
    name, and finally one unambiguous prefix-name match.
    """
    available = {row.get("id"): row for row in current}
    matched = []
    for incoming in incoming_rows:
        candidates = []
        code = clean(incoming.get("code"))
        if code:
            candidates = [row for row in available.values()
                          if clean(row.get("code")) == code]
        if len(candidates) != 1:
            phones = _phone_keys(incoming)
            candidates = [row for row in available.values()
                          if phones and phones.intersection(_phone_keys(row))]
        if len(candidates) != 1:
            key = name_key(incoming.get("name"))
            candidates = [row for row in available.values()
                          if key and name_key(row.get("name")) == key]
        if len(candidates) != 1:
            key = name_key(incoming.get("name"))
            candidates = [row for row in available.values()
                          if key and min(len(key), len(name_key(row.get("name")))) >= 12
                          and (key.startswith(name_key(row.get("name")))
                               or name_key(row.get("name")).startswith(key))]
        found = candidates[0] if len(candidates) == 1 else None
        matched.append(found)
        if found:
            available.pop(found.get("id"), None)
    return matched


def _conscript_plan(data, rows):
    current = list(data.get("conscripts") or [])
    by_name = {name_key(row.get("name")): row for row in current}
    matched = {by_name[name_key(row["name"])]["id"] for row in rows
               if name_key(row["name"]) in by_name}
    removed = [row for row in current if row.get("id") not in matched]
    return {"kind": "conscripts", "rows": rows, "current": len(current),
            "matched": len(matched), "added": len(rows) - len(matched),
            "removed": len(removed), "removed_rows": removed, "warnings": []}


def _personnel_plan(data, rows):
    current = list(data.get("personnel") or [])
    active = [row for row in current if row.get("status", "active") == "active"]
    matched_ids = {found["id"] for found in _personnel_matches(active, rows) if found}
    missing_codes = [row for row in rows if not row["code"]]
    duplicate_codes = [code for code, count in _counts(r["code"] for r in rows).items()
                       if count > 1]
    warnings = []
    if missing_codes:
        warnings.append(f"{len(missing_codes)} سجل بلا رقم شرطة؛ سيُحفظ له رقم مؤقت واضح حتى لا يُفقد من الكشف.")
    if duplicate_codes:
        warnings.append("أرقام شرطة مكررة في المصدر: " + "، ".join(duplicate_codes[:5]))
    return {"kind": "personnel", "rows": rows, "current": len(active),
            "matched": len(matched_ids), "added": len(rows) - len(matched_ids),
            "removed": len(active) - len(matched_ids), "warnings": warnings}


def preview(data, kind, blob, filename, supplied_date=""):
    rows = parse_conscripts(blob) if kind == "conscripts" else parse_personnel(blob)
    plan = _conscript_plan(data, rows) if kind == "conscripts" else _personnel_plan(data, rows)
    plan.update(source_hash=source_hash(blob), effective_date=effective_date(filename, supplied_date),
                filename=filename, total=len(rows))
    plan.pop("removed_rows", None)
    return plan


def _detach_removed_conscripts(data, removed):
    names = {row.get("id"): row.get("name", "") for row in removed}
    touched = 0
    for sheet in (data.get("day_campaign") or {}).values():
        for section in ("lorries", "services", "daily"):
            for entry in (sheet.get(section) or {}).values():
                driver = entry.get("driver") or {}
                if driver.get("kind") == "conscript" and driver.get("id") in names:
                    entry["driver"] = {"kind": "text", "id": "", "name": names[driver["id"]]}
                    touched += 1
    return touched


def _apply_conscripts(data, rows):
    plan = _conscript_plan(data, rows)
    by_name = {name_key(row.get("name")): row for row in data.get("conscripts") or []}
    replacement = []
    for row in rows:
        old = by_name.get(name_key(row["name"]))
        replacement.append({"id": old["id"] if old else reserve_id(data, "CON", replacement),
                            "name": row["name"], "phone": row["phone"], "is_driver": True})
    detached = _detach_removed_conscripts(data, plan["removed_rows"])
    data["conscripts"] = replacement
    return {**{key: plan[key] for key in ("current", "matched", "added", "removed")},
            "total": len(replacement), "detached_history_refs": detached}


def _temporary_code(row):
    return "غير مسجل-" + sha256(name_key(row["name"]).encode("utf-8")).hexdigest()[:8].upper()


def _apply_personnel(data, rows, on_date):
    plan = _personnel_plan(data, rows)
    all_rows = list(data.get("personnel") or [])
    active = [row for row in all_rows if row.get("status", "active") == "active"]
    archive = [row for row in all_rows if row.get("status", "active") != "active"]
    matches = _personnel_matches(active, rows)
    matched_ids, replacement = set(), []
    for incoming, old in zip(rows, matches):
        if old:
            matched_ids.add(old["id"])
        base = dict(old or {})
        base.update({"id": old["id"] if old else reserve_id(data, "IND", all_rows + replacement),
                     "name": incoming["name"], "role": incoming["role"],
                     "code": incoming["code"] or _temporary_code(incoming),
                     "phone": incoming["phone"], "post": incoming["post"],
                     "address": incoming["address"], "status": "active",
                     "join_date": (old or {}).get("join_date") or on_date,
                     "leave_date": "", "leave_reason": ""})
        base.setdefault("other_phones", [])
        base.setdefault("history", [])
        replacement.append(base)
    # leave_date هو آخر يوم كان فيه الشخص على القوة. نجعله يوم سريان الكشف
    # نفسه كي تظل تكليفات ذلك اليوم التاريخية قابلة للعرض، ويخرج من قوة
    # اليوم التالي تلقائيًا.
    leave_day = on_date
    for old in active:
        if old.get("id") in matched_ids:
            continue
        archived = dict(old)
        archived.update(status="archived", leave_date=max(old.get("join_date") or leave_day, leave_day),
                        leave_reason=f"استبدال كشف الأفراد بتاريخ {on_date}")
        archive.append(archived)
    data["personnel"] = replacement + archive
    return {**{key: plan[key] for key in ("current", "matched", "added", "removed")},
            "total": len(replacement), "archive_total": len(archive)}


def apply(data, kind, blob, filename, supplied_date=""):
    rows = parse_conscripts(blob) if kind == "conscripts" else parse_personnel(blob)
    on_date = effective_date(filename, supplied_date)
    result = (_apply_conscripts(data, rows) if kind == "conscripts"
              else _apply_personnel(data, rows, on_date))
    result.update(kind=kind, effective_date=on_date, source_hash=source_hash(blob))
    return result
