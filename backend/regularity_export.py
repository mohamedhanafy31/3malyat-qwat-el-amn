"""تصدير يومية انتظام الخدمات من قالب الأرشيف الرسمي."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from .board_export import WEEKDAY_BY_INDEX, order_xml_children
from .export_fit import PageFitError, TextMeter, estimate_table_height

TEMPLATE = Path(__file__).with_name("docx_templates") / "service_regularity.docx"
LAST_FIT = {}
# أصغر خط 7pt: ستة أيام (23-29 خدمة) ما كانتش بتدخل صفحة واحدة عند 8pt،
# والأرشيف نفسه بيكتبها بـ8-9pt مع نصوص أقصر. الصفحة الواحدة إلزامية.
_SIZES = (12, 11, 10, 9, 8.5, 8, 7.5, 7)
# الـ74 القديمة كانت بتحسب فراغ العنوان/التوقيعات بتحفظ زايد. مع صفوف
# الملاحظات الطويلة بتاعة 023 الفرق ده كان بيرفض 22/8 عند 7pt بـحوالي 3pt،
# مع إن الصفوف نفسها سطر واحد. القيمة دي مقاسة على القالب الفعلي، وفضلنا
# نفصلها عن حساب الصفوف عشان يفضل محافظ وواضح.
_FIXED_PAGE_CHROME_PT = 70


def _set_cell(cell, text, size):
    cell.text = str(text or "")
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.name = "Arial"
            rpr = run._r.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                from docx.oxml import OxmlElement
                fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                fonts.set(qn("w:" + attr), "Arial")
            for tag in ("w:sz", "w:szCs"):
                node = rpr.find(qn(tag))
                if node is None:
                    from docx.oxml import OxmlElement
                    node = OxmlElement(tag); rpr.append(node)
                node.set(qn("w:val"), str(int(round(size * 2))))


def _row(table, prototype, values, size, note=False):
    tr = deepcopy(prototype)
    cells = table._tbl.__class__  # keep lxml namespace registration loaded
    del cells
    physical = tr.findall(qn("w:tc"))
    if note:
        _set_cell(type(table.rows[0].cells[0])(physical[0], table), values, size)
        return tr
    for tc, text in zip(physical, values):
        # _Cell needs a parent table proxy; assigning XML directly through a
        # temporary Document cell is unnecessary, so use the XML text nodes.
        texts = tc.findall(".//" + qn("w:t"))
        if texts:
            texts[0].text = str(text or "")
            for extra in texts[1:]: extra.text = ""
        else:
            from docx.oxml import OxmlElement
            p = tc.find(qn("w:p"))
            r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = str(text or "")
            r.append(t); p.append(r)
        for node in tc.iter(qn("w:sz")):
            node.set(qn("w:val"), str(int(round(size * 2))))
        for node in tc.iter(qn("w:szCs")):
            node.set(qn("w:val"), str(int(round(size * 2))))
    return tr


def _set_row_text(tr, values, size, note=False):
    cells = tr.findall(qn("w:tc"))
    if note:
        values = [values] * len(cells)
    for tc, text in zip(cells, values):
        texts = tc.findall(".//" + qn("w:t"))
        if not texts:
            continue
        texts[0].text = str(text or "")
        for node in texts[1:]: node.text = ""
        for node in tc.iter(qn("w:sz")):
            node.set(qn("w:val"), str(int(round(size * 2))))
        for node in tc.iter(qn("w:szCs")):
            node.set(qn("w:val"), str(int(round(size * 2))))


def _signature_boxes(document, payload):
    names = payload.get("signatories") or {}
    labels = {"مدير إدارة": names.get("director") or {},
              "وكيل إدارة": names.get("deputy") or {},
              "رئيس ق": names.get("operations") or {}}
    for box in document.part.element.iter(qn("w:txbxContent")):
        joined = "".join(node.text or "" for node in box.iter(qn("w:t")))
        item = next((v for key, v in labels.items() if key in joined), None)
        if item is None:
            continue
        if "مدير" in joined:
            text = f"مدير إدارة قـــوات الأمن\n{item.get('rank', '')} /\n({item.get('name', '')})"
        elif "وكي" in joined:
            text = f"وكيـل إدارة قوات الأمن\n{item.get('rank', '')}/\n({item.get('name', '')})"
        else:
            text = f"رئيس قســم العمليات\n{item.get('rank', '')}/\n({item.get('name', '')})"
        nodes = list(box.iter(qn("w:t")))
        if nodes:
            nodes[0].text = text
            for node in nodes[1:]: node.text = ""


def _fit(payload, table1, table2):
    rows1 = [[{"text": "م", "size_pt": 10}, {"text": r["service"], "size_pt": 10},
              {"text": r["people"], "size_pt": 10}, {"text": r["phone"], "size_pt": 10},
              {"text": r["vehicle"], "size_pt": 10}, {"text": r["place"], "size_pt": 10},
              {"text": r["time"], "size_pt": 10}, {"text": r["count"], "size_pt": 10}]
             for r in payload.get("service_rows") or []]
    rows2 = [[{"text": "", "size_pt": 10}, {"text": r["rank"], "size_pt": 10},
              {"text": r["name"], "size_pt": 10}, {"text": r["service"], "size_pt": 10}]
             for r in payload.get("officer_rows") or []]
    rows2 += [[{"text": n, "size_pt": 10, "colspan": 4}] for n in payload.get("notes") or []]
    w1 = [max(1, (cell.width.pt if cell.width else 45)) for cell in table1.columns]
    w2 = [max(1, (cell.width.pt if cell.width else 70)) for cell in table2.columns]
    meter = TextMeter()
    usable = table1._parent.part.document.sections[0].page_height.pt - 90
    for size in _SIZES:
        a = [[{**cell, "size_pt": size} for cell in row] for row in rows1]
        b = [[{**cell, "size_pt": size} for cell in row] for row in rows2]
        h1 = estimate_table_height(a, w1, meter)[0] if a else 0
        h2 = estimate_table_height(b, w2, meter)[0] if b else 0
        if h1 + h2 + _FIXED_PAGE_CHROME_PT <= usable * .96:
            return size, h1 + h2
    raise PageFitError("تعذر احتواء يومية انتظام الخدمات في صفحة واحدة؛ قلل عدد الصفوف أو طول النصوص.")


def build_docx(payload):
    document = Document(TEMPLATE)
    table1, table2 = document.tables[:2]
    fit_size, total = _fit(payload, table1, table2)
    proto1 = deepcopy(table1._tbl.tr_lst[1])
    proto2 = deepcopy(table2._tbl.tr_lst[1])
    note_proto = deepcopy(table2._tbl.tr_lst[-1])
    for tr in list(table1._tbl.tr_lst[1:]): table1._tbl.remove(tr)
    for tr in list(table2._tbl.tr_lst[1:]): table2._tbl.remove(tr)
    for index, row in enumerate(payload.get("service_rows") or [], 1):
        # خانة «م» في القالب ترقيم تلقائي (numbering list) زي الأصل — كتابة
        # الرقم كمان كانت بتطلّعه مرتين («1.» وتحتها «1»).
        tr = _row(table1, proto1, ["", row["service"], row["people"], row["phone"],
                                   row["vehicle"], row["place"], row["time"], row["count"]], fit_size)
        table1._tbl.append(tr)
    for row in payload.get("officer_rows") or []:
        tr = _row(table2, proto2, ["", row["rank"], row["name"], row["service"]], fit_size)
        table2._tbl.append(tr)
    for note in payload.get("notes") or []:
        tr = deepcopy(note_proto)
        _set_row_text(tr, note, fit_size, note=True)
        table2._tbl.append(tr)
    day = payload.get("day") or payload.get("date")
    if day:
        d = date.fromisoformat(day)
        title = f"يومية انتظام الخدمات عن يوم {WEEKDAY_BY_INDEX[d.weekday()]} الموافق {d.day}/{d.month}/{d.year}م"
        for paragraph in document.sections[0].header.paragraphs:
            if paragraph.text or paragraph.runs:
                if paragraph.runs:
                    paragraph.runs[0].text = title
                    for run in paragraph.runs[1:]: run.text = ""
                break
    _signature_boxes(document, payload)
    order_xml_children(document)
    LAST_FIT.clear(); LAST_FIT.update(size_pt=fit_size, total_pt=total, page_count=1)
    buf = BytesIO(); document.save(buf); buf.seek(0); return buf
