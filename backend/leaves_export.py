"""تصدير كشوف راحات الضباط بنفس هندسة النماذج الورقية المعتمدة."""
from copy import deepcopy
from datetime import date, timedelta
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import (WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE,
                             WD_TABLE_ALIGNMENT, WD_TABLE_DIRECTION)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .board_export import _set_rtl_doc, set_run_font
from .labels import RANK_SHORT, first_words, officer_label
from .people import effective
from .repo import PeopleRepo


FONT = "Sultan bold"
DATE_FONT = "Times New Roman"
HEADER_FILL = "D9D9D9"
MONTHS = ["", "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
          "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
WEEKDAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
WEEKLY_GROUPS = [
    ("السبت", "الأحد"), ("الأحد", "الاثنين"), ("الاثنين", "الثلاثاء"),
    ("الثلاثاء", "الأربعاء"), ("الأربعاء", "الخميس"), ("الخميس", "الجمعة"),
    ("الجمعة", "السبت"),
]


def _insert_before(parent, node, successors):
    successor_tags = {qn(f"w:{name}") for name in successors}
    for index, child in enumerate(parent):
        if child.tag in successor_tags:
            parent.insert(index, node)
            return
    parent.append(node)


def _rtl_paragraph(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        _insert_before(ppr, OxmlElement("w:bidi"), (
            "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
            "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
        ))


def _shade(cell, fill=HEADER_FILL):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        _insert_before(tc_pr, shd, ("noWrap", "tcMar", "textDirection", "tcFitText",
                                    "vAlign", "hideMark", "headers", "cellIns", "cellDel"))
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def _borders(table, size=6, outer=None):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        _insert_before(table._tbl.tblPr, borders, (
            "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
        ))
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(outer if outer and edge in {"top", "left", "bottom", "right"} else size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)


def _cell_borders(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        _insert_before(tc_pr, borders, ("shd", "noWrap", "tcMar", "textDirection",
                                        "tcFitText", "vAlign", "hideMark"))
    for edge, size in edges.items():
        old = borders.find(qn(f"w:{edge}"))
        if old is not None:
            borders.remove(old)
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)


def _set_width(cell, width_cm):
    tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        cell._tc.get_or_add_tcPr().append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(round(Cm(width_cm).emu / 635)))


def _set_cell(cell, text="", *, size=14, font=FONT, fill=None, bold=True):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if fill:
        _shade(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _rtl_paragraph(paragraph)
    run = paragraph.add_run(str(text or ""))
    set_run_font(run, font, size, bold)
    run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _fixed_row(row, height_cm):
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_grid(table, widths):
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(round(Cm(width).emu / 635)))
        grid.append(col)
    table.autofit = False
    table.table_direction = WD_TABLE_DIRECTION.RTL
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            _set_width(cell, width)


def _landscape_document(left=1.5, right=1.5, top=.65, bottom=.65):
    doc = Document()
    _set_rtl_doc(doc)
    for doc_section in doc.sections:
        sect_pr = doc_section._sectPr
        bidi = sect_pr.find(qn("w:bidi"))
        if bidi is not None:
            sect_pr.remove(bidi)
            _insert_before(sect_pr, bidi, ("rtlGutter", "docGrid", "printerSettings", "sectPrChange"))
    zoom = doc.settings._element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin, section.right_margin = Cm(left), Cm(right)
    section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
    return doc


def _save(doc):
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def _short_name(person):
    return str(person.get("short_name") or first_words(person.get("name"), 2)).strip()


def _rank(role):
    return RANK_SHORT.get(str(role or "").strip(), str(role or "").strip())


def _date_text(value):
    if not value:
        return ""
    parsed = date.fromisoformat(value)
    return f"{parsed.year}/{parsed.month}/{parsed.day}م"


def _latest_plan(leaves, officer_id, rest_system, on):
    rows = [lv for lv in leaves if lv.get("person_id") == officer_id
            and lv.get("type") == rest_system and lv.get("start") and lv.get("end")]
    if not rows:
        return None
    rows.sort(key=lambda lv: (lv["start"], lv.get("id", "")))
    not_finished = [lv for lv in rows if (lv.get("return_date") or lv["end"]) >= on]
    return deepcopy(not_finished[0] if not_finished else rows[-1])


def monthly_report(data, on):
    on_iso = on.isoformat()
    groups = {"شهرية": [], "نصف شهرية": []}
    leaves = data.get("leaves") or []
    people = PeopleRepo(data)
    force_ids = people.ids_on_force(on_iso, "officers")
    for raw in people.bucket("officers", "active") + people.bucket("officers", "archive"):
        person_id = raw.get("id", "")
        if person_id not in force_ids:
            continue
        eff = effective(raw, on_iso)
        system = eff.get("rest_system", "")
        if system not in groups:
            continue
        leave = _latest_plan(leaves, person_id, system, on_iso)
        start = leave.get("start", "") if leave else ""
        end = leave.get("end", "") if leave else ""
        returned = (leave.get("return_date") or
                    ((date.fromisoformat(end) + timedelta(days=1)).isoformat() if end else "")) if leave else ""
        taqseera = (date.fromisoformat(start) - timedelta(days=1)).isoformat() if start else ""
        groups[system].append({
            "id": person_id, "name": _short_name(raw), "role": _rank(eff.get("role")),
            "taqseera": taqseera, "start": start, "end": end, "return_date": returned,
        })
    starts = [row["start"] for rows in groups.values() for row in rows if row["start"]]
    if starts:
        first, last = date.fromisoformat(min(starts)), date.fromisoformat(max(starts))
        month_label = MONTHS[first.month] if first.month == last.month else f"{MONTHS[first.month]} - {MONTHS[last.month]}"
        year_label = str(last.year) if first.year == last.year else f"{first.year} - {last.year}"
    else:
        month_label, year_label = MONTHS[on.month], str(on.year)
    return {"on": on_iso, "month_label": month_label, "year_label": year_label, "groups": groups}


def weekly_report(data, on):
    on_iso = on.isoformat()
    groups = {rest_day: [] for _, rest_day in WEEKLY_GROUPS}
    people = PeopleRepo(data)
    force_ids = people.ids_on_force(on_iso, "officers")
    for raw in people.bucket("officers", "active") + people.bucket("officers", "archive"):
        person_id = raw.get("id", "")
        if person_id not in force_ids:
            continue
        eff = effective(raw, on_iso)
        rest_day = eff.get("rest_day", "")
        if eff.get("rest_system") != "أسبوعية" or rest_day not in groups:
            continue
        label = officer_label({**raw, "role": eff.get("role", "")})
        groups[rest_day].append({"id": person_id, "label": label})
    return {"on": on_iso, "groups": groups}


def _title_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_grid(table, [17.2])
    _borders(table, size=8)
    _set_cell(table.cell(0, 0), text, size=22, fill=HEADER_FILL)
    _fixed_row(table.rows[0], 1.15)
    return table


def _spacer(doc, points):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(points)


def _monthly_table(doc, rows):
    widths = [.9, 1.9, 4.9, 4.3, 4.1, 4.2, 4.7]
    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_grid(table, widths)
    _borders(table, size=6)
    headers = ["م", "الرتبة", "الاســــــــــم", "التقصيرة", "الراحة من", "إلى", "العودة"]
    for cell, value in zip(table.rows[0].cells, headers):
        _set_cell(cell, value, size=12, fill=HEADER_FILL)
    _fixed_row(table.rows[0], .68)
    for index, item in enumerate(rows, 1):
        row = table.add_row()
        values = [f"{index}.", item["role"], item["name"], _date_text(item["taqseera"]),
                  _date_text(item["start"]), _date_text(item["end"]), _date_text(item["return_date"])]
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            _set_cell(cell, value, size=11, font=DATE_FONT if col >= 3 else FONT)
        _fixed_row(row, .69)
    return table


def build_monthly_docx(data, on):
    report = monthly_report(data, on)
    doc = _landscape_document(top=.75, bottom=.65)
    groups = report["groups"]
    title_suffix = f"عن شهر {report['month_label']} {report['year_label']}"
    _title_box(doc, f"كشف الراحات الشهرية {title_suffix}")
    _spacer(doc, 10)
    _monthly_table(doc, groups["شهرية"])
    _spacer(doc, 11)
    _title_box(doc, f"كشف الراحات النصف شهرية {title_suffix}")
    _spacer(doc, 5)
    _monthly_table(doc, groups["نصف شهرية"])
    return _save(doc)


def _weekly_title(doc, on):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)
    ppr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "000000")
    border.append(bottom)
    _insert_before(ppr, border, ("shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
                                  "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
                                  "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "jc"))
    _rtl_paragraph(paragraph)
    run = paragraph.add_run(f"كشف الراحات الاسبوعية للضباط بتاريخ {on.year}/{on.month}/{on.day}")
    set_run_font(run, FONT, 26, True)
    run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _weekly_header(table, pairs, *, blank_after=0):
    top, days = table.add_row(), table.add_row()
    for group_index, (taqseera_day, rest_day) in enumerate(pairs):
        offset = group_index * 3
        number = top.cells[offset].merge(days.cells[offset])
        _set_cell(number, "م", size=12, fill=HEADER_FILL)
        _set_cell(top.cells[offset + 1], "تقصيرة", size=12, fill=HEADER_FILL)
        _set_cell(top.cells[offset + 2], "راحة", size=12, fill=HEADER_FILL)
        _set_cell(days.cells[offset + 1], taqseera_day, size=12, fill=HEADER_FILL)
        _set_cell(days.cells[offset + 2], rest_day, size=12, fill=HEADER_FILL)
    if blank_after:
        start = len(pairs) * 3
        blank = top.cells[start].merge(days.cells[-1])
        _set_cell(blank, "", fill=HEADER_FILL)
    for row in (top, days):
        _fixed_row(row, .75)
        for cell in row.cells:
            _cell_borders(cell, top=20)


def _weekly_body(table, pairs, groups, slots=6, *, blank_after=0):
    for index in range(slots):
        row = table.add_row()
        for group_index, (_, rest_day) in enumerate(pairs):
            offset = group_index * 3
            item = groups.get(rest_day, [])[index] if index < len(groups.get(rest_day, [])) else None
            _set_cell(row.cells[offset], f"{index + 1}.", size=12)
            name = row.cells[offset + 1].merge(row.cells[offset + 2])
            _set_cell(name, item["label"] if item else "", size=12)
        if blank_after:
            start = len(pairs) * 3
            blank = row.cells[start]
            for cell in row.cells[start + 1:]:
                blank = blank.merge(cell)
            _set_cell(blank, "", fill=HEADER_FILL)
        _fixed_row(row, .675)


def build_weekly_docx(data, on):
    report = weekly_report(data, on)
    doc = _landscape_document(left=2.2, right=1.1, top=.45, bottom=.55)
    _weekly_title(doc, on)
    widths = [1.1, 4.25, 3.4] * 3
    table = doc.add_table(rows=0, cols=9)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    _set_grid(table, widths)
    _borders(table, size=5, outer=20)
    for start in (0, 3):
        pairs = WEEKLY_GROUPS[start:start + 3]
        _weekly_header(table, pairs)
        slots = max(6, *(len(report["groups"].get(rest_day, [])) for _, rest_day in pairs))
        _weekly_body(table, pairs, report["groups"], slots)
    pair = WEEKLY_GROUPS[6:7]
    _weekly_header(table, pair, blank_after=6)
    slots = max(2, len(report["groups"].get("السبت", [])))
    _weekly_body(table, pair, report["groups"], slots, blank_after=6)
    return _save(doc)
