"""كل جداول التصدير: bold + RTL، والعناصر بترتيب المخطط."""
import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell


@pytest.mark.parametrize("kind", ["board", "duty", "afraad", "campaign"])
def test_every_table_is_bold_and_rtl(client, kind):
    response = client.get(f"/api/{kind}/2026-03-01/export.docx")
    assert response.status_code == 200
    document = Document(io.BytesIO(response.data))
    assert document.tables
    runs = 0
    for table in document.tables:
        assert table._tbl.tblPr.find(qn("w:bidiVisual")) is not None
        for row in table.rows:
            for tc in row._tr.tc_lst:
                for paragraph in _Cell(tc, table).paragraphs:
                    assert paragraph._p.pPr.find(qn("w:bidi")) is not None
                    for run in paragraph.runs:
                        runs += 1
                        rpr = run._r.rPr
                        assert run.bold is True and rpr.find(qn("w:bCs")).get(qn("w:val")) == "1"
                        assert rpr.find(qn("w:rtl")) is not None
    assert runs


def test_run_and_paragraph_children_follow_schema_order(client):
    document = Document(io.BytesIO(client.get("/api/duty/2026-03-01/export.docx").data))
    for rpr in document.element.body.iter(qn("w:rPr")):
        names = [child.tag.split("}")[1] for child in rpr]
        for first, second in (("rFonts", "sz"), ("b", "sz"), ("sz", "rtl")):
            if first in names and second in names:
                assert names.index(first) < names.index(second)
    for ppr in document.element.body.iter(qn("w:pPr")):
        names = [child.tag.split("}")[1] for child in ppr]
        if "bidi" in names and "jc" in names:
            assert names.index("bidi") < names.index("jc")
