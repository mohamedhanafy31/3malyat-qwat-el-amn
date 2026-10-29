from docx import Document


def _vehicle(client, **extra):
    body = {"plate_series": "ب12", "plate_number": "4997", "type": "بيك اب"}
    body.update(extra)
    return client.post("/api/vehicles", json=body).get_json()


def _person(client, name="فرد حملة", code="991", **extra):
    body = {"type": "personnel", "name": name, "code": code,
            "phone": "0101", "join_date": "2020-01-01", "role": "أمين شرطة",
            "post": "سائق"}
    body.update(extra)
    return client.post("/api/person", json=body).get_json()


def _conscript(client, name="مجند حملة"):
    return client.post("/api/conscripts", json={"name": name}).get_json()


def _assignment(client, day, name="تفتيش فيصل", section="الخدمات الطارئة"):
    return client.post(f"/api/assignments/{day}", json={
        "name": name, "kind": "خارجية", "section": section,
    }).get_json()


def _campaign(client, day):
    response = client.get(f"/api/campaign/{day}")
    assert response.status_code == 200, response.get_json()
    return response.get_json()


def test_campaign_fixed_rows_assignment_and_attach(client):
    day = "2026-07-10"
    assert len(client.get(f"/api/campaign/{day}").get_json()["lorries"]) == 4
    vehicle = _vehicle(client)
    conscript = _conscript(client)
    person = _person(client)
    assignment = _assignment(client, day)
    payload = {"vehicle_id": vehicle["id"], "driver": {"kind": "personnel", "id": person["id"]}}
    assert client.put(f"/api/campaign/{day}/lorry/LOR-3", json=payload).status_code == 200
    assert client.post(f"/api/campaign/{day}/lorry/LOR-3/service", json={"assignment_id": assignment["id"]}).status_code == 200
    view = client.get(f"/api/campaign/{day}").get_json()
    lorry = next(x for x in view["lorries"] if x["id"] == "LOR-3")
    assert lorry["name"] == "فيصل + تفتيش فيصل"
    assert not any(x["id"] == assignment["id"] for x in view["services"])
    assert client.delete(f"/api/campaign/{day}/lorry/LOR-3/service", json={"assignment_id": assignment["id"]}).status_code == 200
    assert any(x["id"] == assignment["id"] for x in client.get(f"/api/campaign/{day}").get_json()["services"])
    assert client.delete(f"/api/conscripts/{conscript['id']}").status_code == 200


def test_campaign_service_reads_legacy_vehicle_and_saves_split_people(client):
    day = "2026-07-10"
    first, second = _person(client, "فرد أول", "992"), _person(client, "فرد ثان", "993")
    assignment = _assignment(client, day, "خدمة مركبتين")
    assert client.patch(f"/api/assignments/{day}/{assignment['id']}", json={
        "personnel_ids": [first["id"], second["id"]]}).status_code == 200
    vehicle_a, vehicle_b = _vehicle(client), _vehicle(client, plate_number="4998")
    # الشكل القديم بيتقري كمركبة واحدة وبكل أفراد التكليف.
    from backend import store, campaign
    data = store.load_data(lambda view: {day})
    data.setdefault("day_campaign", {}).setdefault(day, {}).setdefault("services", {})[assignment["id"]] = {
        "vehicle_id": vehicle_a["id"]}
    assert [item["vehicle_id"] for item in campaign.service_vehicles(data, day, assignment["id"])] == [vehicle_a["id"]]
    response = client.put(f"/api/campaign/{day}/service/{assignment['id']}", json={
        "vehicles": [{"vehicle_id": vehicle_a["id"], "personnel_ids": [first["id"]]},
                     {"vehicle_id": vehicle_b["id"], "personnel_ids": [second["id"]]}]})
    assert response.status_code == 200, response.get_json()
    row = next(item for item in response.get_json()["services"] if item["id"] == assignment["id"])
    assert [item["vehicle_id"] for item in row["vehicles"]] == [vehicle_a["id"], vehicle_b["id"]]
    assert [item["personnel_ids"] for item in row["vehicles"]] == [[first["id"]], [second["id"]]]


def test_campaign_service_vehicle_validation_rejects_duplicates_and_foreign_people(client):
    day = "2026-07-10"
    person = _person(client, "فرد الخدمة", "994")
    foreign = _person(client, "فرد غريب", "995")
    assignment = _assignment(client, day, "خدمة تحقق")
    assert client.patch(f"/api/assignments/{day}/{assignment['id']}", json={
        "personnel_ids": [person["id"]]}).status_code == 200
    vehicle = _vehicle(client)
    base = f"/api/campaign/{day}/service/{assignment['id']}"
    duplicate = client.put(base, json={"vehicles": [{"vehicle_id": vehicle["id"]}, {"vehicle_id": vehicle["id"]}]})
    assert duplicate.status_code == 400 and "تكرار" in duplicate.get_json()["error"]
    foreign_response = client.put(base, json={"vehicles": [{"vehicle_id": vehicle["id"], "personnel_ids": [foreign["id"]]}]})
    assert foreign_response.status_code == 400 and "الخدمة" in foreign_response.get_json()["error"]


def test_campaign_service_resolver_returns_direct_and_lorry_plates():
    from backend import campaign
    data = {
        "vehicles": [
            {"id": "VEH-1", "plate_series": "ب12", "plate_number": "1001", "type": "بيك اب"},
            {"id": "VEH-2", "plate_series": "ب13", "plate_number": "2002", "type": "بيك اب"},
        ],
        "day_assignments": {"2026-07-10": [{"id": "AS-1", "name": "خدمة", "section": "الخدمات الطارئة",
                                               "personnel_ids": [], "officer_ids": []}]},
        "day_campaign": {"2026-07-10": {"services": {"AS-1": {"vehicles": [{"vehicle_id": "VEH-1"}]}},
                                          "lorries": {"LOR-1": {"vehicle_id": "VEH-2",
                                                                    "attached": [{"assignment_id": "AS-1"}]}}}},
        "personnel": [], "officers": [], "conscripts": [], "reference_lists": {},
    }
    rows = campaign.service_vehicles(data, "2026-07-10", "AS-1")
    assert [(row["source"], row["plate"]) for row in rows] == [
        ("direct", "ب12/1001"), ("lorry", "ب13/2002")]


def test_campaign_appointment_lorries_assign_attach_and_inherit(client):
    first, later = "2026-07-10", "2026-07-11"
    vehicle = _vehicle(client)
    person = _person(client, name="سائق تعيينات", code="994")
    assignment = _assignment(client, first, "تفتيش المديرية")
    view = _campaign(client, first)
    assert [(row["id"], row["name"]) for row in view["appointments"]] == [
        ("LOR-A1", "المديرية"), ("LOR-A2", "الأربعين"), ("LOR-A3", "الإدارة الجديدة")]
    assert view["pickers"]["types"]
    assert client.put(f"/api/campaign/{first}/lorry/LOR-A2", json={
        "vehicle_id": vehicle["id"],
        "driver": {"kind": "personnel", "id": person["id"]},
    }).status_code == 200
    assert client.post(f"/api/campaign/{first}/lorry/LOR-A3/service", json={
        "assignment_id": assignment["id"]}).status_code == 200
    view = _campaign(client, first)
    assert next(row for row in view["appointments"] if row["id"] == "LOR-A2")["vehicle"]["id"] == vehicle["id"]
    assert next(row for row in view["appointments"] if row["id"] == "LOR-A3")["name"] == "الإدارة الجديدة + تفتيش المديرية"
    assert not any(row["id"] == assignment["id"] for row in view["services"])
    inherited = next(row for row in _campaign(client, later)["appointments"] if row["id"] == "LOR-A2")
    assert inherited["inherited"] and inherited["vehicle"]["id"] == vehicle["id"]
    assert client.put(f"/api/campaign/{first}/lorry/UNKNOWN", json={}).status_code == 404


def test_campaign_attachment_is_unique_across_both_lorry_groups(client):
    day = "2026-07-10"
    assignment = _assignment(client, day, "تفتيش مشترك")
    assert client.post(f"/api/campaign/{day}/lorry/LOR-1/service", json={
        "assignment_id": assignment["id"]}).status_code == 200
    response = client.post(f"/api/campaign/{day}/lorry/LOR-A1/service", json={
        "assignment_id": assignment["id"]})
    assert response.status_code == 409 and "ملحقة" in response.get_json()["error"]
    assert not any(row["id"] == assignment["id"] for row in _campaign(client, day)["services"])


def test_campaign_appointment_references_block_vehicle_and_conscript_delete(client):
    day = "2026-07-10"
    vehicle = _vehicle(client)
    conscript = _conscript(client, "سائق تعيينات مرجعي")
    assert client.put(f"/api/campaign/{day}/lorry/LOR-A2", json={
        "vehicle_id": vehicle["id"],
        "driver": {"kind": "conscript", "id": conscript["id"]},
    }).status_code == 200
    assert client.delete(f"/api/vehicles/{vehicle['id']}").status_code == 409
    assert client.delete(f"/api/conscripts/{conscript['id']}").status_code == 409


def test_campaign_default_rows_have_revision_and_empty_values(client):
    view = _campaign(client, "2026-07-10")
    assert len(view["lorries"]) == 4
    assert view["revision"]
    assert all(row["vehicle"] is None and row["driver"] is None for row in view["lorries"])


def test_campaign_resolves_conscript_and_text_drivers(client):
    day = "2026-07-10"
    vehicle = _vehicle(client)
    conscript = _conscript(client, "مجند معروف")
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "vehicle_id": vehicle["id"],
        "driver": {"kind": "conscript", "id": conscript["id"]},
    }).status_code == 200
    assert client.put(f"/api/campaign/{day}/lorry/LOR-2", json={
        "driver": {"kind": "text", "name": "سائق حر"},
    }).status_code == 200
    rows = {row["id"]: row for row in _campaign(client, day)["lorries"]}
    assert rows["LOR-1"]["driver"]["name"] == "مجند معروف"
    assert rows["LOR-1"]["vehicle"]["plate"] == vehicle["plate"]
    assert rows["LOR-1"]["vehicle"]["type"] == vehicle["type"]
    assert rows["LOR-2"]["driver"]["name"] == "سائق حر"


def test_campaign_unknown_driver_and_vehicle_are_rejected(client):
    day = "2026-07-10"
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "vehicle_id": "VEH-999",
    }).status_code == 400
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "driver": {"kind": "conscript", "id": "CON-999"},
    }).status_code == 400


def test_campaign_clear_overrides_inherited_vehicle_and_driver(client):
    first, later = "2026-07-10", "2026-07-11"
    vehicle = _vehicle(client)
    person = _person(client)
    payload = {"vehicle_id": vehicle["id"], "driver": {"kind": "personnel", "id": person["id"]}}
    client.put(f"/api/campaign/{first}/lorry/LOR-1", json=payload)
    client.put(f"/api/campaign/{later}/lorry/LOR-1", json={"vehicle_id": "", "driver": None})
    row = next(row for row in _campaign(client, later)["lorries"] if row["id"] == "LOR-1")
    assert row["vehicle"] is None and row["driver"] is None
    row = next(row for row in _campaign(client, first)["lorries"] if row["id"] == "LOR-1")
    assert row["vehicle"]["id"] == vehicle["id"] and row["driver"]["id"] == person["id"]


def test_campaign_daily_effective_versions_and_inheritance(client):
    first, later = "2026-07-10", "2026-07-11"
    added = client.post(f"/api/campaign/{first}/daily-rows", json={"name": "التدخل السريع"}).get_json()
    row = next(x for x in added["daily"] if x["name"] == "التدخل السريع")
    assert client.put(f"/api/campaign/{first}/daily/{row['id']}", json={"driver": {"kind": "text", "name": "سائق"}}).status_code == 200
    later_view = client.get(f"/api/campaign/{later}").get_json()
    assert next(x for x in later_view["daily"] if x["id"] == row["id"])["inherited"] is True
    assert client.delete(f"/api/campaign/{first}/daily-rows/{row['id']}").status_code == 200
    assert not any(x["id"] == row["id"] for x in client.get(f"/api/campaign/{later}").get_json()["daily"])
    assert not any(x["id"] == row["id"] for x in client.get(f"/api/campaign/{first}").get_json()["daily"])


def test_campaign_daily_duplicate_names_get_distinct_ids(client):
    day = "2026-07-10"
    first = client.post(f"/api/campaign/{day}/daily-rows", json={"name": "التدخل السريع"}).get_json()
    second = client.post(f"/api/campaign/{day}/daily-rows", json={"name": "التدخل السريع"}).get_json()
    rows = [row for row in second["daily"] if row["name"] == "التدخل السريع"]
    assert len(rows) == 2 and rows[0]["id"] != rows[1]["id"]


def test_campaign_later_daily_add_does_not_appear_earlier(client):
    later = "2026-07-11"
    added = client.post(f"/api/campaign/{later}/daily-rows", json={"name": "قول لاحق"}).get_json()
    assert any(row["name"] == "قول لاحق" for row in added["daily"])
    assert not any(row["name"] == "قول لاحق" for row in _campaign(client, "2026-07-10")["daily"])


def test_campaign_attach_free_text_unknown_and_non_inherited(client):
    first, later = "2026-07-10", "2026-07-11"
    assert client.post(f"/api/campaign/{first}/lorry/LOR-3/service", json={"name": "خدمة حرة"}).status_code == 200
    assert client.post(f"/api/campaign/{first}/lorry/LOR-3/service", json={"assignment_id": "AS-9999"}).status_code == 404
    assert not any("خدمة حرة" in row["name"] for row in _campaign(client, later)["lorries"])


def test_campaign_services_include_emergency_and_custom_sections_without_inheritance(client):
    first, later = "2026-07-10", "2026-07-11"
    emergency = _assignment(client, first, "طوارئ")
    custom = _assignment(client, first, "تفتيشات خاصة", "تفتيشات")
    vehicle = _vehicle(client)
    client.put(f"/api/campaign/{first}/service/{emergency['id']}", json={"vehicle_id": vehicle["id"]})
    names = {row["name"] for row in _campaign(client, first)["services"]}
    assert {"طوارئ", "تفتيشات خاصة"} <= names
    assert _campaign(client, later)["services"] == []


def test_campaign_lorry_and_daily_values_inherit_then_change_locally(client):
    first, later = "2026-07-10", "2026-07-11"
    vehicle = _vehicle(client)
    person = _person(client)
    client.put(f"/api/campaign/{first}/lorry/LOR-1", json={"vehicle_id": vehicle["id"], "driver": {"kind": "personnel", "id": person["id"]}})
    added = client.post(f"/api/campaign/{first}/daily-rows", json={"name": "قول"}).get_json()
    row_id = next(row["id"] for row in added["daily"] if row["name"] == "قول")
    client.put(f"/api/campaign/{first}/daily/{row_id}", json={"vehicle_id": vehicle["id"]})
    client.put(f"/api/campaign/{later}/lorry/LOR-1", json={"driver": {"kind": "text", "name": "جديد"}})
    later_lorry = next(row for row in _campaign(client, later)["lorries"] if row["id"] == "LOR-1")
    earlier_lorry = next(row for row in _campaign(client, first)["lorries"] if row["id"] == "LOR-1")
    assert later_lorry["inherited"] and later_lorry["vehicle"]["id"] == vehicle["id"]
    assert earlier_lorry["driver"]["id"] == person["id"]
    later_daily = next(row for row in _campaign(client, later)["daily"] if row["id"] == row_id)
    assert later_daily["inherited"] and later_daily["vehicle"]["id"] == vehicle["id"]


def test_campaign_warnings_are_row_scoped_and_flattened(client):
    day = "2026-07-10"
    vehicle = _vehicle(client, status="disabled", fault="عطل")
    person = _person(client, name="فرد خارج القوة", code="992", join_date="2027-01-01")
    client.put(f"/api/campaign/{day}/lorry/LOR-1", json={"vehicle_id": vehicle["id"], "driver": {"kind": "personnel", "id": person["id"]}})
    client.put(f"/api/campaign/{day}/lorry/LOR-2", json={"vehicle_id": vehicle["id"], "driver": {"kind": "personnel", "id": person["id"]}})
    view = _campaign(client, day)
    texts = [item["text"] for item in view["warnings"]]
    assert any("أكثر من صف" in text for text in texts)
    assert "المركبة معطلة." in texts
    assert any("ليس على القوة" in text for text in texts)


def test_campaign_export_and_stale_revision(client):
    day = "2026-07-10"
    view = client.get(f"/api/campaign/{day}").get_json()
    response = client.get(f"/api/campaign/{day}/export.docx")
    assert response.status_code == 200
    doc = Document(__import__("io").BytesIO(response.data))
    text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "خطــــــــــوط سيــــــــــر توصيــــــــــل الخدمــــــــــات" in text
    assert "لواري التعيينات" not in text
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "revision": "old", "vehicle_id": "",
    }).get_json()["code"] == "stale_revision"


def test_campaign_export_contains_assigned_operational_text(client):
    day = "2026-07-10"
    vehicle = _vehicle(client)
    person = _person(client)
    assignment = _assignment(client, day, "خدمة تصدير")
    client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "vehicle_id": vehicle["id"], "driver": {"kind": "personnel", "id": person["id"]},
    })
    client.put(f"/api/campaign/{day}/service/{assignment['id']}", json={
        "vehicle_id": vehicle["id"], "driver": {"kind": "text", "name": "سائق خدمة"},
    })
    response = client.get(f"/api/campaign/{day}/export.docx")
    doc = Document(__import__("io").BytesIO(response.data))
    text = " ".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    assert all(value in text for value in ("خدمة تصدير", "سائق خدمة", "4997"))


def test_campaign_invalid_day_and_closed_day_write(client):
    assert client.put("/api/campaign/not-a-day/lorry/LOR-1", json={}).status_code == 400
    day = "2026-07-10"
    client.post(f"/api/day-status/{day}/close", json={})
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1", json={"vehicle_id": ""}).status_code == 409


def test_campaign_references_block_vehicle_and_conscript_delete(client):
    day = "2026-07-10"
    vehicle = _vehicle(client)
    conscript = client.post("/api/conscripts", json={"name": "سائق مرجعي"}).get_json()
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "vehicle_id": vehicle["id"], "driver": {"kind": "conscript", "id": conscript["id"]}}).status_code == 200
    assert client.delete(f"/api/vehicles/{vehicle['id']}").status_code == 409
    assert client.delete(f"/api/conscripts/{conscript['id']}").status_code == 409


def test_campaign_personnel_delete_detaches_driver_to_text(client):
    day = "2026-07-10"
    person = _person(client, name="سائق سيُحذف", code="993")
    client.put(f"/api/campaign/{day}/lorry/LOR-1", json={
        "driver": {"kind": "personnel", "id": person["id"]},
    })
    assert client.post(f"/api/person/{person['id']}/remove", json={
        "leave_date": day, "reason": "انتهاء الخدمة",
    }).status_code == 200
    assert client.delete(f"/api/person/{person['id']}").status_code == 200
    driver = next(row for row in _campaign(client, day)["lorries"] if row["id"] == "LOR-1")["driver"]
    assert driver["kind"] == "text" and driver["name"] == "سائق سيُحذف"


def test_campaign_posts_add_update_inherit_clear_and_remove(client):
    first, later = "2026-07-10", "2026-07-11"
    vehicle = _vehicle(client)
    added = client.post(f"/api/campaign/{first}/posts", json={
        "name": "بوابة الاختبار", "number": "6.", "driver_text": "قوات الأمن صبح , ليل",
    }).get_json()
    post = added["posts"][0]
    assert post["id"].startswith("CP-") and post["number"] == "6."
    assert client.put(f"/api/campaign/{first}/post/{post['id']}", json={
        "vehicle_id": vehicle["id"], "plate_text": "3279ب13", "type_text": "مايك ايفكو",
    }).status_code == 200
    inherited = next(row for row in _campaign(client, later)["posts"] if row["id"] == post["id"])
    assert inherited["inherited"] and inherited["vehicle"]["id"] == vehicle["id"]
    assert client.put(f"/api/campaign/{later}/post/{post['id']}", json={
        "vehicle_id": "", "driver_text": "",
    }).status_code == 200
    cleared = next(row for row in _campaign(client, later)["posts"] if row["id"] == post["id"])
    assert cleared["vehicle"] is None and cleared["driver_text"] == ""
    assert client.delete(f"/api/campaign/{first}/posts/{post['id']}").status_code == 200
    assert not _campaign(client, first)["posts"] and not _campaign(client, later)["posts"]


def test_campaign_post_identity_edit_starts_today_and_keeps_inherited_values(client):
    first, later = "2026-07-10", "2026-07-11"
    vehicle = _vehicle(client)
    added = client.post(f"/api/campaign/{first}/posts", json={
        "name": "ارتكاز البوابة", "number": "1.", "driver_text": "سائق الليل",
    }).get_json()
    post_id = added["posts"][0]["id"]
    client.put(f"/api/campaign/{first}/post/{post_id}", json={"vehicle_id": vehicle["id"]})

    inherited = next(row for row in _campaign(client, later)["posts"] if row["id"] == post_id)
    assert inherited["inherited"] is True
    assert inherited["driver_text"] == "سائق الليل"
    assert inherited["vehicle"]["id"] == vehicle["id"]

    changed = client.patch(f"/api/campaign/{later}/posts/{post_id}", json={
        "name": "ارتكاز البوابة الرئيسية", "number": "2.",
    })
    assert changed.status_code == 200
    old_row = next(row for row in _campaign(client, first)["posts"] if row["id"] == post_id)
    new_row = next(row for row in changed.get_json()["posts"] if row["id"] == post_id)
    assert (old_row["name"], old_row["number"]) == ("ارتكاز البوابة", "1.")
    assert (new_row["name"], new_row["number"]) == ("ارتكاز البوابة الرئيسية", "2.")
    assert new_row["driver_text"] == "سائق الليل" and new_row["vehicle"]["id"] == vehicle["id"]


def test_campaign_posts_warnings_plate_register_and_vehicle_reference(client):
    day, later = "2026-07-10", "2026-07-11"
    vehicle = _vehicle(client, status="disabled", fault="عطل")
    first = client.post(f"/api/campaign/{day}/posts", json={"name": "أ", "driver_text": "سائق"}).get_json()
    second = client.post(f"/api/campaign/{day}/posts", json={"name": "ب"}).get_json()
    ids = [row["id"] for row in first["posts"] + second["posts"]]
    for ident in ids:
        assert client.put(f"/api/campaign/{day}/post/{ident}", json={"vehicle_id": vehicle["id"]}).status_code == 200
    view = _campaign(client, day)
    assert any("أكثر من ارتكاز" in text for text in [w["text"] for w in view["warnings"]])
    assert any("معطلة" in text for text in [w["text"] for w in view["warnings"]])
    assert client.delete(f"/api/vehicles/{vehicle['id']}").status_code == 409
    assert client.put(f"/api/campaign/{day}/plates", json={"rows": [["ب13/9416", "ب13/9417"]]}).status_code == 200
    plate_view = _campaign(client, later)["plates"]
    assert plate_view["inherited"] and plate_view["rows"][0][0] == "ب13/9416"
    assert client.put(f"/api/campaign/{day}/plates", json={"rows": [[]] + [["x"] * 13]}).status_code == 400


def test_campaign_post_write_guards(client):
    day = "2026-07-10"
    client.post(f"/api/day-status/{day}/close", json={})
    assert client.post(f"/api/campaign/{day}/posts", json={"name": "مغلق"}).status_code == 409
    view = _campaign(client, "2026-07-11")
    assert client.put("/api/campaign/2026-07-11/plates", json={"revision": "old", "rows": []}).get_json()["code"] == "stale_revision"


def test_export_prints_personnel_driver_with_short_label(client):
    day = "2026-07-10"
    person = client.post("/api/person", json={
        "type": "personnel", "name": "محمد عطية منصور نصار", "code": "9301", "phone": "01010101010",
        "join_date": "2020-01-01", "role": "معاون شرطة ثالث", "post": "سائق"}).get_json()
    assert client.put(f"/api/campaign/{day}/lorry/LOR-1",
                      json={"driver": {"kind": "personnel", "id": person["id"]}}).status_code == 200
    doc = Document(__import__("io").BytesIO(client.get(f"/api/campaign/{day}/export.docx").data))
    text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "م.ش/ محمد عطية منصور" in text and "نصار" not in text
