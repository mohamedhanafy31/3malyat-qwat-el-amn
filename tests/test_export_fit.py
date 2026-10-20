import io

import pytest
from docx import Document
from openpyxl import Workbook

from backend.export_fit import (PageFitError, TextMeter, configure_xlsx_one_page,
                                 estimate_table_height, fit_table)


def test_meter_single_line_and_wrap_height():
    total, heights = estimate_table_height([[{"text": "نص", "size_pt": 10}]], [100])
    assert 11.4 <= heights[0] <= 15
    wrapped, _ = estimate_table_height([[{"text": "كلمة " * 30, "size_pt": 10}]], [30])
    assert wrapped > total


def test_estimator_handles_colspan_and_vertical_merge():
    total, heights = estimate_table_height([
        [{"text": "عنوان طويل", "colspan": 2, "size_pt": 10}],
        [{"text": "", "rowspan": 2, "size_pt": 24}, {"text": "x", "size_pt": 6}],
        [{"text": "y", "size_pt": 6}],
    ], [30, 70])
    assert total == pytest.approx(sum(heights))
    assert heights[1] + heights[2] >= 24 * 1.1499


def test_fit_table_shrinks_and_rejects_impossible_table():
    result = fit_table({"rows": [[{"text": "كلمة " * 8, "size_pt": 18}] for _ in range(8)],
                        "grid_widths_pt": [80]}, 500)
    assert result.scale < 1
    with pytest.raises(PageFitError):
        fit_table({"rows": [[{"text": "سطر " * 8, "size_pt": 40}] for _ in range(3)],
                   "grid_widths_pt": [10]}, 20)


def test_xlsx_one_page_configuration():
    wb = Workbook()
    ws = wb.active
    ws["B2"] = "x"
    configure_xlsx_one_page(ws)
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 1)
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 9
    assert "$A$1" in str(ws.print_area) and "$B$2" in str(ws.print_area)


@pytest.mark.parametrize("url", [
    "/api/board/2026-04-10/export.docx",
    "/api/duty/2026-04-10/export.docx",
    "/api/afraad/2026-04-10/export.docx",
    "/api/campaign/2026-04-10/export.docx",
])
def test_daily_exports_are_a4_and_reopen(client, url):
    response = client.get(url)
    assert response.status_code == 200
    doc = Document(io.BytesIO(response.data))
    section = doc.sections[0]
    assert (section.page_width.twips, section.page_height.twips) == (11906, 16838)
    assert doc.paragraphs is not None
