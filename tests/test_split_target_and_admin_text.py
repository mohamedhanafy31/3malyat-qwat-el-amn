"""تقسيم تغطية الهدف بالوقت (`periods`) + نص «عمل بالإدارة» الحر لكل ضابط
(`admin_work_text`) — الحفظ، اللوحة، ملف Word، يومية الضباط، والتوافق."""
import io

from docx import Document

from backend.store import load_data

DAY = "2026-04-10"
TARGET = "سوميد"
SPLIT = [{"officer_id": "OFF-001", "until": "16:00"},
         {"officer_id": "OFF-002", "from": "16:00"}]


def _put_target(client, body, name=TARGET, day=DAY):
    return client.put(f"/api/board/{day}/target/{name}", json=body)


def _board_target(client, name=TARGET, day=DAY):
    sections = client.get(f"/api/board/{day}").get_json()["sections"]
    section = next(s for s in sections if s["name"] == "الأهداف")
    return next(r for r in section["rows"] if r["name"] == name)


def _docx_text(client, day=DAY):
    doc = Document(io.BytesIO(client.get(f"/api/board/{day}/export.docx").data))
    out = []

    def walk(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    out.extend(p.text for p in cell.paragraphs)
                    walk(cell.tables)

    walk(doc.tables)
    return "\n".join(out)


def _duty_row(client, person_id, day=DAY):
    return next(r for r in client.get(f"/api/duty/{day}").get_json()["rows"]
                if r["id"] == person_id)


# ---------------------------------------------------------------- تقسيم الهدف

def test_split_round_trip_stores_periods_and_both_officer_ids(client):
    response = _put_target(client, {"periods": SPLIT})
    assert response.status_code == 200, response.get_data(as_text=True)
    stored = next(r for r in load_data([DAY])["day_assignments"][DAY] if r["name"] == TARGET)
    assert stored["periods"] == SPLIT
    assert stored["officer_ids"] == ["OFF-001", "OFF-002"]

    row = _board_target(client)
    assert [o["id"] for o in row["officers"]] == ["OFF-001", "OFF-002"]
    assert [o["period_label"] for o in row["officers"]] == ["حتى 4م", "من 4م"]
    assert row["periods"] == SPLIT


def test_split_time_is_normalised_from_word_style_text(client):
    _put_target(client, {"periods": [{"officer_id": "OFF-001", "until": "4م"},
                                     {"officer_id": "OFF-002", "from": "4:00م"}]})
    assert _board_target(client)["periods"][0]["until"] == "16:00"


def test_split_validation_rejects_bad_shapes(client):
    one = [SPLIT[0]]
    three = [*SPLIT, {"officer_id": "OFF-001", "from": "16:00"}]
    bad_order = [SPLIT[1], SPLIT[0]]
    mismatch = [SPLIT[0], {"officer_id": "OFF-002", "from": "17:00"}]
    same = [SPLIT[0], {"officer_id": "OFF-001", "from": "16:00"}]
    bad_time = [{"officer_id": "OFF-001", "until": "xx"}, SPLIT[1]]
    unknown = [SPLIT[0], {"officer_id": "OFF-999", "from": "16:00"}]
    for periods in (one, three, bad_order, mismatch, same, bad_time, unknown, "x", [1, 2]):
        response = _put_target(client, {"periods": periods})
        assert response.status_code in (400, 404), (periods, response.get_data(as_text=True))
    assert _board_target(client)["vacant"] is True


def test_split_officer_must_be_on_force_that_day(client):
    # OFF-002 انضم 2020 فتمام؛ يوم قبل الانضمام مرفوض
    response = _put_target(client, {"periods": SPLIT}, day="2019-06-01")
    assert response.status_code == 400


def test_split_must_match_explicit_officer_ids(client):
    response = _put_target(client, {"periods": SPLIT, "officer_ids": ["OFF-001"]})
    assert response.status_code == 400


def test_changing_officers_without_periods_drops_stale_split(client):
    _put_target(client, {"periods": SPLIT})
    _put_target(client, {"officer_ids": ["OFF-001"]})
    stored = next(r for r in load_data([DAY])["day_assignments"][DAY] if r["name"] == TARGET)
    assert "periods" not in stored
    row = _board_target(client)
    assert row["periods"] == [] and "period_label" not in row["officers"][0]


def test_resending_same_officers_keeps_the_split(client):
    _put_target(client, {"periods": SPLIT})
    _put_target(client, {"officer_ids": ["OFF-002", "OFF-001"]})
    assert _board_target(client)["periods"] == SPLIT


def test_explicit_empty_periods_clears_the_split_but_keeps_officers(client):
    _put_target(client, {"periods": SPLIT})
    _put_target(client, {"officer_ids": ["OFF-001", "OFF-002"], "periods": []})
    row = _board_target(client)
    assert row["periods"] == [] and len(row["officers"]) == 2


def test_generic_assignment_route_accepts_and_validates_periods(client):
    _put_target(client, {"officer_ids": ["OFF-001"]})
    row_id = _board_target(client)["id"]
    bad = client.patch(f"/api/assignments/{DAY}/{row_id}",
                       json={"periods": [SPLIT[0]]})
    assert bad.status_code == 400
    ok = client.patch(f"/api/assignments/{DAY}/{row_id}", json={"periods": SPLIT})
    assert ok.status_code == 200, ok.get_data(as_text=True)
    assert _board_target(client)["periods"] == SPLIT
    # مش هدف: مرفوض
    made = client.post(f"/api/assignments/{DAY}", json={
        "name": "تدخل سريع", "kind": "خارجية", "shift": "صباحية",
        "section": "الخدمات أساسية", "officer_ids": ["OFF-001"], "periods": SPLIT})
    assert made.status_code == 400


def test_both_split_officers_count_as_guards_in_officers_daily(client):
    _put_target(client, {"periods": SPLIT})
    for person in ("OFF-001", "OFF-002"):
        assert _duty_row(client, person)["group"] == "حراسات"
    summary = client.get(f"/api/duty/{DAY}").get_json()["summary"]
    assert summary["حراسات"] == 2 and summary["balanced"]


def test_board_word_export_prints_until_and_from_with_taqseera_mark(client):
    client.put(f"/api/duty/{DAY}/OFF-001", json={"taqseera": True})
    _put_target(client, {"periods": SPLIT})
    text = _docx_text(client)
    assert "عقيد/ أحمد محمد حتى 4م + ت" in text
    assert "نقيب/ محمود علي من 4م" in text
    assert "من 4م + ت" not in text


def test_unsplit_target_is_unchanged_everywhere(client):
    _put_target(client, {"officer_ids": ["OFF-001", "OFF-002"]})
    row = _board_target(client)
    assert row["periods"] == []
    assert all("period_label" not in o for o in row["officers"])
    stored = next(r for r in load_data([DAY])["day_assignments"][DAY] if r["name"] == TARGET)
    assert "periods" not in stored
    text = _docx_text(client)
    assert "حتى 4م" not in text and "من 4م" not in text


def test_confirm_flags_a_split_change_as_pending(client):
    _put_target(client, {"officer_ids": ["OFF-001", "OFF-002"]})
    client.post(f"/api/board/{DAY}/confirm")
    assert client.get(f"/api/board/{DAY}").get_json()["confirm"]["pending"] is False
    _put_target(client, {"periods": SPLIT})
    assert client.get(f"/api/board/{DAY}").get_json()["confirm"]["pending"] is True


# --------------------------------------------------------- نص عمل بالإدارة

def _admin_rows(client, day=DAY):
    sections = client.get(f"/api/board/{day}").get_json()["sections"]
    return next(s for s in sections if s["name"] == "عمل بالإدارة")["rows"]


TEXT = "متابعة أعمال الشئون حتى 4م + مشرف الأهداف من 4م"


def test_admin_work_text_round_trip_and_board_row(client):
    response = client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": f"  {TEXT}  "})
    assert response.status_code == 200, response.get_data(as_text=True)
    assert _duty_row(client, "OFF-002")["admin_work_text"] == TEXT
    stored = load_data([DAY])["day_officers"][DAY]["OFF-002"]
    assert stored["admin_work_text"] == TEXT
    row = next(r for r in _admin_rows(client) if r["id"] == "OFF-002")
    assert row["text"] == TEXT and row["free_text"] is True


def test_admin_work_text_is_not_duplicated_with_service_or_flag(client):
    client.post(f"/api/assignments/{DAY}", json={
        "name": "تدخل سريع", "kind": "خارجية", "shift": "صباحية",
        "section": "الخدمات أساسية", "officer_ids": ["OFF-002"]})
    client.put(f"/api/duty/{DAY}/OFF-002", json={
        "extra": "إضافي", "admin_work": True, "admin_work_text": TEXT})
    rows = [r for r in _admin_rows(client) if r["id"] == "OFF-002"]
    assert len(rows) == 1 and rows[0]["text"] == TEXT
    # يومية الضباط لسه بنصها التشغيلي المعتاد
    assert "تدخل سريع" in _duty_row(client, "OFF-002")["text"]


def test_admin_work_text_alone_needs_no_flag_and_no_operational_text(client):
    assert client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work": True}).status_code == 400
    ok = client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work": True, "admin_work_text": TEXT})
    assert ok.status_code == 200


def test_clearing_admin_work_text_removes_key_and_row(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": TEXT})
    client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": ""})
    assert not any(r["id"] == "OFF-002" for r in _admin_rows(client))
    assert "OFF-002" not in (load_data([DAY]).get("day_officers") or {}).get(DAY, {})


def test_other_state_fields_do_not_touch_admin_work_text(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": TEXT})
    client.put(f"/api/duty/{DAY}/OFF-002", json={"taqseera": True})
    assert _duty_row(client, "OFF-002")["admin_work_text"] == TEXT


def test_admin_work_text_too_long_is_rejected(client):
    response = client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": "س" * 501})
    assert response.status_code == 400


def test_admin_work_text_does_not_apply_to_outside_officers(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": TEXT, "status": "غياب"})
    assert not any(r["id"] == "OFF-002" for r in _admin_rows(client))


def test_admin_work_text_is_printed_exactly_in_word(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": TEXT})
    text = _docx_text(client)
    assert TEXT in text and text.count(TEXT) == 1


def test_automatic_director_row_uses_the_free_text_when_present(client):
    client.patch("/api/command", json={"مدير الإدارة": "OFF-001"})
    row = next(r for r in _admin_rows(client) if r["id"] == "OFF-001")
    assert row["text"] == "مدير الإدارة" and row["automatic"] is True
    client.put(f"/api/duty/{DAY}/OFF-001", json={"admin_work_text": "مدير الادارة + المرور"})
    row = next(r for r in _admin_rows(client) if r["id"] == "OFF-001")
    assert row["text"] == "مدير الادارة + المرور" and row["automatic"] is True


def test_change_log_text_mentions_admin_work_text(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work_text": TEXT})
    blob = str(client.get("/api/changes").get_json()["entries"])
    assert "نص عمل بالإدارة" in blob and TEXT in blob
