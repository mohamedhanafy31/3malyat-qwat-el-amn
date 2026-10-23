from datetime import date
from io import BytesIO

from docx import Document

from backend.leaves_export import monthly_report, weekly_report
from backend.store import load_data


def _officer(client, name, role, rest_system, rest_day=""):
    serial = sum(ord(char) for char in name)
    response = client.post("/api/person", json={
        "type": "officer", "name": name, "role": role, "code": str(serial),
        "phone": f"010{serial:08d}"[-11:], "join_date": "2020-01-01", "rest_system": rest_system,
        "rest_day": rest_day,
    })
    assert response.status_code == 201, response.get_json()
    return response.get_json()["id"]


def test_monthly_report_groups_rows_and_uses_latest_plan(client):
    half = _officer(client, "محمد ميهوب", "رائد", "نصف شهرية")
    monthly = _officer(client, "أحمد السمان", "رائد", "شهرية")
    for officer_id, kind, start, end in (
        (half, "نصف شهرية", "2026-08-28", "2026-08-30"),
        (monthly, "شهرية", "2026-09-08", "2026-09-14"),
    ):
        assert client.post("/api/leaves", json={
            "person_id": officer_id, "type": kind, "start": start, "end": end,
        }).status_code == 201
    report = monthly_report(load_data(()), date(2026, 9, 9))
    assert report["groups"]["نصف شهرية"][0]["taqseera"] == "2026-08-27"
    assert report["groups"]["شهرية"][0]["return_date"] == "2026-09-15"


def test_monthly_word_export_matches_two_section_layout(client):
    monthly = _officer(client, "أحمد السمان", "رائد", "شهرية")
    client.post("/api/leaves", json={"person_id": monthly, "type": "شهرية",
                                      "start": "2026-09-08", "end": "2026-09-14"})
    response = client.get("/api/leaves/monthly/export.docx?on=2026-09-09")
    assert response.status_code == 200
    document = Document(BytesIO(response.data))
    assert len(document.tables) == 4
    text = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "كشف الراحات الشهرية" in text
    assert "كشف الراحات النصف شهرية" in text
    assert "2026/9/7م" in text and "2026/9/15م" in text


def test_weekly_word_export_has_seven_day_pairs_and_officers(client):
    sunday = _officer(client, "محمد مصطفى", "مقدم", "أسبوعية", "الأحد")
    saturday = _officer(client, "علي فاروق", "نقيب", "أسبوعية", "السبت")
    report = weekly_report(load_data(()), date(2026, 8, 15))
    assert {row["id"] for row in report["groups"]["الأحد"]} == {sunday}
    assert {row["id"] for row in report["groups"]["السبت"]} >= {saturday, "OFF-001"}

    response = client.get("/api/leaves/weekly/export.docx?on=2026-08-15")
    assert response.status_code == 200
    document = Document(BytesIO(response.data))
    assert len(document.tables) == 1 and len(document.tables[0].rows) == 20
    text = " ".join(p.text for p in document.paragraphs)
    text += " " + " ".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "2026/8/15" in text and "مقدم/ محمد مصطفى" in text
    assert all(day in text for day in ("السبت", "الأحد", "الاثنين", "الثلاثاء",
                                       "الأربعاء", "الخميس", "الجمعة"))


def test_leave_exports_reject_invalid_report_date(client):
    assert client.get("/api/leaves/monthly/export.docx?on=bad").status_code == 400
    assert client.get("/api/leaves/weekly/export.docx?on=bad").status_code == 400
