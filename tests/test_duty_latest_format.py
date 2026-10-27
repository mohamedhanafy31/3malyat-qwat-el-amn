from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.duty_export import LAST_FIT, _daily_text, build_docx
from tools.fidelity.duty import _structure, load_docx


def _summary():
    return {
        "أصل القوة": 1,
        "خارجية": {"صباحية": 0, "ليلية": 0, "بحث": 0},
        "داخلية": {"صباحية": 0, "ليلية": 0},
        "طبية": {"موجود": 0, "راحة": 0},
        "خوارج": {key: 0 for key in ("تقصيرة", "راحة", "طارئة", "غياب", "مرضي", "فرقة", "انتداب")},
        "حراسات": 0,
        "صافي": 1,
        "net_names": ["ignored by exporter"],
    }


def _duty(rest_system="أسبوعية", rest_day=None):
    return {"date": "2026-09-29", "summary": _summary(), "rows": [{
        "section": "القوة", "role": "ملازم أول", "name": "الاسم الكامل للضابط",
        "post": "عمل بالإدارة", "rest_system": rest_system, "rest_day": rest_day,
        "group": "صافي", "services": [], "status": "", "note": "",
    }]}


ARCHIVE_DAY = Path(__file__).parents[3] / "2026/10/21/21-10-2026.docx"
ROSTER_GRID = [573, 572, 2410, 2539, 3842, 971, 7]


def _build(duty=None):
    return Document(BytesIO(build_docx(duty or _duty()).getvalue()))


def _runs(table):
    return list(table._tbl.iter(qn("w:r")))


def _sz(run):
    node = run.find(qn("w:rPr") + "/" + qn("w:sz"))
    return int(node.get(qn("w:val"))) if node is not None else None


def test_latest_geometry_and_one_page_fit():
    output = build_docx(_duty()).getvalue()
    document = Document(BytesIO(output))
    assert len(document.tables[0].rows) == 37
    assert len(document.tables[1].rows) == 4
    # شبكة الملف الحقيقي (أكتوبر 2026) لا القالب القديم
    assert [int(g.get(qn("w:w"))) for g in document.tables[0]._tbl.tblGrid] == ROSTER_GRID
    assert len(document.tables[1].columns) == 17          # 17 عمودًا من يوم 16
    assert document.tables[1].columns[16].width.twips == 1416
    assert LAST_FIT["ratio"] <= 1.0
    assert len(Document(BytesIO(output)).tables) == 2


def test_latest_mapping_uses_full_roster_name_and_dash_rest():
    document = _build()
    row = document.tables[0].rows[1]
    assert row.cells[1].text == "م.اول"
    assert row.cells[2].text == "الاسم الكامل للضابط"
    assert row.cells[5].text == "------------"
    assert document.tables[1].rows[0].cells[16].text == "الصافي"


def test_summary_has_patrol_and_mission_columns_in_real_order():
    document = _build()
    headers = [c.text for c in document.tables[1].rows[1].cells]
    assert headers[9:15] == ["دورية", "غياب", "مرضي", "فرقة", "مأمورية", "انتداب"]
    summary = _summary()
    summary["خوارج"]["مأمورية"] = 3
    summary["خوارج"]["طارئة"] = 2
    summary["حراسات"] = 8
    values = [c.text for c in _build({**_duty(), "summary": summary}).tables[1].rows[2].cells]
    assert values[9] == "2" and values[13] == "3" and values[15] == "8"


def test_row_heights_are_at_least_never_exact_and_row_15_matches_real():
    document = _build()
    heights = [int(h.get(qn("w:val"))) for h in document.tables[0]._tbl.iter(qn("w:trHeight"))]
    assert heights[15] == 191
    for node in document.element.body.iter(qn("w:trHeight")):
        assert node.get(qn("w:hRule")) is None


def test_tcw_gridspan_and_grid_are_exactly_the_real_files():
    """tcW/gridSpan/tblGrid/trHeight لا تتغير عن الملف الحقيقي (القالب)."""
    from backend.duty_export import TEMPLATE
    out, template = _build(), Document(TEMPLATE)
    for table_out, table_tpl in zip(out.tables, template.tables):
        for tag in ("w:tcW", "w:gridSpan", "w:gridCol", "w:trHeight", "w:vMerge"):
            assert ([dict(n.attrib) for n in table_out._tbl.iter(qn(tag))]
                    == [dict(n.attrib) for n in table_tpl._tbl.iter(qn(tag))])
    # كل صف بيانات في الجدول الأول مجموع tcW فيه 5000 (نسبة مئوية كاملة)
    for tr in out.tables[0]._tbl.findall(qn("w:tr"))[1:25]:
        assert sum(int(w.get(qn("w:w"))) for w in tr.iter(qn("w:tcW"))) == 5000


def test_runs_keep_template_theme_fonts_and_never_force_arial():
    document = _build()
    for run in document.element.body.iter(qn("w:r")):
        fonts = run.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        if fonts is None:
            continue
        assert fonts.get(qn("w:asciiTheme")) == "majorBidi"
        assert fonts.get(qn("w:ascii")) is None and fonts.get(qn("w:hAnsi")) is None
        assert "Arial" not in fonts.attrib.values()


def test_font_sizes_are_9pt_roster_8pt_summary_without_global_scaling():
    document = _build()
    assert {_sz(r) for r in _runs(document.tables[0])} == {18}
    assert {_sz(r) for r in _runs(document.tables[1])} == {16}
    # لا تكبير/تصغير على الفقرة العلوية (الشعارات) ولا الفقرات الفارغة
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            assert _sz(run._r) in (None, 20, 28)


def test_only_a_long_cell_shrinks_to_8_or_7pt():
    duty = _duty()
    duty["rows"][0]["post"] = "رئيس قسم العلاقات العامة + رئيس قسم التعيينات + مشرف المخبز + مشرف المطبخ " * 3
    document = _build(duty)
    sizes = [_sz(r) for r in _runs(document.tables[0])]
    assert sizes.count(18) > 0.9 * len(sizes)
    long_cell = document.tables[0].rows[1].cells[3]._tc
    assert {_sz(r) for r in long_cell.iter(qn("w:r"))} <= {16, 14}
    assert LAST_FIT["shrunk"]
    # باقي صفوف الصف نفسه بقي 9pt
    assert {_sz(r) for r in document.tables[0].rows[1].cells[2]._tc.iter(qn("w:r"))} == {18}


def test_cell_paragraph_spacing_is_276_without_forced_before_after():
    document = _build()
    for table in document.tables:
        for ppr in table._tbl.iter(qn("w:pPr")):
            spacing = ppr.find(qn("w:spacing"))
            if spacing is None:
                continue
            assert spacing.get(qn("w:before")) is None and spacing.get(qn("w:after")) is None
            assert spacing.get(qn("w:lineRule")) in (None, "auto")
            assert spacing.get(qn("w:line")) in (None, "276")
    cell_ppr = document.tables[0].rows[1].cells[2]._tc.find(qn("w:p")).find(qn("w:pPr"))
    assert cell_ppr.find(qn("w:spacing")).get(qn("w:line")) == "276"
    assert cell_ppr.find(qn("w:jc")).get(qn("w:val")) == "center"


def test_title_is_14pt_underlined_sultan_with_real_spacing_and_hamza_weekday():
    duty = _duty()
    duty["date"] = "2026-10-18"                           # أحد: بالهمزة بقرار المالك
    document = _build(duty)
    title = document.paragraphs[1]
    assert title.text == ("يومية تشغيل الضبـــاط بإدارة قـوات أمن السويس عـن يوم "
                          "الأحد الموافق 18 /10/2026م")
    for run in title.runs:
        rpr = run._r.rPr
        assert _sz(run._r) == 28 and rpr.find(qn("w:u")).get(qn("w:val")) == "single"
        assert rpr.find(qn("w:rFonts")).get(qn("w:cs")) == "Sultan bold"
        assert rpr.find(qn("w:b")) is not None


def test_body_order_signature_between_tables_and_two_trailing_paragraphs():
    body = list(Document(BytesIO(build_docx(_duty()).getvalue())).element.body)
    order = [el.tag.split("}")[1] for el in body]
    assert order == ["p", "p", "tbl", "p", "tbl", "p", "p", "sectPr"]
    signature = "".join(body[3].itertext())
    assert "رئيس" in signature or body[3].find(".//" + qn("w:pict")) is not None


def test_logo_anchors_match_real_file():
    document = _build()
    xml = document.paragraphs[0]._p.xml
    for needle in ("508483", "432671", "57785", "257429", "872592", "702520", "6341821", "7766"):
        assert needle in xml


def test_summary_names_are_one_numbered_paragraph_per_officer():
    duty = _duty()
    duty["rows"] = duty["rows"] + [{**duty["rows"][0], "name": "محمد فؤاد حسن علي", "role": "عقيد"}]
    cell = _build(duty).tables[1].rows[1]._tr.tc_lst[16]
    paragraphs = cell.findall(qn("w:p"))
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        assert paragraph.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None
        assert {_sz(r) for r in paragraph.iter(qn("w:r"))} == {16}
    assert "".join(t.text for t in paragraphs[1].iter(qn("w:t"))) == "عقيد / محمد فؤاد"


def test_vmerge_cells_are_not_scaled_twice():
    document = _build()
    for tc in document.tables[1]._tbl.iter(qn("w:tc")):
        for run in tc.iter(qn("w:r")):
            assert _sz(run) == 16


def test_course_note_is_not_repeated_when_spelling_differs():
    from backend.duty import course_with_note
    course = {"name": "إعداد مدربي القوات الخاصة — تدريب مشترك", "text": "الاسم الكامل المكرر"}
    note = "فرقة اعداد مدربي القوات الخاصة تدريب مشترك خلال الفترة"
    assert course_with_note(course, note) == note


def test_open_status_prints_name_and_start_day_only():
    row = {"status": "مرضي", "status_source": "period", "persistent_status": {
        "status": "مرضي", "start": "2026-10-23", "end": "",
    }}
    assert _daily_text(row) == "مرضي من يوم 23/10/2026"


def test_bounded_status_keeps_start_and_end_days():
    row = {"status": "طارئة", "status_source": "period", "persistent_status": {
        "status": "طارئة", "start": "2026-10-23", "end": "2026-10-25",
    }}
    assert _daily_text(row) == "طارئة من 23/10/2026 حتى 25/10/2026"


def test_archive_structure_signature_matches_latest_template_when_available():
    if not ARCHIVE_DAY.exists():
        pytest.skip("latest archive is not present")
    # Compare the template's fixed XML geometry directly; content is filled in
    # separately and must not change grids, merges, or page setup.
    from backend.duty_export import TEMPLATE
    assert _structure(load_docx(ARCHIVE_DAY)) == _structure(load_docx(TEMPLATE))
