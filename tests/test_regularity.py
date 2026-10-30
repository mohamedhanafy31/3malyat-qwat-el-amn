import shutil
from pathlib import Path
from io import BytesIO

import pytest
from docx import Document

from backend.regularity import DEFAULT_NOTES, PLACEHOLDER, build
from backend.regularity_export import PageFitError, build_docx
from backend import store
from backend.utils import weekday_name
from tools.compare_regularity import compare_day, _norm

DAY = "2026-10-20"


def person(ident, name, role="مقدم", phone="01000000000", post=""):
    return {"id": ident, "name": name, "role": role, "phone": phone,
            "post": post, "join_date": "2020-01-01", "history": []}


def data_for(assignments=None):
    return {
        "officers": [person("OFF-1", "مدير سمير", "عقيد"), person("OFF-2", "وكيل فؤاد", "عقيد"),
                     person("OFF-3", "رئيس عمليات", "رائد", post="رئيس قسم العمليات"),
                     person("OFF-4", "ضابط خدمة", "مقدم")],
        "personnel": [person("IND-1", "فرد أحمد", "معاون")], "conscripts": [],
        "command": {"مدير الإدارة": "OFF-1", "وكيل الإدارة": "OFF-2"},
        "command_history": [], "day_assignments": {DAY: assignments or []},
        "day_officers": {}, "leaves": [], "day_afraad": {}, "day_campaign": {},
        "reference_lists": {"afraad_basic": [{"from": "", "items": [
            {"id": "AFB-19", "name": "الإسعاف", "location_text": "الحماية المدنية", "conscript_count": 3}
        ]}], "campaign_daily": [{"from": "", "items": []}]},
        "vehicles": [], "inspection_schedule": {}, "inspection_seeded_days": {},
        "officer_medical_states": [], "courses": [], "course_terms": [],
    }


def assignment(name, section="الخدمات الطارئة", **extra):
    return {"id": extra.pop("id", "AS-1"), "name": name, "section": section,
            "time": extra.pop("time", "8ص"), "party": extra.pop("party", "المديرية"),
            "conscript_count": extra.pop("conscript_count", 2),
            "personnel_ids": extra.pop("personnel_ids", ["IND-1"]),
            "officer_ids": extra.pop("officer_ids", []), "conscripts": [], **extra}


def test_selection_excludes_prison_includes_inspection_and_special_and_orders_time():
    rows = [assignment("مباراة", section="انتشار", time="12ظ"),
            assignment("سجن", section="خدمات سجن قوات الامن", time="5ص"),
            assignment("تفتيش", section="تفتيشات", time="6ص")]
    payload = build(data_for(rows), "2026-10-20")
    names = [r["service"] for r in payload["service_rows"]]
    assert "سجن" not in names and names[:2] == ["تفتيش", "مباراة"]


def test_placeholders_and_count_source_and_compact_midnight():
    row = assignment("خدمة", party="", time="00:00", conscript_count=7,
                     personnel_ids=[], officer_ids=[])
    service = build(data_for([row]), "2026-10-20")["service_rows"][0]
    assert service["place"] == PLACEHOLDER and service["time"] == "12ل"
    assert service["count"] == "7" and service["people"] == PLACEHOLDER
    assert any("مكان الانتظام" in warning for warning in service["warnings"])


def test_supervisory_switch_and_deputy_extra_only_when_supplied():
    data = data_for([])
    data["leaves"] = [{"id": "LV-1", "person_id": "OFF-1", "type": "راحة",
                       "start": "2026-10-20", "end": "2026-10-20", "return_date": "2026-10-21"}]
    payload = build(data, "2026-10-20", {"officer_extra": {"OFF-2": "مرور تعقبي"}})
    rows = payload["officer_rows"]
    assert rows[0]["id"] == "OFF-2" and "مرور تعقبي" in rows[0]["service"]
    assert not any(row["id"] == "OFF-1" for row in rows)


def test_docx_structure_and_fit_error():
    payload = build(data_for([assignment("خدمة")]), "2026-10-20")
    payload["day"] = "2026-10-20"
    document = Document(BytesIO(build_docx(payload).getvalue()))
    assert [(len(t.rows), len(t.columns)) for t in document.tables] == [(3, 8), (5, 4)]
    assert document.sections[0].page_width < document.sections[0].page_height
    assert "الثلاثاء" in " ".join(p.text for p in document.sections[0].header.paragraphs)
    huge = dict(payload, service_rows=payload["service_rows"] * 40)
    with pytest.raises(PageFitError):
        build_docx(huge)


def vehicle(ident, series, number):
    return {"id": ident, "plate_series": series, "plate_number": number,
            "type": "ميك", "model": "ايفكو", "status": "operational",
            "fault": "", "notes": "", "archived": False}


def two_people_data():
    data = data_for([])
    data["personnel"].append(person("IND-2", "فرد محمود", "معاون", "01111111111"))
    data["conscripts"] = [{"id": "CON-1", "name": "مجند واحد", "phone": ""},
                           {"id": "CON-2", "name": "مجند اثنان", "phone": ""}]
    data["vehicles"] = [vehicle("VEH-1", "ب11", "1111"), vehicle("VEH-2", "ب12", "2222")]
    return data


def test_multi_vehicle_split_and_single_vehicle_people_fallback():
    data = two_people_data()
    row = assignment("خدمة مركبتين", id="AS-M", personnel_ids=["IND-1", "IND-2"])
    data["day_assignments"][DAY] = [row]
    data["day_campaign"][DAY] = {"services": {"AS-M": {"vehicles": [
        {"vehicle_id": "VEH-1", "personnel_ids": ["IND-1"]},
        {"vehicle_id": "VEH-2", "personnel_ids": ["IND-2"]},
    ]}}}
    rows = [r for r in build(data, DAY)["service_rows"] if r["service"] == "خدمة مركبتين"]
    assert len(rows) == 2
    assert {r["vehicle"].split(" — ")[0] for r in rows} == {"ب11/1111", "ب12/2222"}
    assert any("فرد أحمد" in r["people"] and "01000000000" in r["phone"] for r in rows)
    assert any("فرد محمود" in r["people"] and "01111111111" in r["phone"] for r in rows)

    data["day_campaign"][DAY] = {"services": {"AS-M": {"vehicle_id": "VEH-1"}}}
    rows = [r for r in build(data, DAY)["service_rows"] if r["service"] == "خدمة مركبتين"]
    assert len(rows) == 1 and "فرد أحمد" in rows[0]["people"] and "فرد محمود" in rows[0]["people"]


def test_lorry_attached_vehicle_and_missing_vehicle_warning():
    data = two_people_data()
    row = assignment("خدمة لوري", id="AS-L", personnel_ids=["IND-1"])
    data["day_assignments"][DAY] = [row]
    data["day_campaign"][DAY] = {"lorries": {"LOR-1": {
        "vehicle_id": "VEH-1", "attached": [{"assignment_id": "AS-L"}],
    }}}
    generated = [r for r in build(data, DAY)["service_rows"] if r["service"] == "خدمة لوري"][0]
    assert generated["vehicle"].startswith("ب11/1111")
    assert "فرد أحمد" in generated["people"]

    data["day_campaign"][DAY] = {}
    generated = [r for r in build(data, DAY)["service_rows"] if r["service"] == "خدمة لوري"][0]
    assert generated["vehicle"] == PLACEHOLDER
    assert any("المركبة" in warning for warning in generated["warnings"])


def test_ambulance_uses_basic_and_linked_campaign_daily_row():
    data = two_people_data()
    data["day_afraad"] = {DAY: {"AFB-19": {
        "morning_person_id": "IND-1", "morning_name": "فرد أحمد",
        "morning_phone": "01000000000", "schedule": "11 ص حتى الانتهاء",
    }}}
    data["reference_lists"]["campaign_daily"] = [{"from": "", "items": [
        {"id": "CD-005", "name": "ارتكاز الاسعاف", "basic_entry_id": "AFB-19"}
    ]}]
    data["day_campaign"][DAY] = {"daily": {"CD-005": {"vehicle_id": "VEH-1"}}}
    row = next(r for r in build(data, DAY)["service_rows"] if "اسعاف" in r["service"])
    assert "فرد أحمد" in row["people"]
    assert row["phone"] == "01000000000"
    assert row["time"] == "11ص" and row["place"] == "الحماية المدنية"
    assert row["count"] == "3" and row["vehicle"] == "ب11/1111"


@pytest.mark.parametrize("alias", ["خدمات السجن", "سجن قوات الامن", "فرد تأمين فترة ليلية"])
def test_all_prison_aliases_are_excluded(alias):
    rows = build(data_for([assignment("لا يظهر", section=alias)]), DAY)["service_rows"]
    assert not any(row["service"] == "لا يظهر" for row in rows)


def test_custom_normal_and_special_sections_are_included_and_ties_are_stable():
    rows = [assignment("الثاني", section="فرد + مج", time="8ص"),
            assignment("الأول", section="مناسبة خاصة", time="8ص")]
    generated = [r["service"] for r in build(data_for(rows), DAY)["service_rows"]]
    assert generated[:2] == ["الثاني", "الأول"]


@pytest.mark.parametrize(("raw", "expected"), [("6ص", "6ص"), ("12ظ", "12ظ"),
                                                  ("5م", "5م"), ("12ليل", "12ل")])
def test_regular_time_format(raw, expected):
    row = assignment("وقت", time=raw)
    generated = next(r for r in build(data_for([row]), DAY)["service_rows"] if r["service"] == "وقت")
    assert generated["time"] == expected


def test_count_is_not_sum_of_conscripts():
    row = assignment("عدد منظم", conscript_count=1, conscripts=["CON-1", "CON-2"], personnel_ids=[])
    data = two_people_data()
    data["day_assignments"][DAY] = [row]
    generated = next(r for r in build(data, DAY)["service_rows"] if r["service"] == "عدد منظم")
    assert generated["count"] == "1"


def test_supervisory_switch_big_column_override_notes_and_signatories():
    data = data_for([assignment("قول مكبر", section="الخدمات أساسية", id="AS-B", shift="صباحية",
                                officer_ids=["OFF-4"], personnel_ids=[])])
    data["day_officers"][DAY] = {"OFF-1": {"status": "غياب"}}
    inputs = {"big_column_supervisors": {"am": {"rank": "مقدم", "name": "أحمد هلال", "phone": "0102"}},
              "notes": ["ملاحظة واحدة"]}
    payload = build(data, DAY, inputs)
    assert payload["officer_rows"][0]["id"] == "OFF-2"
    assert any(row["name"] == "أحمد هلال (0102)" and "مشرف القول المكبر" in row["service"]
               for row in payload["officer_rows"])
    assert not any(row["id"] == "OFF-4" and "مشرف القول المكبر" in row["service"]
                   for row in payload["officer_rows"])
    assert payload["notes"] == ["ملاحظة واحدة"]
    assert payload["signatories"]["operations"]["name"] == "رئيس عمليات"
    assert build(data_for([]), DAY)["notes"] == DEFAULT_NOTES
    assert build(data_for([]), DAY)["signatories"]["operations"]["name"] == "رئيس عمليات"


def test_dashboard_export_contract(client):
    html = client.get("/").get_data(as_text=True)
    assert 'data-export="docx"' in html
    dashboard_js = Path("static/js/dashboard.js").read_text(encoding="utf-8")
    assert "window.exportDocxUrl" in dashboard_js
    assert "/api/ops/${day}/regularity.docx" in dashboard_js


def test_dashboard_day_navigation_and_regularitiy_ui_contract(client):
    html = client.get("/?day=2026-10-20").get_data(as_text=True)
    assert 'id="dashboardDate"' in html and 'id="dayPrev"' in html
    assert 'id="dashboardDayWarning"' in html
    assert "regularityWarnings" not in html and "regularityOperationsChief" not in html
    dashboard_js = Path("static/js/dashboard.js").read_text(encoding="utf-8")
    assert "new URLSearchParams(location.search).get(\"day\")" in dashboard_js
    assert "/api/ops/${day}" in dashboard_js


def test_regularity_route_download_and_invalid_day(client):
    response = client.get(f"/api/ops/{DAY}/regularity.docx")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "20..docx" in response.headers["Content-Disposition"]
    assert Document(BytesIO(response.data)).tables
    assert client.get("/api/ops/not-a-day/regularity.docx").status_code == 400


def test_regularity_route_previews_inspections_without_store_write(client):
    from backend.inspection_schedule import schedule as inspection_schedule

    def seed(data):
        inspection_schedule(data)[weekday_name(DAY)] = [{
            "id": "INS-1", "name": "تفتيش اختباري", "weapon": "خرطوش", "count": 2,
        }]

    store.with_data(seed, [DAY])
    before = store.assemble()
    response = client.get(f"/api/ops/{DAY}/regularity.docx")
    assert response.status_code == 200
    document = Document(BytesIO(response.data))
    services = [cell.text for row in document.tables[0].rows for cell in row.cells]
    assert any("تفتيش اختباري" in text for text in services)
    assert before == store.assemble()


EXPECTED_DIFFS = {
    "2026-10-20": [
        {"kind": "service_set", "reason": "missing campaign vehicles / تصوير vs المديرية / طلبة امداد 3 vs 2"},
        {"kind": "row", "service": "ماموريهامداد", "field": "people",
         "reason": "two مأمورية امداد rows vs one assignment"},
        {"kind": "row", "service": "نيابهحدثوحريم", "field": "people",
         "reason": "historical personnel gap"},
        {"kind": "notes", "reason": "نيابة سياسي note"},
        {"kind": "inputs", "reason": "morning big-column supervisor and deputy extra text"},
    ],
    "2026-09-25": [
        {"kind": "service_set", "reason": "missing campaign vehicles / imported day without seeded inspections"},
        {"kind": "row", "service": "ارتكازالمسجدالكبير", "field": "people",
         "reason": "قول مكبر ليلية has two stored officers"},
        {"kind": "notes", "reason": "archive note count 5"},
        {"kind": "inputs", "reason": "morning big-column supervisor and deputy extra text"},
    ],
}


@pytest.mark.parametrize("day", ["2026-10-20", "2026-09-25"])
def test_archive_comparison_on_core_and_days_copy_only(tmp_path, monkeypatch, day):
    archive_root = Path("../../2026").resolve()
    real_data = Path("data")
    if not archive_root.exists() or not (real_data / "core.json").exists():
        pytest.skip("archive or real data folder is absent")
    temp_data = tmp_path / "data"
    temp_data.mkdir()
    shutil.copy2(real_data / "core.json", temp_data / "core.json")
    shutil.copytree(real_data / "days", temp_data / "days")
    monkeypatch.setattr(store, "DATA_DIR", temp_data)
    store._cache.clear()
    result = compare_day(day, str(archive_root), store.load_data([day]))
    assert _norm(result["generated"][0]) == _norm(result["archived"][0])
    generated_headers = [_norm(value) for value in result["generated"][1][0]]
    archive_headers = [_norm(value) for value in result["archived"][1][0]]
    assert generated_headers[0] == archive_headers[0]
    assert generated_headers[1] == archive_headers[1] or {generated_headers[1], archive_headers[1]} == {"الخدمه", "الخدم"}
    assert generated_headers[2:] == archive_headers[2:]
    generated_names = set(result["generated_service_names"])
    archive_names = set(result["archive_service_names"])
    assert set(result["shared_service_names"]) == generated_names & archive_names
    allowed = EXPECTED_DIFFS[day]
    assert allowed and all(isinstance(item, dict) for item in allowed)
    if generated_names != archive_names:
        assert any(item["kind"] == "service_set" for item in allowed)
    generated_notes = [_norm(row[1]) for row in result["generated"][2][-3:] if len(row) > 1]
    archive_notes = [_norm(value) for value in result["archived"][3][-3:]]
    if generated_notes != archive_notes:
        assert any(item["kind"] == "notes" for item in allowed)
    for check in result["row_checks"]:
        for field, equal in check["fields"].items():
            if not equal:
                assert any(item["kind"] == "row" and item["service"] == check["service"]
                           and item["field"] == field for item in allowed), \
                    f"unexpected {day} difference: {check['service']} {field}"
