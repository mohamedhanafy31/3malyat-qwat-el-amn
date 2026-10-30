"""مدخلات يومية انتظام الخدمات: توريث، تحقق، ودخولها في التصدير."""
import io

from docx import Document


def _inputs(client, day):
    response = client.get(f"/api/ops/{day}/regularity")
    assert response.status_code == 200
    return response.get_json()


def _text(doc):
    return "\n".join(
        paragraph.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )


def test_regularity_inputs_inherit_and_own_entry_wins_and_reset_falls_back(client):
    assert client.patch("/api/command", json={"مدير الإدارة": "OFF-001"}).status_code == 200
    assert client.patch("/api/command", json={"وكيل الإدارة": "OFF-002"}).status_code == 200
    first = _inputs(client, "2026-04-10")
    value = {"officer_extra": {"OFF-001": "مرور مسائي"},
             "big_column_supervisors": {"am": {"rank": "مقدم", "name": "مشرف خارجي", "phone": "0109"}},
             "notes": ["ملاحظة موروثة"]}
    saved = client.put("/api/ops/2026-04-10/regularity",
                       json={"revision": first["revision"], "inputs": value})
    assert saved.status_code == 200

    inherited = _inputs(client, "2026-04-11")
    assert inherited["inherited_from"] == "2026-04-10"
    assert inherited["inputs"] == value

    own = {**value, "notes": ["تخصيص اليوم"]}
    saved = client.put("/api/ops/2026-04-11/regularity",
                       json={"revision": inherited["revision"], "inputs": own})
    assert saved.status_code == 200
    assert _inputs(client, "2026-04-11")["inherited_from"] == "2026-04-11"

    current = _inputs(client, "2026-04-11")
    reset = client.put("/api/ops/2026-04-11/regularity",
                       json={"revision": current["revision"], "reset": True})
    assert reset.status_code == 200
    after = _inputs(client, "2026-04-11")
    assert after["inherited_from"] == "2026-04-10"
    assert after["inputs"] == value


def test_regularity_inputs_reject_invalid_ids_lengths_and_empty_notes(client):
    current = _inputs(client, "2026-04-10")
    base = current["inputs"]
    invalid = {**base, "big_column_supervisors": {"am": "OFF-999"}}
    response = client.put("/api/ops/2026-04-10/regularity",
                          json={"revision": current["revision"], "inputs": invalid})
    assert response.status_code == 400

    current = _inputs(client, "2026-04-10")
    invalid = {**current["inputs"], "notes": ["   "]}
    response = client.put("/api/ops/2026-04-10/regularity",
                          json={"revision": current["revision"], "inputs": invalid})
    assert response.status_code == 400

    current = _inputs(client, "2026-04-10")
    invalid = {**current["inputs"], "notes": [str(i) for i in range(13)]}
    response = client.put("/api/ops/2026-04-10/regularity",
                          json={"revision": current["revision"], "inputs": invalid})
    assert response.status_code == 400

    current = _inputs(client, "2026-04-10")
    invalid = {**current["inputs"], "officer_extra": {"OFF-001": "x" * 501}}
    response = client.put("/api/ops/2026-04-10/regularity",
                          json={"revision": current["revision"], "inputs": invalid})
    assert response.status_code == 400


def test_regularity_put_get_round_trip_and_preview(client):
    current = _inputs(client, "2026-04-10")
    value = {"officer_extra": {}, "big_column_supervisors": {},
             "notes": ["ملاحظة اختبار"]}
    response = client.put("/api/ops/2026-04-10/regularity",
                          json={"revision": current["revision"], "inputs": value})
    assert response.status_code == 200
    result = response.get_json()
    assert result["inputs"] == value
    assert result["preview"]["notes"] == ["ملاحظة اختبار"]
    assert isinstance(result["preview"]["service_rows"], int)
    assert {row["id"] for row in result["officers"]} == {"OFF-001", "OFF-002"}
    assert "warnings" not in result["preview"]
    assert "operations_chief" not in result["inputs"]


def test_old_supervisor_id_is_read_as_free_text_and_operations_chief_is_ignored(client):
    current = _inputs(client, "2026-04-10")
    response = client.put("/api/ops/2026-04-10/regularity", json={
        "revision": current["revision"],
        "inputs": {"officer_extra": {}, "big_column_supervisors": {"am": "OFF-002"},
                   "notes": ["ملاحظة"], "operations_chief": "OFF-001"},
    })
    assert response.status_code == 200
    inputs = response.get_json()["inputs"]
    assert inputs["big_column_supervisors"]["am"] == {
        "rank": "نقيب", "name": "محمود علي", "phone": "0101"
    }
    assert "operations_chief" not in inputs


def test_regularity_docx_uses_saved_inputs(client):
    assert client.patch("/api/command", json={"مدير الإدارة": "OFF-001"}).status_code == 200
    assert client.patch("/api/command", json={"وكيل الإدارة": "OFF-002"}).status_code == 200
    current = _inputs(client, "2026-04-10")
    value = {"officer_extra": {"OFF-002": "مرور تعقبي محفوظ"},
             "big_column_supervisors": {"am": {"rank": "مقدم", "name": "مشرف محفوظ", "phone": ""}},
             "notes": ["ملاحظة مخصصة محفوظة"]}
    response = client.put("/api/ops/2026-04-10/regularity",
                          json={"revision": current["revision"], "inputs": value})
    assert response.status_code == 200

    exported = client.get("/api/ops/2026-04-10/regularity.docx")
    assert exported.status_code == 200
    text = _text(Document(io.BytesIO(exported.data)))
    assert "مرور تعقبي محفوظ" in text
    assert "ملاحظة مخصصة محفوظة" in text
    assert "محمود علي" in text
