from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from backend.campaign_export import CHECKPOINT_SPANS, DELIVERY_SPANS, GRID_WIDTHS, LAST_FIT, build_docx


def _text(doc):
    return [cell.text for row in doc.tables[0].rows for cell in row.cells]


def test_latest_format_shape_and_mapping():
    payload = {
        "date": "2026-09-29",
        "posts": [{"number": "10.", "name": "بوابة", "driver_text": "مجند حر",
                    "plate_text": "3279ب13", "type_text": "مايك ايفكو", "vehicle": None}],
        "lorries": [{"name": "الإفطار + المحور", "driver": {"kind": "text", "name": "سائق"},
                     "vehicle": {"plate": "ب15/5217", "type": "لوري", "origin": "internal"}}],
        "services": [], "daily": [], "appointments": [],
        "plates": {"rows": [["ب13/9416", "", "ب14/7895", "", "", "", ""]]},
    }
    doc = Document(BytesIO(build_docx(payload).read()))
    table = doc.tables[0]
    assert table.rows[0].cells[0]._tc.find(qn('w:tcPr')).find(qn('w:gridSpan')).get(qn('w:val')) == '15'
    spans = [tc.find(qn('w:tcPr')).find(qn('w:gridSpan'))
             for tc in table.rows[1]._tr.findall(qn('w:tc'))]
    assert [node.get(qn('w:val')) if node is not None else None for node in spans[:5]] == [None, '3', '4', '3', '4']
    assert len(table.rows[-1]._tr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc')) == 7
    assert table.rows[0].cells[0].text == "يومية المركبات اليوم الثلاثاء 29/9/2026م"
    values = _text(doc)
    assert "10." in values and "3279ب13" in values and "5217" in values
    assert "مايك ايفكو" in values and "ـــــ" not in values
    assert [c.width.twips for c in table.rows[0].cells[:1]] == [sum(GRID_WIDTHS)]
    assert LAST_FIT["ratio"] <= .94


def test_latest_format_empty_day_is_reopenable_and_one_page():
    payload = {"date": "2026-09-29", "posts": [], "lorries": [], "services": [],
               "daily": [], "appointments": [], "plates": {"rows": []}}
    doc = Document(BytesIO(build_docx(payload).read()))
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 4
    assert LAST_FIT["ratio"] <= .94


def test_latest_format_worst_case_fits():
    post = {"number": "11.", "name": "ارتكاز طويل", "driver_text": "سائق", "vehicle": None,
            "plate_text": "1234ب13", "type_text": "بيك اب"}
    delivery = {"name": "خدمة طويلة", "driver": {"kind": "text", "name": "سائق"},
                "vehicle": {"plate": "ب15/5217", "type": "بيك اب دوبل كبينة", "origin": "internal"}}
    payload = {"date": "2026-09-29", "posts": [post] * 11, "lorries": [delivery] * 4,
               "services": [delivery] * 20, "daily": [delivery] * 16,
               "appointments": [delivery] * 3,
               "plates": {"rows": [["x"] * 7] * 3}}
    Document(BytesIO(build_docx(payload).read()))
    assert LAST_FIT["ratio"] <= .94


def _delivery_rows(payload):
    doc = Document(BytesIO(build_docx(payload).read()))
    rows = doc.tables[0].rows
    start = next(i for i, row in enumerate(rows) if row.cells[0].text == "م"
                 and row.cells[1].text == "الخدمة") + 1
    return [[c.text for c in dict.fromkeys(row.cells)] for row in rows[start:]]


def _payload(**extra):
    base = {"date": "2026-09-29", "posts": [], "lorries": [], "services": [],
            "daily": [], "appointments": [], "plates": {"rows": []}}
    base.update(extra)
    return base


def test_driver_name_is_two_words_without_rank():
    # الأرشيف: اسمين بدون رتبة (1,076 من 1,418 اسم)، «عبد X» كلمة واحدة
    lorries = [
        {"name": "فيصل", "driver": {"kind": "personnel", "name": "محمود عبدالعاطي ابراهيم حسن", "role": "أمين شرطة"}},
        {"name": "القنصلية", "driver": {"kind": "personnel", "name": "علي وحيد علي", "role": "معاون أمن"}},
        {"name": "x", "driver": {"kind": "conscript", "name": "محمد عبد الرحمن سيد"}},
        {"name": "y", "driver": {"kind": "text", "name": "طارق"}},
        {"name": "z", "driver": None},
    ]
    rows = _delivery_rows(_payload(lorries=lorries))
    assert [r[2] for r in rows] == ["محمود عبدالعاطي", "علي وحيد", "محمد عبد الرحمن", "طارق", "ـــــ"]


def test_delivery_plate_is_bare_number_by_default():
    lorry = {"name": "فيصل", "driver": None,
             "vehicle": {"plate": "ب13/9416", "type": "لوري", "origin": "internal"}}
    assert _delivery_rows(_payload(lorries=[lorry]))[0][3] == "9416"


def test_delivery_plate_series_switch(monkeypatch):
    from backend import campaign_export
    monkeypatch.setattr(campaign_export, "DELIVERY_PLATE_WITH_SERIES", True)
    lorry = {"name": "فيصل", "driver": None,
             "vehicle": {"plate": "ب13/9416", "type": "لوري", "origin": "internal"}}
    assert _delivery_rows(_payload(lorries=[lorry]))[0][3] == "ب13/9416"


def test_service_label_wording_matches_archive():
    lorries = [{"name": "كـ109"}, {"name": "المحور+المحكمة"},
               {"name": "فيصل + كـ109 +تفتيش الجناين"}]
    services = [{"name": "ترحيلة العاشر", "time": "7ص", "driver": None, "vehicle": None},
                {"name": "مستشفى 9ص", "time": "9ص", "driver": None, "vehicle": None},
                {"name": "قول السلام", "time": "", "driver": None, "vehicle": None},
                {"name": "نيابة عامة 1", "time": "8ص", "party": "عتاقة", "driver": None, "vehicle": None}]
    labels = [r[1] for r in _delivery_rows(_payload(lorries=lorries, services=services))]
    assert labels == ["كـــ 109", "المحور + المحكمة", "فيصل + كـــ 109 + تفتيش الجناين",
                      "ترحيلة العاشر 7ص", "مستشفى 9ص", "قول السلام", "نيابة عامة 8ص"]


def test_empty_appointment_lorries_are_not_printed():
    empty = {"name": "المديرية", "driver": None, "vehicle": None, "attached": []}
    with_driver = {"name": "الأربعين", "driver": {"kind": "text", "name": "سائق جديد"},
                   "vehicle": None, "attached": []}
    with_vehicle = {"name": "الإدارة الجديدة", "driver": None, "attached": [],
                    "vehicle": {"plate": "ب15/5217", "type": "لوري", "origin": "internal"}}
    with_service = {"name": "المديرية + تفتيش", "driver": None, "vehicle": None,
                    "attached": [{"name": "تفتيش"}]}
    rows = _delivery_rows(_payload(appointments=[empty, with_driver, with_vehicle, with_service]))
    assert [r[1] for r in rows] == ["الأربعين", "الإدارة الجديدة", "المديرية + تفتيش"]
