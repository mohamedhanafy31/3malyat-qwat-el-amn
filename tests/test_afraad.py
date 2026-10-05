"""يومية الأفراد — الأساسية قايمة مقفولة (نفس أسماء وترتيب
`22-6-2026 افراد.docx` بالظبط، 22 خدمة)، وكل تفاصيلها (القائم بها/
التليفون/القوام/التسليح/الانتظام) بتتغيّر يوميًا. الطارئة + أي قسم
مخصّص بييجوا من اليومية التفصيلية مرتبين بمعاد الانتظام."""
from backend.afraad import BASIC_SERVICE_NAMES

DAY = "2026-04-10"
FIRST_ID = "AFB-01"  # مدرعة المديرية


def _add_assignment(client, day=DAY, **over):
    body = {"name": "خدمة", "kind": "خارجية", "section": "الخدمات الطارئة", "shift": "صباحية"}
    body.update(over)
    r = client.post(f"/api/assignments/{day}", json=body)
    assert r.status_code == 201, r.get_json()
    return r.get_json()


def test_basic_section_always_has_the_22_fixed_services_in_order(client):
    d = client.get(f"/api/afraad/{DAY}").get_json()
    assert [r["name"] for r in d["basic"]] == BASIC_SERVICE_NAMES
    assert len(BASIC_SERVICE_NAMES) == 22


def test_basic_entries_start_vacant_with_no_details(client):
    d = client.get(f"/api/afraad/{DAY}").get_json()
    row = d["basic"][0]
    assert row["id"] == FIRST_ID and row["name"] == "مدرعة المديرية"
    assert row["morning"] == {"name": "", "phone": ""}
    assert row["night"] == {"name": "", "phone": ""}
    assert row["count"] == "" and row["weapon"] == "" and row["schedule"] == ""


def test_setting_a_basic_entry_is_scoped_to_one_day(client):
    other_day = "2026-04-11"
    r = client.put(f"/api/afraad/{DAY}/basic/{FIRST_ID}", json={
        "morning_name": "م.ش/ محمد عطية", "morning_phone": "0121",
        "count": "1 مجند صبح + 1 سائق", "weapon": "خرطوش + كلبش", "schedule": "8ص / 8م"})
    assert r.status_code == 200, r.get_json()
    row = next(b for b in r.get_json()["basic"] if b["id"] == FIRST_ID)
    assert row["morning"] == {"name": "م.ش/ محمد عطية", "phone": "0121"}
    assert row["count"] == "1 مجند صبح + 1 سائق"
    assert row["weapon"] == "خرطوش + كلبش"
    assert row["schedule"] == "8ص / 8م"

    other = client.get(f"/api/afraad/{other_day}").get_json()
    other_row = next(b for b in other["basic"] if b["id"] == FIRST_ID)
    assert other_row["morning"] == {"name": "", "phone": ""}


def test_new_day_inherits_only_strength_weapon_and_schedule(client):
    client.put(f"/api/afraad/{DAY}/basic/{FIRST_ID}", json={
        "morning_name": "فرد اليوم السابق", "morning_phone": "0100",
        "count": "2 مجند", "weapon": "آلي + كلبش", "schedule": "8ص / 8م"})

    following = client.get("/api/afraad/2026-04-11").get_json()
    row = next(item for item in following["basic"] if item["id"] == FIRST_ID)
    assert row["count"] == "2 مجند"
    assert row["weapon"] == "آلي + كلبش"
    assert row["schedule"] == "8ص / 8م"
    assert row["morning"] == {"name": "", "phone": ""}
    assert row["inherited_from"] == DAY
    assert set(row["inherited_fields"]) == {"count", "weapon", "schedule"}

    changed = client.put(f"/api/afraad/2026-04-11/basic/{FIRST_ID}",
                         json={"count": "3 مجند", "weapon": "", "schedule": "9ص"})
    assert changed.status_code == 200
    row = next(item for item in changed.get_json()["basic"] if item["id"] == FIRST_ID)
    assert (row["count"], row["weapon"], row["schedule"]) == ("3 مجند", "", "9ص")
    assert row["inherited_fields"] == []


def test_linked_personnel_phone_is_filled_from_force_record(client):
    person = client.post("/api/person", json={
        "type": "personnel", "name": "فرد تجريبي", "code": "901",
        "phone": "01012345678", "join_date": "2020-01-01", "role": "أمين شرطة",
    }).get_json()
    response = client.put(f"/api/afraad/{DAY}/basic/{FIRST_ID}", json={
        "morning_person_id": person["id"], "morning_name": "أمين شرطة/ فرد تجريبي",
        "morning_phone": "00000000000",
    })
    assert response.status_code == 200, response.get_json()
    row = next(item for item in response.get_json()["basic"] if item["id"] == FIRST_ID)
    assert row["morning"]["person_id"] == person["id"]
    assert row["morning"]["phone"] == "01012345678"
    assert any(item["id"] == person["id"] and item["phone"] == "01012345678"
               for item in response.get_json()["personnel"])

    client.patch(f"/api/person/{person['id']}", json={"phone": "01112345678"})
    refreshed = client.get(f"/api/afraad/{DAY}").get_json()
    row = next(item for item in refreshed["basic"] if item["id"] == FIRST_ID)
    assert row["morning"]["phone"] == "01112345678"


def test_setting_an_entry_outside_the_fixed_list_404s(client):
    r = client.put(f"/api/afraad/{DAY}/basic/AFB-99", json={"morning_name": "x"})
    assert r.status_code == 404


def test_setting_a_basic_entry_on_a_closed_day_is_blocked(client):
    client.post(f"/api/day-status/{DAY}/close", json={})
    r = client.put(f"/api/afraad/{DAY}/basic/{FIRST_ID}", json={"morning_name": "x"})
    assert r.status_code == 409


def test_occasional_section_reflects_the_board(client):
    _add_assignment(client, name="تسفير الأربعين", time="9ص")
    d = client.get(f"/api/afraad/{DAY}").get_json()
    assert len(d["occasional"]) == 1
    assert d["occasional"][0]["label"] == "تسفير الأربعين"


def test_occasional_section_excludes_the_officer_boards_basic_and_targets(client):
    """قسم «الخدمات أساسية» و«الأهداف» بتوع اليومية التفصيلية (الضباط)
    مالهمش علاقة بيومية الأفراد — أساسية الأفراد قايمتها المقفولة
    المستقلة الخاصة بيها."""
    _add_assignment(client, name="دورية", section="الخدمات أساسية")
    _add_assignment(client, name="طارئة حقيقية", section="الخدمات الطارئة")
    d = client.get(f"/api/afraad/{DAY}").get_json()
    labels = {r["label"] for r in d["occasional"]}
    assert labels == {"طارئة حقيقية"}


def test_a_custom_section_lands_in_the_occasional_table_too(client):
    """أي تصنيف زيادة عن «الخدمات الطارئة» (زي «خدمات انتشار») بيتحط في
    نفس جدول الطوارئ هنا — نفس مبدأ `counts.py` واللوحة بالظبط."""
    _add_assignment(client, name="خدمة انتشار", section="خدمات انتشار")
    d = client.get(f"/api/afraad/{DAY}").get_json()
    labels = {r["label"] for r in d["occasional"]}
    assert "خدمة انتشار" in labels


def test_occasional_rows_are_sorted_by_schedule(client):
    _add_assignment(client, name="متأخرة", time="9ص")
    _add_assignment(client, name="مبكرة", time="7ص")
    _add_assignment(client, name="بلا معاد")
    d = client.get(f"/api/afraad/{DAY}").get_json()
    assert [r["label"] for r in d["occasional"]] == ["مبكرة", "متأخرة", "بلا معاد"]


def test_bootstrap_ships_the_afraad_page(client):
    r = client.get("/api/bootstrap/afraad")
    assert r.status_code == 200, r.get_json()


def test_export_returns_a_real_docx_file(client):
    import io

    from docx import Document

    client.put(f"/api/afraad/{DAY}/basic/{FIRST_ID}", json={"morning_name": "م.ش/ محمد عطية"})
    _add_assignment(client, name="تسفير الأربعين", time="9ص")
    r = client.get(f"/api/afraad/{DAY}/export.docx")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    doc = Document(io.BytesIO(r.data))
    assert len(doc.tables) == 1
    assert len(doc.tables[0].columns) == 14
    assert any("10/4/2026" in paragraph.text for paragraph in doc.paragraphs)
    text = "\n".join(p.text for t in doc.tables for row in t.rows for c in row.cells
                     for p in c.paragraphs)
    text = text.replace("ـ", "")
    assert "مدرعة المديرية" in text
    assert "كمين شرق النفق المستحدث" in text
    assert "تسفير الأربعين" in text
    assert "الخدمات الطارئة" in text
