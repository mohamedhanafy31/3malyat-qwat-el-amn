"""تصدير التفصيلية: نص الراحة/الإجازة بمدتها، لا قسم تقصيرات، وكل الخلايا في النص."""
import io

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import _Cell

from backend.board_export import build_docx, leave_text, status_period_text


def test_leave_text_formats_rest_and_vacation_with_dates():
    assert leave_text("شهرية", "2026-09-17", "2026-09-21", "2026-09-22") == "راحة شهرية (من 17/9 عودة 22/9)"
    assert leave_text("نصف شهرية", "2026-09-08", "2026-09-10", None) == "راحة نصف شهرية (من 8/9 عودة 11/9)"
    assert leave_text("إجازة مصيف", "2026-09-11", "2026-09-15", "2026-09-16") == "إجازة مصيف (من 11/9 عودة 16/9)"
    assert leave_text("أسبوعية", "2026-09-21", "2026-09-21", None) == "راحة أسبوعية"
    assert leave_text("شهرية", "2026-12-30", "2026-12-31", None) == "راحة شهرية (من 30/12 عودة 1/1)"


def test_status_period_text_hides_stop_wording_for_open_statuses():
    assert status_period_text({"start": "2026-10-23", "end": ""}) == "من يوم 23/10"
    assert status_period_text({"start": "2026-10-23", "end": "2026-10-25"}) == "من 23/10 حتى 25/10"


def _board(rests=(), outsiders=(), taqseera=(), admin=(), net=(), basic=()):
    return {"date": "2026-09-09", "sections": [
        {"name": "الخدمات أساسية", "type": "services", "rows": list(basic)},
        {"name": "عمل بالإدارة", "type": "officers", "rows": list(admin)},
        {"name": "الصافي", "type": "officers", "rows": list(net)},
        {"name": "الراحات", "type": "officers", "rows": list(rests)},
        {"name": "التقصيرات", "type": "officers", "rows": list(taqseera)},
        {"name": "الخوارج", "type": "officers", "rows": list(outsiders)}]}


def _tail_text(board):
    document = Document(io.BytesIO(build_docx(board).getvalue()))
    table = document.tables[0]
    return document, "\n".join(_Cell(tc, table).text for row in table.rows for tc in row._tr.tc_lst)


def test_export_prints_rest_dates_vacations_and_has_no_taqseera_section():
    board = _board(
        rests=[{"id": "O1", "name": "أ", "role": "مقدم", "type": "شهرية", "start": "2026-09-08",
                "end": "2026-09-14", "return_date": "2026-09-15"}],
        outsiders=[{"id": "O2", "name": "ب", "role": "رائد", "reason": "إجازة طارئة", "note": "",
                    "leave_type": "إجازة طارئة", "start": "2026-09-21", "end": "2026-09-21",
                    "return_date": "2026-09-22"}],
        taqseera=[{"id": "O3", "name": "جـ", "role": "نقيب", "note": "", "text": "متابعة عنابر",
                   "has_service": False},
                  {"id": "O4", "name": "د", "role": "نقيب", "note": "", "text": "على خدمة",
                   "has_service": True}])
    _document, text = _tail_text(board)
    assert "راحة شهرية (من 8/9 عودة 15/9)" in text
    assert "إجازة طارئة (من 21/9 عودة 22/9)" in text
    assert "التقصيرات" not in text and "تقصيرة" not in text
    assert "متابعة عنابر" not in text
    assert "على خدمة" not in text


def test_export_prints_taqseera_marker_beside_service_and_omits_net():
    row = {
        "name": "دورية", "label": "دورية صبح", "officers": [
            {"id": "O1", "name": "أحمد", "role": "نقيب", "taqseera": True}],
        "personnel": [], "conscripts": [], "vacant": False,
    }
    board = _board(
        basic=[row],
        net=[{"id": "O2", "name": "محمود", "role": "رائد", "taqseera": True}],
    )
    _document, text = _tail_text(board)
    # النص الكامل يشمل الجزء العلوي للخدمات والذيل الخاص بالصافي.
    document = Document(io.BytesIO(build_docx(board).getvalue()))
    all_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "نقيب/ أحمد + ت" in all_text
    # «الصافي» لا يُطبع في ملف Word إطلاقًا (لا عنوان ولا اسم)
    assert "رائد/ محمود" not in all_text
    assert "الصافي" not in text


def test_every_cell_is_centered_horizontally_and_vertically():
    document, _ = _tail_text(_board(admin=[{"id": "O1", "name": "أ", "role": "عقيد", "text": "عمل"}]))
    table = document.tables[0]
    for row in table.rows:
        for tc in row._tr.tc_lst:
            cell = _Cell(tc, table)
            assert cell.vertical_alignment == WD_ALIGN_VERTICAL.CENTER
            assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in cell.paragraphs)


def test_other_exports_center_their_cells_too(client):
    for kind in ("duty", "afraad", "campaign"):
        response = client.get(f"/api/{kind}/2026-03-01/export.docx")
        assert response.status_code == 200
        document = Document(io.BytesIO(response.data))
        table = document.tables[0]
        sample = [_Cell(tc, table) for row in table.rows[:6] for tc in row._tr.tc_lst]
        assert all(c.vertical_alignment == WD_ALIGN_VERTICAL.CENTER for c in sample), kind
