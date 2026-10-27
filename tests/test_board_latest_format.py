"""التفصيلية: تطابق تنسيق ملف Word الحقيقي (قالب نسخة من 2026/10/20)."""
import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell

from backend import board_export
from backend.board_export import LAST_FIT, build_docx
from backend.export_fit import PageFitError
from tools.fidelity.board import read_docx

ARCHIVE = Path("../../2026/10/20/20.docx")
GRID = [1835, 2975, 851, 2125, 2977]


def _doc(board=None):
    payload = board or {"date": "2026-10-20", "sections": []}
    return Document(io.BytesIO(build_docx(payload).getvalue()))


def _unique(row):
    return [_Cell(tc, row.table) for tc in row._tr.tc_lst]


def _board(admin=0, outsiders=0, basic=0, emergency=0, targets=0, prison=0):
    def svc(prefix, i):
        return {"id": f"{prefix}{i}", "label": f"{prefix} {i}", "name": f"{prefix} {i}",
                "officers": [], "personnel": [], "conscripts": [], "conscript_count": 1}
    return {"date": "2026-10-05", "sections": [
        {"name": "الخدمات أساسية", "type": "services", "rows": [svc("اساسية", i) for i in range(basic)]},
        {"name": "الخدمات الطارئة", "type": "services", "rows": [svc("طارئة", i) for i in range(emergency)]},
        {"name": "الأهداف", "type": "targets", "rows": [svc("هدف", i) for i in range(targets)]},
        {"name": "خدمات سجن قوات الأمن", "type": "services", "rows": [svc("سجن", i) for i in range(prison)]},
        {"name": "عمل بالإدارة", "type": "admin", "rows": [
            {"id": f"O{i}", "name": f"ضابط {i}", "role": "رائد", "text": f"عمل {i}"} for i in range(admin)]},
        {"name": "الراحات", "type": "officers", "rows": [
            {"id": f"R{i}", "name": f"مستريح {i}", "role": "نقيب", "type": "شهرية"}
            for i in range(outsiders)]},
    ]}


def _texts(document):
    table = document.tables[0]
    return "\n".join(" | ".join(c.text for c in _unique(row)) for row in table.rows)


def test_grid_tblw_and_floating_table_come_from_the_real_template():
    document = _doc(_board(basic=3, emergency=3, admin=2, outsiders=3))
    tbl = document.tables[0]._tbl
    grid = [int(c.get(qn("w:w"))) for c in tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(grid) == 10763 and len(grid) == 5
    for width, (low, high) in zip(grid, board_export.COLUMN_RANGES):
        assert low <= width <= high
    pr = tbl.tblPr
    assert pr.find(qn("w:tblW")).get(qn("w:w")) == "10763"
    assert pr.find(qn("w:tblpPr")) is not None and pr.find(qn("w:tblBorders")) is not None
    assert pr.find(qn("w:bidiVisual")) is not None
    assert pr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"


def test_every_tcw_equals_sum_of_spanned_grid_columns():
    document = _doc(_board(basic=4, emergency=6, targets=3, prison=2, admin=3, outsiders=5))
    GRID = [int(c.get(qn("w:w"))) for c in document.tables[0]._tbl.tblGrid.findall(qn("w:gridCol"))]
    for tr in document.tables[0]._tbl.tr_lst:
        start = 0
        for tc in tr.tc_lst:
            span = tc.tcPr.find(qn("w:gridSpan"))
            span = int(span.get(qn("w:val"))) if span is not None else 1
            assert int(tc.tcPr.find(qn("w:tcW")).get(qn("w:w"))) == sum(GRID[start:start + span])
            start += span
        assert start == len(GRID)          # كل صف يملأ الشبكة (بلا gridAfter)


def test_row_heights_are_at_least_never_exact_or_line_exact():
    document = _doc(_board(basic=4, emergency=6, admin=3, outsiders=5))
    for tr in document.tables[0]._tbl.tr_lst:
        height = tr.trPr.find(qn("w:trHeight"))
        assert height.get(qn("w:hRule")) is None
    assert 'lineRule="exact"' not in document.element.xml


def test_title_lives_in_the_header_part_and_no_leading_empty_paragraph():
    document = _doc({"date": "2026-10-21", "sections": []})
    header = document.sections[0].header
    assert header.paragraphs[0].text == "الخدمات عن يوم الأربعاء 21/10/2026م"
    run = header.paragraphs[0].runs[0]._r.rPr
    assert run.find(qn("w:sz")).get(qn("w:val")) == "36"
    assert run.find(qn("w:rFonts")).get(qn("w:cs")) == "Sultan bold"
    assert run.find(qn("w:u")) is not None
    first = document.element.body[0]
    assert first.tag == qn("w:tbl")              # الجدول أول عنصر: لا فقرة فارغة قبله


def test_run_fonts_follow_the_real_file():
    document = _doc(_board(basic=2, emergency=2, admin=1, outsiders=1))
    runs = list(document.tables[0]._tbl.iter(qn("w:r")))
    assert runs
    for run in runs:
        fonts = run.rPr.find(qn("w:rFonts"))
        assert fonts.get(qn("w:ascii")) == fonts.get(qn("w:hAnsi")) == "Bahnschrift SemiBold"
        assert fonts.get(qn("w:cs")) == "Times New Roman"
        assert fonts.get(qn("w:eastAsia")) is None
        assert run.rPr.find(qn("w:b")) is not None and run.rPr.find(qn("w:bCs")) is not None
        assert run.rPr.find(qn("w:rtl")) is not None
        assert run.rPr.find(qn("w:sz")).get(qn("w:val")) == run.rPr.find(qn("w:szCs")).get(qn("w:val"))


def test_font_size_is_in_the_real_range_and_shrinks_with_row_count():
    _doc(_board(basic=3, emergency=3, outsiders=1))
    small = LAST_FIT["size_half_points"]
    _doc(_board(basic=12, emergency=12, targets=8, prison=6, admin=14, outsiders=14))
    big = LAST_FIT["size_half_points"]
    assert small == 24 and 16 <= big < small
    assert "ratio" in LAST_FIT


def test_page_setup_and_section_order_match_the_real_file():
    document = _doc()
    sect = document.element.body.find(qn("w:sectPr"))
    assert [child.tag.split("}")[1] for child in sect] == [
        "headerReference", "pgSz", "pgMar", "pgBorders", "cols", "bidi", "rtlGutter", "docGrid"]
    mar = sect.find(qn("w:pgMar"))
    assert (mar.get(qn("w:top")), mar.get(qn("w:right")), mar.get(qn("w:bottom")),
            mar.get(qn("w:left"))) == ("249", "720", "4", "567")
    assert document.styles["Normal"].element.pPr.find(qn("w:bidi")) is not None
    assert 'compatibilityMode' in document.settings.element.xml
    assert 'w:val="15"' in document.settings.element.xml
    assert document.core_properties.author != "python-docx"


def test_cell_properties_and_paragraphs_match_real_cells():
    document = _doc(_board(basic=2, emergency=2))
    for tc in document.tables[0]._tbl.iter(qn("w:tc")):
        assert tc.tcPr.find(qn("w:vAlign")).get(qn("w:val")) == "center"
        ppr = tc.find(qn("w:p")).pPr
        assert ppr.find(qn("w:jc")).get(qn("w:val")) == "center"
        assert ppr.find(qn("w:spacing")) is None and ppr.find(qn("w:bidi")) is None
    title = document.tables[0]._tbl.tr_lst[0]
    assert all(tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "D9D9D9" for tc in title.tc_lst)
    assert all(tc.tcPr.find(qn("w:tcBorders")) is not None for tc in title.tc_lst)


def test_empty_board_has_no_padding_rows_or_empty_header_sections():
    document = _doc()
    texts = _texts(document)
    assert len(document.tables[0].rows) == 4     # عنوان + خوارج + صفّا المعسكر الفرعي
    assert "—" not in texts
    for header in ("الصافي", "عمل بالإدارة", "خدمات سجن", "الاهداف"):
        assert header not in texts


def test_sections_with_data_still_print_and_tail_grows_without_dropping_rows():
    document = _doc(_board(admin=16, outsiders=18, targets=10, prison=9))
    texts = _texts(document)
    for i in range(16):
        assert f"عمل {i}" in texts
    for i in range(18):
        assert f"مستريح {i}" in texts
    for i in range(9):
        assert f"سجن {i}" in texts
    for i in range(1, 10):
        assert f"هدف {i}" in texts
    assert "خدمات سجن قوات الامن" in texts


def test_status_column_is_vertically_merged_per_consecutive_group():
    document = _doc(_board(outsiders=6))
    marks = []
    for tr in document.tables[0]._tbl.tr_lst:
        if len(tr.tc_lst) == 5:
            node = tr.tc_lst[2].tcPr.find(qn("w:vMerge"))
            marks.append(None if node is None else node.get(qn("w:val")) or "continue")
    merged = [m for m in marks if m]
    assert merged[0] == "restart" and all(m == "continue" for m in merged[1:])
    assert len(merged) == 6


def test_estimate_slightly_over_page_degrades_to_smallest_size_not_error(monkeypatch):
    """تقدير 1.02 لا يفشل التصدير: أصغر مقاس حقيقي (8pt)."""
    monkeypatch.setattr(board_export, "_measure", lambda rows, grid, half, head: 735 * 1.02)
    build_docx(_board(basic=3))
    assert LAST_FIT["size_half_points"] == 16
    assert LAST_FIT["ratio"] == pytest.approx(1.02, abs=.01)


def test_clearly_impossible_estimate_still_raises(monkeypatch):
    monkeypatch.setattr(board_export, "_measure", lambda rows, grid, half, head: 735 * 1.5)
    with pytest.raises(PageFitError):
        build_docx(_board(basic=3))


def test_grid_matches_latest_real_archive_when_available():
    if not ARCHIVE.is_file():
        pytest.skip("archive is not available")
    path = Path("/tmp/board-latest-structure-test.docx")
    path.write_bytes(build_docx({"date": "2026-10-20", "sections": []}).getvalue())
    real, ours = read_docx(ARCHIVE), read_docx(path)
    assert len(ours["grid_widths"]) == len(real["grid_widths"]) == 5   # العرض الآن يتبع المحتوى
    for key in ("w", "h", "orient", "top", "right", "bottom", "left"):
        assert ours["page"][key] == real["page"][key]


# ---------------------------------------------------------------- متطلبات المالك الجديدة

def _person(pid, name, role="رائد", short=""):
    return {"id": pid, "name": name, "role": role, "short_name": short, "category": "officers"}


def _svc(name, officers=(), **extra):
    return {"id": name, "label": name, "name": name, "officers": list(officers), "personnel": [],
            "conscripts": [], "conscript_count": 0, **extra}


def _sections(basic=(), emergency=(), targets=(), admin=(), net=(), great=(), security=(), rests=()):
    return {"date": "2026-10-27", "sections": [
        {"name": "الخدمات أساسية", "type": "services", "rows": list(basic)},
        {"name": "الخدمات الطارئة", "type": "services", "rows": list(emergency)},
        {"name": "الأهداف", "type": "targets", "rows": list(targets)},
        {"name": "عمل بالإدارة", "type": "admin", "rows": list(admin)},
        {"name": "الصافي", "type": "officers", "rows": list(net)},
        {"name": "ضابط عظيم الإدارة", "type": "slots", "rows": list(great)},
        {"name": "ضابط الأمن بالإدارة", "type": "slots", "rows": list(security)},
        {"name": "الراحات", "type": "officers", "rows": list(rests)},
    ]}


def _grid_rows(document):
    out = []
    for tr in document.tables[0]._tbl.tr_lst:
        cells = []
        for tc in tr.tc_lst:
            vm = tc.tcPr.find(qn("w:vMerge"))
            cells.append(("".join(t.text or "" for t in tc.iter(qn("w:t"))),
                          None if vm is None else vm.get(qn("w:val")) or "continue"))
        out.append(cells)
    return out


def test_net_is_not_exported_at_all():
    board = _sections(net=[{"id": "N1", "name": "صافي الاسم", "role": "نقيب"}],
                      admin=[{"id": "A1", "name": "إداري", "role": "عقيد", "text": "مدير الإدارة"}])
    text = _texts(_doc(board))
    assert "الصافي" not in text and "صافي الاسم" not in text
    assert "مدير الإدارة" in text


def test_multi_officer_service_is_split_into_rows_with_vmerged_name():
    board = _sections(basic=[
        _svc("هدف كهرباء عتاقة", [_person("1", "ا", short="اول"), _person("2", "ب", short="ثان")]),
        _svc("خدمة فردية", [_person("3", "ج", short="ثالث")])])
    rows = _grid_rows(_doc(board))
    first = next(i for i, r in enumerate(rows) if r[0][0] == "هدف كهرباء عتاقة")
    assert rows[first][0][1] == "restart" and rows[first][1][0] == "رائد/ اول"
    assert rows[first + 1][0] == ("", "continue") and rows[first + 1][1][0] == "رائد/ ثان"
    assert rows[first + 2][0] == ("خدمة فردية", None)


def test_multi_officer_target_and_emergency_split_too():
    board = _sections(
        emergency=[_svc("طارئة", [_person("1", "ا", short="اول"), _person("2", "ب", short="ثان")],
                        conscript_count=2)],
        targets=[_svc("مشرف", [_person("3", "ج", short="ثالث"), _person("4", "د", short="رابع")])])
    text = _texts(_doc(board))
    for label in ("رائد/ ثالث", "رائد/ رابع", "رائد/ ثان"):
        assert label in text
    rows = _grid_rows(_doc(board))
    assert any(cell[1] == "restart" for row in rows for cell in row)


def test_great_and_security_officers_listed_flat_under_single_admin_header():
    shifts = lambda name, a, b: [
        {"id": None, "name": name, "shift": "صباحية", "slot": True, "officers": [a] if a else []},
        {"id": None, "name": name, "shift": "ليلية", "slot": True, "officers": [b] if b else []}]
    p1, p2 = _person("1", "ا", short="اول"), _person("2", "ب", short="ثان")
    board = _sections(great=shifts("ضابط عظيم", p1, p2), security=shifts("ضابط أمن", p1, p2),
                      admin=[{"id": "9", "name": "إداري", "role": "عقيد", "text": "وكيل الإدارة"}])
    text = _texts(_doc(board))
    assert text.count("عمل بالإدارة") == 1
    assert "ضابط عظيم الإدارة  " not in text            # لا عنوان مستقل
    assert "ضابط الأمن بالإدارة  " not in text
    assert "ضابط عظيم الإدارة فترة صباحية + ضابط الأمن بالإدارة فترة صباحية" in text
    assert "ضابط عظيم الإدارة فترة ليلية" in text


def test_short_name_is_used_in_every_cell_never_the_full_name():
    full = "محمد ابراهيم سليمان عبدالله"
    officer = _person("1", full, short="سليمان عبدالله")
    board = _sections(basic=[_svc("خ", [officer])], emergency=[_svc("ط", [officer])],
                      targets=[_svc("هـ", [officer])],
                      admin=[{**officer, "text": "عمل"}],
                      rests=[{**officer, "type": "شهرية"}])
    text = _texts(_doc(board))
    assert full not in text and "محمد ابراهيم" not in text
    assert "رائد/ سليمان عبدالله" in text


def test_columns_follow_content_within_real_ranges():
    long_names = [_person(str(i), "ن", short="اسم شهرة طويل جدا جدا جدا ومعه كلمات اكثر") for i in range(1)]
    wide = _doc(_sections(basic=[_svc("خدمة", long_names)] * 6, emergency=[_svc("ط")] * 6))
    grid_wide = list(LAST_FIT["grid"])
    narrow = _doc(_sections(basic=[_svc("خ", [_person("1", "ن", short="س")])] * 6))
    grid_narrow = list(LAST_FIT["grid"])
    assert grid_wide != grid_narrow
    for grid in (grid_wide, grid_narrow):
        assert sum(grid) == 10763
        for width, (low, high) in zip(grid, board_export.COLUMN_RANGES):
            assert low <= width <= high
    assert grid_wide[1] >= grid_narrow[1]


def test_clinic_and_search_members_are_removed_from_the_board_file():
    clinic, search, normal = (_person("C", "طبيب", short="الطبيب"), _person("S", "باحث", short="الباحث"),
                              _person("N", "عادي", short="العادي"))
    board = _sections(
        basic=[_svc("خدمة", [clinic, normal]), _svc("عيادة", [clinic], kind="طبية")],
        admin=[{**search, "text": "عمل البحث"}], net=[clinic],
        rests=[{**clinic, "type": "شهرية"}])
    text = _texts(_doc_with_excluded(board, {"C", "S"}))
    assert "الطبيب" not in text and "الباحث" not in text and "عمل البحث" not in text
    assert "العادي" in text and "عيادة" not in text.replace("العيادة", "")


def _doc_with_excluded(board, ids):
    return Document(io.BytesIO(build_docx(board, ids).getvalue()))


def test_outsiders_interleaved_kinds_are_grouped_and_merged_once_each():
    def rest(i, kind):
        return {"id": f"r{i}", "name": f"ر {i}", "type": kind}

    def out(i, reason):
        return {"id": f"o{i}", "name": f"خ {i}", "reason": reason}
    board = {"date": "2026-10-20", "sections": [
        {"name": "الراحات", "type": "officers", "rows": [rest(1, "راحة")]},
        {"name": "الخوارج", "type": "officers", "rows": [
            out(1, "فرقة"), out(2, "مرضي"), out(3, "فرقة"), out(4, "فرقة"),
            out(5, "دورة"), out(6, "مرضي")]}]}
    document = _doc(board)
    labels, marks = [], []
    for tr in document.tables[0]._tbl.tr_lst:
        if len(tr.tc_lst) == 5:
            tc = tr.tc_lst[2]
            node = tc.tcPr.find(qn("w:vMerge")) if tc.tcPr is not None else None
            mark = None if node is None else node.get(qn("w:val")) or "continue"
            text = "".join(t.text or "" for t in tc.iter(qn("w:t")))
            if text or mark:
                labels.append(text)
                marks.append(mark)
    restarts = [t for t, m in zip(labels, marks) if m == "restart"]
    singles = [t for t, m in zip(labels, marks) if m is None and t]
    shown = restarts + singles
    assert sorted(shown) == sorted(set(shown))  # no kind label repeated
    assert shown.count("فرقة") == 1 and shown.count("مرضي") == 1


def test_group_outsiders_by_kind_order_and_stability():
    rows = [{"status_type": k, "n": i} for i, k in enumerate(
        ["فرقة", "مرضي", "راحة", "فرقة", "xyz", "إجازة", "مرضي", "abc", "xyz"])]
    out = board_export.group_outsiders_by_kind(rows)
    assert [r["status_type"] for r in out] == [
        "راحة", "إجازة", "فرقة", "فرقة", "مرضي", "مرضي", "xyz", "xyz", "abc"]
    assert [r["n"] for r in out if r["status_type"] == "فرقة"] == [0, 3]
