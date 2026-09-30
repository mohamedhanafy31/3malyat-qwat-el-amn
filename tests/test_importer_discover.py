import csv
import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

from backend import store
from importer.discover import (
    classify_document,
    detect_roster_layout,
    parse_folder_date,
    read_docx,
    scan_archive,
    select_documents,
)
from importer.ledger import Ledger
from importer.textnorm import clean_text, parse_date, weekday_matches


def _docx(path, paragraphs=(), rows=()):
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    if rows:
        table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    document.save(path)


@pytest.mark.parametrize(
    ("name", "paragraphs", "rows", "role"),
    [
        ("1-1-2024.docx", ["يومية عن يوم الاثنين الموافق 1-1-2024م"],
         [["م", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي"]], "roster"),
        ("قديم.docx", ["يومية تشغيل السادة الضباط", "كشف راحات الضباط"], [], "old_officer_sheet"),
        ("#.docx", [], [["الخدمات الأساسية", "الخدمات الطارئة", "الأهداف"]], "board"),
        ("#.docx", [], [["الخدمة", "الخدمة الصباحية", "الخدمة الليلية", "قوام الخدمة", "التسليح", "الانتظام"]], "afraad"),
        ("#..docx", ["انتظام خدمات أمر الخدمة"], [["الخدمة", "الميعاد"]], "duty_list"),
        ("أعداد.docx", ["أعداد الخدمات يوم 1-1-2024"], [["الخدمة", "العدد"]], "counts"),
        ("القائد.docx", ["يومية القائد"], [["الخدمة", "العدد"]], "counts_leader"),
        ("أمر.docx", ["أمر خدمة يوم الاثنين"], [], "command_order"),
        ("ماتش.docx", ["خدمات مباراة اليوم"], [], "match"),
        ("انتشار.docx", ["خطة انتشار القوات"], [], "deployment"),
        ("ضباط.docx", [], [["الرتبة", "الأقدمية", "الاسم", "المسمى الوظيفي", "رقم الهاتف"]], "ref_officers"),
        ("أفراد.docx", [], [["الدرجة", "الاسم", "العمل المسند", "العنوان", "رقم الهاتف"]], "ref_personnel"),
        ("يومية الحملة.docx", [], [["بيان المركبات"]], "out_of_scope"),
        ("ملاحظات.docx", ["نص لا يطابق وثيقة مستوردة"], [], "other"),
    ],
)
def test_content_classifier_covers_every_role(tmp_path, name, paragraphs, rows, role):
    path = tmp_path / name
    _docx(path, paragraphs, rows)
    found, reason, _ = classify_document(name, read_docx(path))
    assert found == role
    if role in {"out_of_scope", "other"}:
        assert reason


@pytest.mark.parametrize(
    ("name", "paragraph", "rows", "role"),
    [
        ("#.docx", "", [["الخدمة", "الخدمة الصباحية", "الخدمة الليلية", "قوام الخدمة", "التسليح", "الانتظام"]], "afraad"),
        ("#.docx", "انتظام خدمات أمر الخدمة", [["الخدمة", "التوقيت"]], "duty_list"),
        ("#.docx", "", [["الخدمات الأساسية", "الخدمات الطارئة", "الأهداف"]], "board"),
    ],
)
def test_hash_filename_is_classified_by_its_three_content_eras(tmp_path, name, paragraph, rows, role):
    path = tmp_path / name
    _docx(path, [paragraph] if paragraph else [], rows)
    assert classify_document(name, read_docx(path))[0] == role


@pytest.mark.parametrize(
    ("header", "version"),
    [
        (["م", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي"], "roster_v1_basic"),
        (["0.", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي"], "roster_v1_corrupt_header"),
        (["م", "الرتبة", "رقم الأقدمية", "الاسم", "العمل المسند إليه", "التشغيل اليومي"], "roster_v2_seniority"),
        (["م", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي", "الراحات"], "roster_v4_rest"),
        (["م", "الرتبة", "الاسم", "الحالة", "رقم التلفون", "التشغيل اليومي"], "roster_temporary_variant"),
    ],
)
def test_roster_layout_signatures(tmp_path, header, version):
    path = tmp_path / "roster.docx"
    _docx(path, [], [header])
    assert detect_roster_layout(read_docx(path)) == version


def test_roster_temporary_reassigned_work_layout_has_no_operation_column(tmp_path):
    path = tmp_path / "roster.docx"
    _docx(
        path,
        ["يومية تشغيل الضباط عن يوم الثلاثاء الموافق 4-8-2026م"],
        [["م", "الرتبة", "الاسم", "العمل المسند إليه", "العمل المسند مجددا", "ملاحظات"]],
    )
    content = read_docx(path)
    assert classify_document(path.name, content)[0] == "roster"
    assert detect_roster_layout(content) == "roster_temporary_variant"


def _record(path, role, folder_date, title_date=None, filename_date=None, body="body", draft=False):
    return {
        "path": path,
        "folder_path": path.rsplit("/", 1)[0],
        "scope": "day",
        "role": role,
        "folder_date": folder_date,
        "draft_folder": draft,
        "title_date": title_date,
        "filename_date": filename_date,
        "extension": ".docx",
        "top_level": True,
        "body_sha1": body,
        "layout_version": None,
        "notes": [],
        "chosen_for": [],
        "stale_for": [],
        "evidence": True,
    }


def test_selection_uses_title_then_filename_and_adjacent_folder_and_marks_stale():
    records = [
        _record("2024/8/30/a.docx", "roster", "2024-08-30", title_date="2024-08-30", body="a"),
        _record("2024/8/30 - Copy/31.docx", "roster", "2024-08-30", title_date="2024-08-31", filename_date="2024-08-31", body="a"),
        _record("2024/8/30 - Copy/31 افراد.docx", "afraad", "2024-08-30", title_date="2024-08-31", filename_date="2024-08-31", body="b"),
        _record("2024/9/1/wrong.docx", "roster", "2024-09-01", filename_date="2024-08-31", body="c"),
    ]
    folders = [
        {"path": "2024/8/30", "kind": "day-folder", "date": "2024-08-30", "copy": False},
        {"path": "2024/8/30 - Copy", "kind": "day-folder", "date": "2024-08-30", "copy": True},
    ]
    rows, metrics = select_documents(records, folders, dt.date(2024, 8, 30), dt.date(2024, 8, 31))
    assert metrics["recovered_dates"] == ["2024-08-31"]
    assert rows[1]["chosen"]["roster"]["path"].endswith("31.docx")
    assert rows[1]["chosen"]["roster"]["stale"] is True
    assert records[1]["evidence"] is False


def test_stale_requires_the_immediately_previous_calendar_date():
    records = [
        _record("2024/1/1/a.docx", "roster", "2024-01-01", body="same"),
        _record("2024/1/3/a.docx", "roster", "2024-01-03", body="same"),
    ]
    folders = [
        {"path": "2024/1/1", "kind": "day-folder", "date": "2024-01-01", "copy": False},
        {"path": "2024/1/3", "kind": "day-folder", "date": "2024-01-03", "copy": False},
    ]
    rows, _ = select_documents(records, folders, dt.date(2024, 1, 1), dt.date(2024, 1, 3))
    assert rows[1]["chosen"]["roster"]["stale"] is False


def test_name_year_typo_is_corrected_to_folder_year_and_cannot_cross_year(tmp_path):
    archive = tmp_path / "archive"
    real = archive / "2024" / "3" / "1"
    typo = archive / "2025" / "3" / "1"
    real.mkdir(parents=True)
    typo.mkdir(parents=True)
    header = [["الخدمة", "الخدمة الصباحية", "الخدمة الليلية", "قوام الخدمة", "التسليح", "الانتظام"]]
    _docx(real / "1-3-2024 افراد.docx", [], header)
    _docx(typo / "1-3-2024 افراد.docx", [], header)
    records, folders = scan_archive(archive)
    typo_record = next(record for record in records if record["path"].startswith("2025/"))
    assert typo_record["original_filename_date"] == "2024-03-01"
    assert typo_record["filename_date"] == "2025-03-01"
    assert typo_record["notes"] == ["name year typo"]
    rows, _ = select_documents(records, folders, dt.date(2024, 3, 1), dt.date(2024, 3, 1))
    assert rows[0]["chosen"]["afraad"]["path"].startswith("2024/")


def test_real_folder_beats_same_date_draft_folder(tmp_path):
    archive = tmp_path / "archive"
    draft = archive / "2023" / "شهر 12" / "23-7-2024"
    real = archive / "2024" / "7" / "23-7-2024"
    draft.mkdir(parents=True)
    real.mkdir(parents=True)
    paragraphs = ["يومية تشغيل الضباط عن يوم الثلاثاء الموافق 23-7-2024م"]
    rows = [["م", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي"]]
    _docx(draft / "23-7-2024.docx", paragraphs, rows)
    _docx(real / "23-7-2024.docx", paragraphs, rows)
    records, folders = scan_archive(archive)
    assert next(folder for folder in folders if folder["path"].startswith("2023/"))["draft"] is True
    assert next(record for record in records if record["path"].startswith("2023/"))["draft_folder"] is True
    rows, _ = select_documents(records, folders, dt.date(2024, 7, 23), dt.date(2024, 7, 23))
    assert rows[0]["chosen"]["roster"]["path"].startswith("2024/")


def test_adjacent_folder_document_dated_another_day_is_not_a_candidate():
    records = [
        _record("2024/10/28/28.docx", "board", "2024-10-28",
                title_date="2024-10-28", filename_date="2024-10-28"),
    ]
    folders = [
        {"path": "2024/10/28", "kind": "day-folder", "date": "2024-10-28", "copy": False},
        {"path": "2024/10/29", "kind": "day-folder", "date": "2024-10-29", "copy": False},
    ]
    rows, _ = select_documents(records, folders, dt.date(2024, 10, 29), dt.date(2024, 10, 29))
    assert "board" not in rows[0]["chosen"]


@pytest.mark.parametrize(
    ("folder_date", "target"),
    [("2024-08-30", "2024-08-31"), ("2025-01-30", "2025-01-31")],
)
def test_copy_folder_can_hold_and_recover_the_next_day(folder_date, target):
    folder = f"{folder_date} - Copy"
    records = [
        _record(f"{folder}/roster.docx", "roster", folder_date,
                title_date=target, filename_date=folder_date),
        _record(f"{folder}/afraad.docx", "afraad", folder_date,
                title_date=target, filename_date=target),
    ]
    folders = [{"path": folder, "kind": "day-folder", "date": folder_date, "copy": True}]
    day = dt.date.fromisoformat(target)
    rows, metrics = select_documents(records, folders, day, day)
    assert metrics["recovered_dates"] == [target]
    assert rows[0]["chosen"]["roster"]["path"].startswith(folder)
    assert rows[0]["chosen"]["afraad"]["path"].startswith(folder)


def test_scan_skips_junk_keeps_hidden_docx_and_does_not_read_nested(tmp_path):
    archive = tmp_path / "archive"
    day = archive / "2024" / "شهر 10" / "1 الجمعة"
    nested = day / "هام"
    nested.mkdir(parents=True)
    _docx(day / ".26.docx", [], [["الخدمات الأساسية", "الخدمات الطارئة", "الأهداف"]])
    _docx(nested / "nested.docx", [], [["الخدمات الأساسية", "الخدمات الطارئة", "الأهداف"]])
    (day / "desktop.ini").write_text("junk")
    (day / "empty.txt").write_text("")
    records, folders = scan_archive(archive)
    assert [record["name"] for record in records] == [".26.docx"]
    assert folders[0]["nested_folders"] == ["هام"]


def test_text_and_folder_date_helpers():
    assert clean_text("\u200f1 ـ 2") == "1 2"
    assert parse_date(" 9 - 5 - 2024 م ") == dt.date(2024, 5, 9)
    assert parse_folder_date(2025, "شهر 11", "28 الجمعة") == dt.date(2025, 11, 28)
    assert weekday_matches("الجمعة", dt.date(2025, 11, 28)) is True


def test_ledger_layout_and_store_ignores_import_folder(tmp_path, monkeypatch):
    data_dir = tmp_path / "data"
    monkeypatch.setattr(store, "DATA_DIR", data_dir)
    store._cache.clear()
    store.explode({"schema": store.SCHEMA_VERSION})
    ledger = Ledger(data_dir, "batch")
    state = ledger.initialise({"from": "2024-01-01"}, {"discover": "1"})
    ledger.mark_stage(state, "discover", "complete", {"files": 0})
    (data_dir / "import" / "staging" / "batch" / "fake.json").write_text('{"day_assignments":{"bad":[]}}')
    assert json.loads((data_dir / "import" / "id_map.json").read_text()) == {}
    assert "bad" not in store._read().get("day_assignments", {})


def test_cli_requires_data_dir_and_extract_requires_discover(tmp_path):
    archive = tmp_path / "archive"
    archive.mkdir()
    missing = subprocess.run(
        [sys.executable, "-m", "importer", "discover", "--archive", str(archive)],
        text=True, capture_output=True,
    )
    assert missing.returncode == 2
    result = subprocess.run(
        [sys.executable, "-m", "importer", "extract", "--archive", str(archive),
         "--data-dir", str(tmp_path / "data")],
        text=True, capture_output=True,
    )
    assert result.returncode == 2
    assert "discover" in result.stderr
