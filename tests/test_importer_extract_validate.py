import datetime as dt
import hashlib
import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from importer.extract import extract_document, read_docx_raw, run_extract
from importer.ledger import Ledger, atomic_write_jsonl
from importer.validate import run_validate, validate_records


DAY = "2026-09-02"


def _docx(path, paragraphs=(), tables=()):
    document = Document()
    for value in paragraphs:
        document.add_paragraph(value)
    for rows in tables:
        table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    document.save(path)


def _meta(path, role, *, layout=None, stale=False, title_date=DAY):
    payload = path.read_bytes()
    return {
        "date": DAY,
        "role": role,
        "path": path.name,
        "sha1": hashlib.sha1(payload).hexdigest(),
        "body_sha1": "discover-body",
        "layout_version": layout,
        "stale": stale,
        "title_date": title_date,
    }


@pytest.mark.parametrize(
    ("header", "layout"),
    [
        (["م", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي"], "roster_v1_basic"),
        (["0.", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي"], "roster_v1_corrupt_header"),
        (["م", "الرتبة", "رقم الأقدمية", "الاسم", "العمل المسند إليه", "التشغيل اليومي"], "roster_v2_seniority"),
        (["م", "الرتبة", "الاسم", "العمل المسند إليه", "التشغيل اليومي", "الراحات"], "roster_v4_rest"),
        (["م", "الرتبة", "الاسم", "الحالة", "رقم التلفون", "التشغيل اليومي"], "roster_temporary_variant"),
    ],
)
def test_roster_extracts_every_layout_with_title_section_and_summary(tmp_path, header, layout):
    path = tmp_path / "roster.docx"
    row = ["1", "رائد", "100", "ضابط تجريبي", "عمل", "تشغيل", "الجمعة"][:len(header)]
    _docx(path, ["يومية يوم الأربعاء الموافق 2-9-2026م"], [
        [header, row, ["", "الحراسات المشددة"]],
        [["أصل القوة", "30"], ["الموجود", "29"]],
    ])
    records = extract_document(tmp_path, _meta(path, "roster", layout=layout))
    assert records[0]["record_type"] == "document"
    assert records[0]["body_hash"] != "discover-body"
    assert next(record for record in records if record["record_type"] == "title")["parsed_dates"] == [DAY]
    assert any(record["record_type"] == "section" for record in records)
    roster = next(record for record in records if record["record_type"] == "roster_row")
    assert roster["layout_version"] == layout
    assert roster["fields"]["الاسم"]
    assert any(record["record_type"] == "summary_cell" and record["header_path"]
               for record in records)


def test_docx_grid_span_is_one_cell_with_all_grid_columns(tmp_path):
    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "مدمج"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 2).text = "منفصل"
    document.save(path)
    row = read_docx_raw(path).tables[0].rows[0]
    assert [(cell.text, cell.columns) for cell in row] == [("مدمج", (0, 1)), ("منفصل", (2,))]


@pytest.mark.parametrize(("header", "layout"), [
    (["م", "الرتبة", "الاسم", "العمل المسند إليه", "ملاحظات"], "roster_v3_notes"),
    (["م", "الرتبة", "الاسم", "الإمضاء"], "roster_v3_signature"),
    (["م", "الرتبة", "الاسم", "العمل المسند إليه", "الإمضاء", "ملاحظات"],
     "roster_v3_signature_notes"),
])
def test_unknown_roster_headers_receive_specific_layouts(tmp_path, header, layout):
    path = tmp_path / "unknown.docx"
    _docx(path, [], [[header, ["1", "رائد", "ضابط تجريبي", "عمل", "", ""][:len(header)]]])
    records = extract_document(tmp_path, _meta(path, "roster", layout="unknown"))
    assert records[0]["layout_version"] == layout
    assert next(record for record in records if record["record_type"] == "roster_row")["fields"]["الاسم"]


def test_old_officer_sheet_keeps_preceding_heading(tmp_path):
    path = tmp_path / "old.docx"
    _docx(path, ["خدمات الإدارة"], [[
        ["رائد/ ضابط تجريبي", "خدمة تجريبية"],
    ]])
    records = extract_document(tmp_path, _meta(path, "old_officer_sheet"))
    row = next(record for record in records if record["record_type"] == "old_officer_row")
    assert row["heading"] == "خدمات الإدارة"
    assert row["rank_name"] == "رائد/ ضابط تجريبي"


def test_board_splits_two_halves_and_preserves_order(tmp_path):
    path = tmp_path / "board.docx"
    _docx(path, [], [[
        ["الخدمات أساسية", "", "الأهداف", ""],
        ["خدمة أولى", "2 فرد", "هدف أول", "فرد"],
        ["الخدمات الطارئة", "", "هدف ثان", "فردان"],
    ]])
    records = extract_document(tmp_path, _meta(path, "board"))
    header = next(record for record in records if record["record_type"] == "header")
    assert header["half_starts"] == [0, 2]
    rows = [record for record in records if record["record_type"] in {"board_row", "board_label"}]
    assert [(record["half"], record["label"]) for record in rows] == [
        (0, "خدمة أولى"), (0, "الخدمات الطارئة"), (1, "هدف أول"), (1, "هدف ثان")
    ]


def test_board_header_without_article_and_spanning_cell_stay_in_one_half():
    from importer.extract import Cell, DocxData, Table, _extract_board
    rows = (
        (Cell("خدمات أساسية", (0, 1, 2)), Cell("الأهداف", (3, 4))),
        (Cell("خدمة أولى", (0,)), Cell("رائد/ ضابط تجريبي", (1, 2)), Cell("هدف أول", (3,)), Cell("فرد", (4,))),
    )
    data = DocxData((), (), (Table(rows),), "x", ("",))
    records = _extract_board({"path": "b.docx", "date": "2025-11-10", "role": "board"}, data)
    header = next(record for record in records if record["record_type"] == "header")
    assert header["half_starts"] == [0, 3] and header["known_section"]
    pairs = [(record["half"], record["label"], record["manning"]) for record in records
             if record["record_type"] == "board_row"]
    assert pairs == [(0, "خدمة أولى", "رائد/ ضابط تجريبي"), (1, "هدف أول", "فرد")]


def test_afraad_merged_shift_is_reassigned_by_content(tmp_path):
    path = tmp_path / "afraad.docx"
    _docx(path, ["يوم الأربعاء الموافق 2-9-2026م"], [[
        ["الخدمة", "الخدمة الصباحية", "الخدمة الليلية", "قوام الخدمة", "التسليح", "الانتظام"],
        ["بوابة", "م.ش/ فرد تجريبي 01000000000", "3 مجند", "آلي", "8ص", ""],
        ["الخدمات الطارئة", "", "", "", "", ""],
        ["مأمورية", "م.ش/ قائد تجريبي", "2 مج", "خرطوش", "9ص", "جهة"],
    ]])
    records = extract_document(tmp_path, _meta(path, "afraad"))
    basic = next(record for record in records if record["record_type"] == "afraad_basic_row")
    assert basic["aligned"] == {
        "الخدمة الصباحية": "م.ش/ فرد تجريبي 01000000000",
        "قوام الخدمة": "3 مجند",
        "التسليح": "آلي",
        "الانتظام": "8ص",
    }
    emergency = next(record for record in records if record["record_type"] == "afraad_emergency_row")
    assert emergency["emergency"] is True


def test_afraad_basic_table_after_separate_emergency_table(tmp_path):
    path = tmp_path / "afraad-split.docx"
    _docx(path, ["يوم الأحد الموافق 9-11-2025م"], [
        [["الخدمات الطارئة", "", "", "", ""],
         ["مأمورية", "م.ش/ فرد طوارئ", "2 مجند", "آلي", "9ص"]],
        [["الخدمة", "الخدمة الصباحية", "الخدمة الليلية", "قوام الخدمة", "التسليح", "الانتظام"],
         ["بوابة", "م.ش/ فرد صبح", "م.ش/ فرد ليل", "2 مجند", "آلي", "8ص / 8م"]],
    ])
    records = extract_document(tmp_path, _meta(path, "afraad"))
    basic = [record for record in records if record["record_type"] == "afraad_basic_row"]
    assert len(basic) == 1
    assert basic[0]["aligned"]["قوام الخدمة"] == "2 مجند"


def test_duty_list_era_a_paragraphs_and_supervisor(tmp_path):
    path = tmp_path / "duty-a.docx"
    _docx(path, ["8ص", "بوابة: فرد تجريبي", "مشرف الخدمات : ضابط تجريبي"])
    records = extract_document(tmp_path, _meta(path, "duty_list"))
    assert any(record["record_type"] == "duty_time_heading" for record in records)
    assert any(record["record_type"] == "duty_supervisor" for record in records)


def test_duty_list_era_b_and_officers_table(tmp_path):
    path = tmp_path / "duty-b.docx"
    _docx(path, [], [
        [
            ["الخدمة", "قائم للخدمة", "ساعة الانتظام", "عدد المجندين", "بند قيام", "بند العودة", "مكان الانتظام", "المركبة", "رقم الهاتف"],
            ["بوابة", "فرد تجريبي", "8ص", "2", "", "", "المعسكر", "سيارة", "01000000000"],
        ],
        [["الرتبة", "الاسم", "الخدمة"], ["رائد", "ضابط تجريبي", "إشراف"]],
    ])
    records = extract_document(tmp_path, _meta(path, "duty_list"))
    assert any(record["record_type"] == "duty_row" and record["era"] == "B" for record in records)
    assert any(record["record_type"] == "duty_officer_row" for record in records)


def test_counts_docx_and_leader_keep_raw_rows(tmp_path):
    for role in ("counts", "counts_leader"):
        path = tmp_path / f"{role}.docx"
        _docx(path, ["أعداد الخدمات"], [[
            ["صباحية", "الخدمة", "عدد المجندين"], ["بوابة", "2"], ["المجموع", "2"]
        ]])
        records = extract_document(tmp_path, _meta(path, role))
        assert any(record["record_type"] == "counts_row" for record in records)
        assert any(record["record_type"] == "counts_total" for record in records)


@pytest.mark.parametrize("rows", [
    [
        ["الخدمة", "عدد المجندين", "الخدمة", "عدد المجندين", "الخدمة", "عدد المجندين"],
        ["بوابة", "2", "دورية", "3", "هدف", "1"],
        ["الإجمالي", "6", "", "", "", ""],
    ],
    [
        ["الخدمات الطارئة", "العدد"], ["مأمورية", "4"], ["المجموع", "4"],
    ],
])
def test_counts_docx_t0_and_t1_layouts_keep_totals(tmp_path, rows):
    path = tmp_path / "counts-layout.docx"
    _docx(path, [], [rows])
    records = extract_document(tmp_path, _meta(path, "counts"))
    total = next(record for record in records if record["record_type"] == "counts_total")
    assert total["raw_cells"] == rows[-1]


def test_counts_xlsx_keeps_values_and_sum_formula(tmp_path):
    path = tmp_path / "counts.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الأعداد"
    sheet.append(["صباحية", "العدد"])
    sheet.append(["بوابة", 2])
    sheet.append(["دورية", 3])
    sheet.append(["المجموع", "=SUM(B2:B3)"])
    workbook.save(path)
    records = extract_document(tmp_path, _meta(path, "counts"))
    total = next(record for record in records if record["record_type"] == "counts_total")
    assert total["formulas"] == {1: "=SUM(B2:B3)"}
    assert all(record["body_hash"] for record in records)


@pytest.mark.parametrize(
    ("role", "paragraphs", "table", "kind"),
    [
        ("command_order", ["بوابة : تفصيل تجريبي"], None, "command_row"),
        ("match", ["مباراة يوم 2-9-2026"], [["الخدمة", "القوة"], ["بوابة", "2"]], "match_row"),
        ("deployment", ["خطة انتشار 2-9-2026"], [["المكان", "القوة"], ["مكان", "2"]], "deployment_row"),
        ("ref_officers", ["كشف 2-9-2026"], [["الرتبة", "الاسم"], ["رائد", "ضابط تجريبي"]], "ref_officers_row"),
        ("ref_personnel", ["كشف 2-9-2026"], [["الدرجة", "الاسم"], ["م.ش", "فرد تجريبي"]], "ref_personnel_row"),
    ],
)
def test_remaining_families_extract_title_and_rows(tmp_path, role, paragraphs, table, kind):
    path = tmp_path / f"{role}.docx"
    _docx(path, paragraphs, [table] if table else [])
    records = extract_document(tmp_path, _meta(path, role))
    row = next(record for record in records if record["record_type"] == kind)
    assert row["raw_text"] or row["raw_cells"]
    if role == "command_order":
        assert row["service"] == "بوابة"
        assert row["detail"] == "تفصيل تجريبي"


def test_parse_failure_becomes_one_error_record(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not-a-zip")
    records = extract_document(tmp_path, _meta(path, "roster"))
    assert len(records) == 1
    assert records[0]["record_type"] == "extract_error"


def _record(kind, role="roster", path="sample.docx", **extra):
    return {
        "record_type": kind, "date": DAY, "role": role, "path": path,
        "file_sha1": "sha", "body_hash": "body", "layout_version": "layout",
        "stale": False, "table_index": None, "row_index": None,
        "column_indexes": [], "raw_text": "", "raw_cells": [], **extra,
    }


def test_validation_date_weekday_stale_and_roster_rules():
    records = [_record("document", stale=True)]
    records.append(_record("title", parsed_dates=[DAY], weekday="الأربعاء"))
    for index in range(20):
        records.append(_record("roster_row", raw_cells=[str(index + 1), "رائد", f"ضابط {index}"],
                               fields={"الاسم": f"ضابط {index}"}))
    records.append(_record("roster_row", raw_cells=["21"], fields={"الاسم": ""}))
    validations, quarantined = validate_records(records)
    assert validations[0]["verdict"] == "warn"
    assert validations[0]["reason_code"] == "stale"
    assert not quarantined


def test_validation_warns_when_roster_has_fewer_than_twenty_officers():
    records = [_record("document")]
    records.extend(_record("roster_row", raw_cells=[str(index), f"ضابط {index}"],
                           fields={"الاسم": f"ضابط {index}"}) for index in range(19))
    validations, _ = validate_records(records)
    assert validations[0]["verdict"] == "warn"
    assert validations[0]["reason_code"] == "roster_too_few_rows"


def test_validation_quarantines_bad_dates_weekdays_and_structures():
    records = [
        _record("document"),
        _record("title", parsed_dates=["2026-09-04"], weekday="الخميس"),
        _record("roster_row", raw_cells=["1", "رائد"], fields={"الاسم": ""}),
    ]
    validations, quarantined = validate_records(records)
    assert validations[0]["verdict"] == "quarantine"
    assert {"title_date_differs_content_unique", "weekday_mismatch", "roster_name_missing"} <= set(validations[0]["reason_codes"])
    assert len(quarantined) == len(records)

    for role, code in (("board", "board_section_missing"), ("afraad", "afraad_header_missing")):
        result, quarantined = validate_records([_record("document", role=role, path=f"{role}.docx")])
        assert result[0]["reason_code"] == code
        assert quarantined


def test_validation_allows_adjacent_date_only_for_three_roles():
    for role, verdict in (("duty_list", "warn"), ("command_order", "warn"),
                          ("afraad", "quarantine"), ("board", "quarantine")):
        records = [_record("document", role=role, path=f"{role}.docx"),
                   _record("title", role=role, path=f"{role}.docx",
                           parsed_dates=["2026-09-01"], weekday="الثلاثاء")]
        if role == "afraad":
            records.append(_record("afraad_header", role=role, path=f"{role}.docx"))
            verdict = "warn"
        validations, _ = validate_records(records)
        assert validations[0]["verdict"] == verdict


def _roster_validation_fixture(day, path, title, duties):
    records = [_record("document", path=path, source_title_date=title)]
    records.extend(_record("roster_row", path=path, raw_cells=[name, duty],
                           fields={"الاسم": name, "التشغيل اليومي": duty})
                   for name, duty in duties)
    return records


def test_weekday_mismatch_with_exact_title_is_only_a_warning():
    duties = [(f"ضابط {index}", f"خدمة {index}") for index in range(20)]
    records = _roster_validation_fixture(DAY, "current.docx", DAY, duties)
    records.append(_record("title", path="current.docx", parsed_dates=[DAY], weekday="الخميس"))
    validations, quarantined = validate_records(records)
    assert validations[0]["verdict"] == "warn"
    assert validations[0]["reason_code"] == "weekday_mismatch"
    assert not quarantined


def test_title_month_or_year_typo_is_warning_when_folder_and_filename_match():
    duties = [(f"ضابط {index}", f"خدمة {index}") for index in range(20)]
    records = _roster_validation_fixture(DAY, "current.docx", "2026-08-02", duties)
    metadata = {(DAY, "roster", "current.docx"): {
        "filename_date": DAY, "folder_date": DAY,
    }}
    validations, quarantined = validate_records(records, metadata)
    assert validations[0]["verdict"] == "warn"
    assert validations[0]["reason_code"] == "title_typo"
    assert not quarantined


@pytest.mark.parametrize(("same_rows", "verdict", "reason"), [
    (19, "quarantine", "copy_of_other_day:2026-09-03"),
    (17, "warn", "title_date_differs_content_unique"),
])
def test_title_date_difference_uses_ninety_percent_content_threshold(same_rows, verdict, reason):
    current = [(f"ضابط {index}", f"خدمة {index}") for index in range(20)]
    other = [(name, duty if index < same_rows else f"خدمة مختلفة {index}")
             for index, (name, duty) in enumerate(current)]
    records = _roster_validation_fixture(DAY, "current.docx", "2026-09-03", current)
    other_records = _roster_validation_fixture("2026-09-03", "other.docx", "2026-09-03", other)
    for record in other_records:
        record["date"] = "2026-09-03"
    records.extend(other_records)
    metadata = {(DAY, "roster", "current.docx"): {
        "filename_date": DAY, "folder_date": DAY,
    }}
    validations, quarantined = validate_records(records, metadata)
    assert validations[0]["verdict"] == verdict
    assert validations[0]["reason_code"] == reason
    comparison = next(item for item in validations[0]["similarities"]
                      if item["date"] == "2026-09-03")
    assert comparison == {
        "date": "2026-09-03", "identical_rows": same_rows,
        "compared_rows": 20, "similarity": same_rows / 20,
    }
    assert bool(quarantined) is (verdict == "quarantine")


def test_xlsx_total_mismatch_warns():
    records = [
        _record("document", role="counts", path="counts.xlsx"),
        _record("counts_row", role="counts", path="counts.xlsx", table_index=0, row_index=0,
                raw_cells=["خدمة", "2"]),
        _record("counts_row", role="counts", path="counts.xlsx", table_index=0, row_index=1,
                raw_cells=["خدمة", "4"]),
        _record("counts_total", role="counts", path="counts.xlsx", table_index=0, row_index=2,
                raw_cells=["المجموع", "5"], formulas={1: "=SUM(B1:B2)"}),
    ]
    validations, _ = validate_records(records)
    assert validations[0]["verdict"] == "warn"
    assert validations[-1]["reason_code"] == "xlsx_total_mismatch"


def test_extract_validate_pipeline_is_deterministic_and_resume_safe(tmp_path):
    archive = tmp_path / "archive"
    archive.mkdir()
    source = archive / "board.docx"
    _docx(source, [], [[
        ["الخدمات أساسية", "", "الأهداف", ""], ["بوابة", "فرد", "هدف", "فرد"]
    ]])
    data_dir = tmp_path / "data"
    ledger = Ledger(data_dir, "batch")
    params = {"archive": str(archive), "from": DAY, "to": DAY, "batch": "batch"}
    state = ledger.initialise(params, {"discover": "1", "extract": "1", "validate": "1"})
    item = _meta(source, "board")
    item.update({"record_type": "file", "chosen_for": [DAY], "stale_for": []})
    atomic_write_jsonl(ledger.staging_path("discover"), [item])
    first = run_extract(archive, ledger, dt.date.fromisoformat(DAY), dt.date.fromisoformat(DAY), state)
    extract_bytes = ledger.staging_path("extract").read_bytes()
    second = run_extract(archive, ledger, dt.date.fromisoformat(DAY), dt.date.fromisoformat(DAY), state)
    assert first == second
    assert ledger.staging_path("extract").read_bytes() == extract_bytes
    result = run_validate(ledger, state)
    validate_bytes = ledger.staging_path("validate").read_bytes()
    assert result["documents"] == 1
    assert result["quarantined_documents"] == 0
    assert run_validate(ledger, state, resume=True) == result
    assert ledger.staging_path("validate").read_bytes() == validate_bytes
    batch = json.loads(ledger.batch_file.read_text(encoding="utf-8"))
    assert batch["extractor_versions"] == {"discover": "1", "extract": "1", "validate": "1"}
