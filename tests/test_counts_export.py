from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from backend.counts_export import build_xlsx, compute_totals


def _payload(am=1, pm=1, recurring=1, emergency=1, custom=1):
    def rows(prefix, n):
        return [{"name": f"{prefix}{i}", "party": "جهة" if i == 0 else "", "count": i + 1}
                for i in range(n)]
    return {"date": "2026-09-29", "basic_am": rows("ص", am), "basic_pm": rows("م", pm),
            "recurring": rows("ت", recurring), "emergency": rows("ي", emergency),
            "custom_sections": [{"name": "فرد + مج", "rows": rows("خ", custom)}] if custom else []}


def _sheet(payload):
    return load_workbook(BytesIO(build_xlsx(payload)), data_only=False).active


def test_default_v2_geometry_formulas_and_page_setup():
    ws = _sheet(_payload(25, 28, 3, 0, 0))
    assert ws.title == "اعداد الخدمات"
    assert ws.sheet_view.rightToLeft is True
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 9
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) == (1, 1)
    assert ws.calculate_dimension() == "A1:H52"
    assert {str(r) for r in ws.merged_cells.ranges} == {
        "B2:H2", "B25:H25", "H4:H12", "H14:H23", "H26:H34", "H35:H36",
        "C41:E48", "F41:F48",
    }
    assert {c.coordinate: c.value for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")} == {
        "H4": "=SUM(C4:C12)+SUM(E4:E12)+SUM(G4:G12)",
        "H14": "=SUM(C14:C23)+SUM(E14:E23)+SUM(G14:G23)",
        "C36": "=SUM(C27:C35)", "E36": "=SUM(E27:E35)", "G36": "=SUM(G27:G35)",
        "H35": "=C36+E36+G36", "F41": "=H4+H14+H35",
    }
    assert ws["B2"].value == "اعداد الخدمات الصباحية والمسائية الاساسية يوم الثلاثاء 29 / 9 / 2026 م"


def test_large_payload_grows_all_blocks_and_keeps_column_major_order():
    ws = _sheet(_payload(40, 35, 20, 30, 10))
    # AM=14, PM=12, emergency=60 -> all downstream coordinates move.
    assert ws["B4"].value == "ص0 (جهة)"
    assert ws["B17"].value == "ص13"
    assert ws["D4"].value == "ص14"
    assert ws["B19"].value == "م0 (جهة)"
    assert ws["D19"].value == "م12"
    assert ws["B34"].value == "ت0 (جهة)"
    assert ws["D34"].value == "ي0 (جهة)"
    assert ws["F34"].value == "ي20"
    assert ws["F44"].value == "خ0 (جهة)"
    # subtotal is after emergency data: AM end 17, PM end 30, emergency end 53.
    assert ws["C54"].value == "=SUM(C34:C53)"
    assert ws["F59"].value == "=H4+H19+H53"


def test_party_rule_and_emergency_order_include_custom_rows():
    ws = _sheet(_payload(1, 1, 1, 1, 1))
    assert ws["B4"].value == "ص0 (جهة)"
    assert ws["B14"].value == "م0 (جهة)"
    assert ws["B27"].value == "ت0 (جهة)"
    assert ws["B28"].value == "ي0 (جهة)"
    assert ws["B29"].value == "خ0 (جهة)"


def test_compute_totals_matches_all_payload_counts():
    payload = _payload(4, 5, 2, 3, 4)
    assert compute_totals(payload) == {"basic_am": 10, "basic_pm": 15,
                                       "emergency": 19, "grand_total": 44}


def test_special_sections_get_separate_titled_blocks_and_grand_total():
    payload = _payload(0, 0, 1, 1, 0)
    payload["custom_sections"] = [
        {"name": "مباراة المصري", "rows": [{"name": "ملعب", "count": 5}]},
        {"name": "خطة الانتشار", "rows": [{"name": "انتشار", "count": 7}]},
    ]
    ws = _sheet(payload)
    titles = [ws.cell(row, 2).value for row in range(1, ws.max_row + 1)]
    assert "خدمات مباراة المصري" in titles
    assert "خدمات خطة الانتشار" in titles
    assert ws["F49"].value == "=H4+H14+H35+H39+H43"
    assert ws["C40"].value == "=SUM(C39:C39)"
    assert ws["C44"].value == "=SUM(C43:C43)"


def test_archive_v2_structure_when_reference_exists():
    path = Path(__file__).resolve().parents[3] / "2026/9/29/اعداد الخدمات يوم 29-9-2026.xlsx"
    if not path.exists():
        import pytest
        pytest.skip("archive reference workbook is absent")
    reference = load_workbook(path, data_only=False).active
    generated = _sheet(_payload(25, 28, 3, 0, 0))
    assert generated.max_row == reference.max_row == 52
    assert generated.max_column == reference.max_column == 8
    assert {str(r) for r in generated.merged_cells.ranges} == {str(r) for r in reference.merged_cells.ranges}
    assert {c.coordinate for row in generated.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")} == {
        c.coordinate for row in reference.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}
    for coord in ("B2", "B25", "B3", "H26"):
        assert (generated[coord].font.name, generated[coord].font.sz) == (reference[coord].font.name, reference[coord].font.sz)


def test_export_route_returns_excel(client):
    response = client.get("/api/counts/2026-04-10/export.xlsx")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "2026" in response.headers.get("Content-Disposition", "")
    assert load_workbook(BytesIO(response.data), data_only=False).active.title == "اعداد الخدمات"
    assert client.get("/api/counts/not-a-day/export.xlsx").status_code == 400
