"""يومية الضباط: خدمة الضابط التلقائية من اللوحة + تشغيل إضافي يدوي،
وإضافته (اختياريًا) لقسم «عمل بالإدارة» في اليومية التفصيلية."""

import io

from docx import Document

DAY = "2026-03-01"


def _row(client, person_id="OFF-002", day=DAY):
    return next(r for r in client.get(f"/api/duty/{day}").get_json()["rows"]
                if r["id"] == person_id)


def _board_section(client, name, day=DAY):
    sections = client.get(f"/api/board/{day}").get_json()["sections"]
    return next(s for s in sections if s["name"] == name)


def _assign(client, officer="OFF-002"):
    r = client.post(f"/api/assignments/{DAY}",
                    json={"name": "تدخل سريع", "kind": "خارجية", "shift": "صباحية",
                          "section": "الخدمات أساسية", "officer_ids": [officer]})
    assert r.status_code in (200, 201), r.get_data(as_text=True)
    return _row(client, officer)["services"][0]


def test_text_is_the_automatic_service_plus_the_manual_extra(client):
    service = _assign(client)
    auto = f"{service['name']} فترة صباحية"
    assert _row(client)["text"] == auto

    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "متابعة اعمال التعيينات"})
    row = _row(client)
    assert row["text"] == f"{auto} + متابعة اعمال التعيينات"
    assert row["extra"] == "متابعة اعمال التعيينات"


def test_extra_without_a_service_is_the_whole_text(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "عمل بالادارة"})
    assert _row(client)["text"] == "عمل بالادارة"


def test_legacy_full_note_replaces_the_service_so_it_is_not_repeated(client):
    service = _assign(client)
    client.put(f"/api/duty/{DAY}/OFF-002", json={"note": f"{service['name']} فترة صباحية + متابعة"})
    assert _row(client)["text"] == f"{service['name']} فترة صباحية + متابعة"
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "اجتماع"})
    assert _row(client)["text"].endswith("+ متابعة + اجتماع")


def test_extra_and_admin_flag_are_independent_of_other_state_fields(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"taqseera": True, "extra": "أ", "admin_work": True})
    client.put(f"/api/duty/{DAY}/OFF-002", json={"note": "ب"})
    row = _row(client)
    assert (row["taqseera"], row["extra"], row["admin_work"], row["note"]) == (True, "أ", True, "ب")
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "", "admin_work": False})
    assert (_row(client)["extra"], _row(client)["admin_work"]) == ("", False)


def test_empty_optional_fields_are_not_stored(client, data_file):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"note": "x", "extra": "", "admin_work": False})
    from backend import store
    stored = store.load_data([DAY])["day_officers"][DAY]["OFF-002"]
    assert "extra" not in stored and "admin_work" not in stored


def test_too_long_extra_is_rejected(client):
    response = client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "ا" * 501})
    assert response.status_code == 400


def test_clear_state_removes_extra_and_admin_flag(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "أ", "admin_work": True})
    assert client.delete(f"/api/duty/{DAY}/OFF-002").status_code == 200
    row = _row(client)
    assert row["extra"] == "" and row["admin_work"] is False


def test_flag_adds_the_full_text_to_admin_work_and_keeps_the_service_place(client):
    service = _assign(client)
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "متابعة", "admin_work": True})
    full = f"{service['name']} فترة صباحية + متابعة"

    admin = _board_section(client, "عمل بالإدارة")["rows"]
    mine = [r for r in admin if r["id"] == "OFF-002"]
    assert len(mine) == 1 and mine[0]["text"] == full and mine[0]["added"] is True

    sections = client.get(f"/api/board/{DAY}").get_json()["sections"]
    service_rows = [row for s in sections if s["type"] in ("services", "slots", "targets")
                    for row in s["rows"]]
    assert any("OFF-002" in [o.get("id") for o in (row.get("officers") or [])]
               for row in service_rows)


def test_unflagged_officer_with_a_service_is_not_in_admin_work(client):
    _assign(client)
    admin = _board_section(client, "عمل بالإدارة")["rows"]
    assert "OFF-002" not in [r["id"] for r in admin]


def test_net_officer_is_listed_once_even_when_flagged(client):
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "متابعة", "admin_work": True})
    admin = _board_section(client, "عمل بالإدارة")["rows"]
    assert [r["id"] for r in admin].count("OFF-002") == 1
    assert next(r for r in admin if r["id"] == "OFF-002")["text"] == "متابعة"
    assert "OFF-002" not in [r["id"] for r in _board_section(client, "الصافي")["rows"]]


def test_net_officer_needs_operation_text_before_moving_to_admin(client):
    response = client.put(f"/api/duty/{DAY}/OFF-002", json={"admin_work": True})
    assert response.status_code == 400
    assert "النص التشغيلي" in response.get_json()["error"]
    assert "OFF-002" in [r["id"] for r in _board_section(client, "الصافي")["rows"]]


def test_flag_is_ignored_for_an_officer_among_the_outsiders(client):
    _assign(client)
    client.put(f"/api/duty/{DAY}/OFF-002", json={"status": "غياب", "admin_work": True})
    admin = _board_section(client, "عمل بالإدارة")["rows"]
    assert "OFF-002" not in [r["id"] for r in admin]


def test_duty_export_prints_service_plus_extra(client):
    _assign(client)
    client.put(f"/api/duty/{DAY}/OFF-002", json={"extra": "اجتماع غرفة الازمات"})
    response = client.get(f"/api/duty/{DAY}/export.docx")
    document = Document(io.BytesIO(response.data))
    text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "فترة صباحية + اجتماع غرفة الازمات" in text


def test_duty_export_lists_net_names_on_separate_lines(client):
    response = client.get(f"/api/duty/{DAY}/export.docx")
    document = Document(io.BytesIO(response.data))
    # الملف الحقيقي (17 عمودًا من يوم 16): خلية الأسماء آخر عمود، فقرة لكل ضابط
    cell = document.tables[1].rows[1].cells[16]
    assert "\n" in cell.text and cell.text.count("/") >= 2
    assert all(" / " in line for line in cell.text.split("\n"))
