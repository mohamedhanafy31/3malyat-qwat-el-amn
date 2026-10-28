"""تصدير اليومية التفصيلية كملف Word حقيقي (`/api/board/<day>/export.docx`)
— بيتولّد من نفس بيانات `/api/board` (`backend/board_export.py`)."""
import io

from docx import Document
from backend import store

DAY = "2026-04-10"


def _download(client, day=DAY):
    r = client.get(f"/api/board/{day}/export.docx")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return Document(io.BytesIO(r.data))


def _all_text(doc):
    """`document.tables` بترجع الجداول العليا بس — جداول الأقسام متداخلة
    جوّه خانات جدول العمودين، فلازم ننزل جوّه كل خانة كمان."""
    out = []

    def walk_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    out.extend(p.text for p in cell.paragraphs)
                    walk_tables(cell.tables)

    walk_tables(doc.tables)
    return "\n".join(out)


def test_export_returns_a_real_docx_file(client):
    doc = _download(client)
    assert len(doc.tables) >= 1


def test_export_rejects_an_invalid_day(client):
    r = client.get("/api/board/not-a-day/export.docx")
    assert r.status_code == 400


def test_export_reflects_an_assigned_target_officer(client):
    client.put("/api/board/2026-04-10/target/سوميد", json={"officer_ids": ["OFF-001"]})
    doc = _download(client)
    assert "سوميد" in _all_text(doc)


def test_export_filename_carries_the_day(client):
    r = client.get(f"/api/board/{DAY}/export.docx")
    disposition = r.headers.get("Content-Disposition", "")
    assert DAY in disposition


def test_export_ignores_preserved_authoritative_roster_and_generates(client):
    path = store.board_source_doc_path(DAY)
    path.parent.mkdir(parents=True, exist_ok=True)
    source = Document()
    source.add_paragraph("النص الأصلي كما ورد في اليومية")
    source.add_table(rows=2, cols=5).cell(0, 0).text = "الخدمات الأساسية"
    source.save(path)

    response = client.get(f"/api/board/{DAY}/export.docx")

    assert response.status_code == 200
    assert response.data != path.read_bytes()
    generated = Document(io.BytesIO(response.data))
    assert len(generated.tables) == 1
    # بلا صفوف حشو: عنوان + «الاهداف» + مشرف الأهداف وأهدافه + عنوان الخوارج + صفّا المعسكر الفرعي
    assert 10 <= len(generated.tables[0].rows) < 34


def test_generated_roster_uses_word_strength_columns_without_service_details(client):
    for payload in (
        {"name": "خدمة أساسية اختبار", "section": "الخدمات أساسية", "kind": "خارجية",
         "shift": "صباحية", "officer_ids": ["OFF-001"], "weapon": "آلي", "time": "8ص"},
        {"name": "خدمة طارئة اختبار", "section": "الخدمات الطارئة", "kind": "خارجية",
         "shift": "صباحية", "officer_ids": ["OFF-002"], "weapon": "آلي", "time": "12ظ",
         "conscript_count": 2},
    ):
        assert client.post(f"/api/assignments/{DAY}", json=payload).status_code == 201
    table = _download(client).tables[0]
    basic = next(row for row in table.rows if row.cells[0].text.strip() == "خدمة أساسية اختبار صبح")
    emergency = next(row for row in table.rows
                     if "خدمة طارئة اختبار" in " | ".join(cell.text for cell in row.cells))
    assert "آلي" not in basic.cells[1].text
    assert "8ص" not in basic.cells[1].text
    assert "آلي" not in " | ".join(cell.text for cell in emergency.cells)
    assert "12ظ" not in " | ".join(cell.text for cell in emergency.cells)


def test_generated_roster_uses_shared_person_labels(client):
    person = client.post("/api/person", json={
        "type": "personnel", "name": "فرد اختبار طويل الاسم", "code": "901",
        "phone": "01012345678", "join_date": "2020-01-01", "role": "أمين شرطة",
    }).get_json()
    client.patch("/api/person/OFF-001", json={"short_name": "مختصر الضابط"})
    for name, payload in (("خدمة فرد", {"personnel_ids": [person["id"]]}),
                          ("خدمة ضابط", {"officer_ids": ["OFF-001"]}),
                          ("خدمة fallback", {"officer_ids": ["OFF-002"]})):
        payload.update({"name": name, "section": "الخدمات الطارئة", "kind": "خارجية"})
        assert client.post(f"/api/assignments/{DAY}", json=payload).status_code == 201
    text = _all_text(_download(client))
    # قوام الطوارئ أعداد («فرد») مش أسماء — الأسماء في يومية الأفراد
    assert "فرد اختبار" not in text
    assert "فرد" in text
    assert "عقيد/ مختصر الضابط" in text
    assert "نقيب/ محمود علي" in text


def test_export_prints_only_the_role_kind_for_required_vacant_commands(client):
    for name, field in (("خدمة ضابط شاغرة", "officer_command_required"),
                        ("خدمة فرد شاغرة", "personnel_command_required")):
        response = client.post(f"/api/assignments/{DAY}", json={
            "name": name, "section": "الخدمات أساسية", "kind": "خارجية",
            "shift": "صباحية", field: True,
        })
        assert response.status_code == 201, response.get_json()
    table = _download(client).tables[0]
    officer = next(row for row in table.rows if row.cells[0].text.strip() == "خدمة ضابط شاغرة صبح")
    personnel = next(row for row in table.rows if row.cells[0].text.strip() == "خدمة فرد شاغرة صبح")
    assert officer.cells[1].text.strip() == "ضابط"
    assert personnel.cells[1].text.strip() == "فرد"
    assert "رئاسة" not in _all_text(_download(client))


def test_export_prints_assigned_command_people_without_role_prefixes(client):
    person = client.post("/api/person", json={
        "type": "personnel", "name": "فرد قيادة اختبار", "code": "902",
        "phone": "01012345679", "join_date": "2020-01-01", "role": "أمين شرطة",
    }).get_json()
    for name, field in (("قيادة ضابط", {"officer_ids": ["OFF-001"]}),
                        ("قيادة فرد", {"personnel_ids": [person["id"]]})):
        payload = {"name": name, "section": "الخدمات أساسية", "kind": "خارجية",
                   "shift": "صباحية", **field}
        response = client.post(f"/api/assignments/{DAY}", json=payload)
        assert response.status_code == 201, response.get_json()
    table = _download(client).tables[0]
    officer = next(row for row in table.rows if row.cells[0].text.strip() == "قيادة ضابط صبح")
    personnel = next(row for row in table.rows if row.cells[0].text.strip() == "قيادة فرد صبح")
    assert officer.cells[1].text.strip() == "عقيد/ أحمد محمد"
    assert personnel.cells[1].text.strip() == "ا.ش/ فرد قيادة اختبار"
    assert "رئاسة ضابط" not in _all_text(_download(client))
    assert "رئاسة فرد" not in _all_text(_download(client))


def test_board_file_drops_clinic_and_search_members_and_web_hides_search_only(client):
    clinic, search = "OFF-001", "OFF-002"
    names = {clinic: "عقيد/ أحمد محمد", search: "نقيب/ محمود علي"}
    for name, oid in (("خدمة العيادة", clinic), ("خدمة البحث", search)):
        r = client.post(f"/api/assignments/{DAY}", json={
            "name": name, "section": "الخدمات الطارئة", "kind": "خارجية", "officer_ids": [oid]})
        assert r.status_code == 201, r.get_json()
    label = names.get
    before = _all_text(_download(client))
    assert label(clinic) in before and label(search) in before
    assert client.patch("/api/command-groups",
                        json={"طبي": [clinic], "بحث": [search]}).status_code == 200
    after = _all_text(_download(client))
    assert label(clinic) not in after and label(search) not in after
    web = str(client.get(f"/api/board/{DAY}").get_json())
    assert clinic in web and search not in web              # الطبي كما هو، والبحث مخفي من الشاشة
