"""فجوات الأرشيف (importer/archive_gaps.py + importer/single_day.py + migrations/023):
استخراج من docx تخليقي لكل فئة، تخطي التعارض، idempotency، وdry-run لا يكتب."""
import hashlib
import sys
from copy import deepcopy
from importlib import import_module

import pytest
from docx import Document

from backend import store
from importer import archive_gaps as gaps
from importer import single_day

DAY = "2026-07-10"
ROSTER = [
    {"id": "OFF-001", "name": "احمد سمير محمد حمدي", "short_name": "احمد سمير", "role": "عقيد"},
    {"id": "OFF-002", "name": "محمد فؤاد وحيد فؤاد", "short_name": "محمد فؤاد", "role": "عقيد"},
    {"id": "OFF-003", "name": "السعيد محمد احمد عبده", "short_name": "السعيد عبده", "role": "مقدم"},
    {"id": "OFF-004", "name": "طاهر سلطان عبد الحليم", "short_name": "طاهر سلطان", "role": "مقدم"},
    {"id": "OFF-005", "name": "محمد مصطفي محمد احمد", "short_name": "", "role": "عقيد"},
    {"id": "OFF-006", "name": "محمد احمد مصطفي عبدالرازق", "short_name": "محمد مصطفي", "role": "مقدم"},
    {"id": "OFF-007", "name": "رامي سامي مطر", "short_name": "رامي سامي", "role": "رائد"},
]


def _officer(raw):
    return {**raw, "code": raw["id"], "phone": "0100", "join_date": "2020-01-01", "post": "", "status": "active",
            "section": "القوة", "rest_system": "—", "rest_day": ""}


def _data(day=DAY, targets=("مشرف الأهداف", "سوميد")):
    return {
        "officers": [_officer(o) for o in ROSTER], "personnel": [], "leaves": [], "change_log": [],
        "command": {"مدير الإدارة": "OFF-001", "وكيل الإدارة": "OFF-002"},
        "day_assignments": {day: [
            {"id": f"AS-000{i}", "section": "الأهداف", "name": name, "kind": "حراسات", "shift": "",
             "officer_ids": ["OFF-003"] if name == "مشرف الأهداف" else [], "personnel_ids": [],
             "conscripts": [], "conscript_count": 0, "note": ""} for i, name in enumerate(targets, 1)]},
        "day_officers": {}, "day_import": {day: {"batch": "old", "sources": []}},
    }


# ───────────── ملفات docx تخليقية ─────────────

def _regularity(root, officers, notes=(), day=DAY, folder_day=None, header=True):
    year, month, dd = day.split("-")
    folder = root / str(int(month)) / str(folder_day or int(dd))
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{int(dd)}..docx"
    document = Document()
    if header:
        document.sections[0].header.paragraphs[0].text = (
            f"يومية انتظام الخدمات عن يوم الجمعة الموافق {int(dd)}/{int(month)}/{year}م")
    services = document.add_table(rows=1, cols=3)
    for cell, text in zip(services.rows[0].cells, ["م", "الخدمة", "قائم للخدمة"]):
        cell.text = text
    table = document.add_table(rows=1, cols=4)
    for cell, text in zip(table.rows[0].cells, ["م", "الرتبة ", "الاسم", "الـخــدمـــــــــــــــة"]):
        cell.text = text
    for rank, name, service in officers:
        row = table.add_row()
        for cell, text in zip(row.cells, ["", rank, name, service]):
            cell.text = text
    for text in notes:
        row = table.add_row()
        row.cells[1].merge(row.cells[3]).text = text
    document.save(path)
    return path


def _board(root, targets=(), admin=(), day=DAY):
    """لوحة بخمس أعمدة: نصف أيمن (0-1) نصفه الأيسر (2-4)؛ الأهداف ثم عمل بالإدارة."""
    year, month, dd = day.split("-")
    folder = root / str(int(month)) / str(int(dd))
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{int(dd)}.docx"
    document = Document()
    table = document.add_table(rows=0, cols=5)

    def heading(text):
        row = table.add_row()
        row.cells[0].merge(row.cells[1]).text = text

    def line(label, value):
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = label, value

    head = table.add_row()
    head.cells[0].merge(head.cells[1]).text = "الخدمات الاساسية"
    head.cells[2].merge(head.cells[4]).text = "الخدمات الطارئة"
    line("تدخل سريع صبح", "رائد / رامي سامي")
    heading("الاهداف")
    for label, value in targets:
        line(label, value)
    heading("عمل بالادارة")
    for label, value in admin:
        line(label, value)
    document.save(path)
    return path


# ───────────── 1) الانتظام ─────────────

OFFICERS = [
    ("عقيد", "احمد سمير", "المتابعة الاشرافية علي جميع خدمات قوات الامن "),
    ("عقيد ", "محمد فؤاد", "المرور التعقبي على جميع خدمات المدينة من 5م : 10م "),
    ("مقدم ", "أحمد هلال السملاوي", 'مشرف القول المكبر فترة صباحية "8 ص الي 8 م"'),
    ("مقدم ", "عبدالرحمن الحسيني( 01090376666  )", 'مشرف القول المكبر فترة ليلية "8 م الي 8 ص"'),
    ("رائد ", "رامي سامي (01017203535)", 'تدخل سريع فترة صباحية "8 ص الي 8 م"'),
]
NOTES = ["عدد 10 مجند حفظ نظام لتأمين محكمه  والباب الرئيسي الساعه 8ص بالمديرية ",
         "تعزيز المجمع الطبي (2مجند) الساعة 2 م وحتي انتهاء زيارات المرضي"]


def test_regularity_extraction_reads_officer_rows_and_free_notes(tmp_path):
    sheet = gaps.extract_regularity(_regularity(tmp_path, OFFICERS, NOTES))
    assert sheet["error"] is None
    assert [o["service"] for o in sheet["officers"]][:2] == [
        "المتابعة الاشرافية علي جميع خدمات قوات الامن", "المرور التعقبي على جميع خدمات المدينة من 5م : 10م"]
    phone_row = sheet["officers"][3]
    assert (phone_row["name"], phone_row["phone"]) == ("عبدالرحمن الحسيني", "01090376666")
    assert sheet["notes"][0] == "عدد 10 مجند حفظ نظام لتأمين محكمه والباب الرئيسي الساعه 8ص بالمديرية"
    assert len(sheet["notes"]) == 2


def test_regularity_day_comes_from_header_not_folder_and_duplicates_are_reported(tmp_path):
    right = _regularity(tmp_path, OFFICERS, day="2026-09-08")
    wrong = _regularity(tmp_path, OFFICERS[:2], day="2026-09-08", folder_day=9)      # نسخة مزاحة في فولدر 9/9
    shifted = _regularity(tmp_path, OFFICERS[:1], day="2026-10-21", folder_day=22)   # 10/22/21..docx = يوم 21
    chosen, notes = gaps.discover_regularity(tmp_path, "2026-08-01")
    assert chosen == {"2026-09-08": right, "2026-10-21": shifted}
    assert [(n["reason"], n["path"]) for n in notes] == [("duplicate", str(wrong))]
    assert gaps.header_day(shifted) == "2026-10-21"


def test_regularity_plan_maps_command_extra_big_column_and_notes(tmp_path):
    data = _data()
    sheet = gaps.extract_regularity(_regularity(tmp_path, OFFICERS, NOTES))
    plan = gaps.plan_regularity(data, DAY, sheet, apply=True)
    assert [op.status for op in plan.ops] == ["add"] * 4
    assert data["regularity_inputs"][DAY] == {
        "officer_extra": {"OFF-002": "المرور التعقبي على جميع خدمات المدينة من 5م : 10م"},
        "big_column_supervisors": {
            "am": {"rank": "مقدم", "name": "أحمد هلال السملاوي", "phone": ""},
            "pm": {"rank": "مقدم", "name": "عبدالرحمن الحسيني", "phone": "01090376666"}},
        "notes": [NOTES[0].replace("  ", " ").strip(), NOTES[1]]}
    assert plan.skipped_basic == 1 and plan.unresolved == []       # صف التدخل السريع مشتق من اللوحة
    # الشكل المخزّن هو بالظبط اللي regularity.effective_inputs/PUT بيقبله
    from backend import regularity
    effective, source = regularity.effective_inputs(data, DAY)
    assert source == DAY and effective["officer_extra"]["OFF-002"].startswith("المرور التعقبي")


def test_supervision_extra_goes_after_the_plus_and_unknown_officers_are_unresolved(tmp_path):
    data = _data()
    rows = [("عقيد", "احمد سمير", "المتابعة الاشرافية الميدانية على جميع خدمات ادارة قوات الأمن + المرور التعقبي 5م حتي 10م"),
            ("نقيب", "فلان الغريب", "المتابعة الاشرافية علي جميع خدمات قوات الامن"),
            ("نقيب", "علان الغريب", "اشراف على نقطة غامضة")]
    plan = gaps.plan_regularity(data, DAY, gaps.extract_regularity(_regularity(tmp_path, rows)), apply=True)
    assert data["regularity_inputs"][DAY]["officer_extra"] == {"OFF-001": "المرور التعقبي 5م حتي 10م"}
    assert [u["reason"] for u in plan.unresolved] == [
        "المتابعة الإشرافية لضابط مش مدير/وكيل اليوم",
        "صف ضابط مش مدير/وكيل ولا مشرف قول مكبر ولا خدمة أساسية"]


def test_regularity_conflict_is_reported_and_stored_value_survives(tmp_path):
    data = _data()
    data["regularity_inputs"] = {DAY: {"officer_extra": {"OFF-002": "نص دخله المستخدم"}, "notes": ["ملاحظة المستخدم"],
                                         "big_column_supervisors": {}}}
    sheet = gaps.extract_regularity(_regularity(tmp_path, OFFICERS, NOTES))
    plan = gaps.plan_regularity(data, DAY, sheet, apply=True)
    assert sorted(op.key for op in plan.ops if op.status == "conflict") == ["notes", "officer_extra:OFF-002"]
    assert data["regularity_inputs"][DAY]["officer_extra"] == {"OFF-002": "نص دخله المستخدم"}
    assert data["regularity_inputs"][DAY]["notes"] == ["ملاحظة المستخدم"]
    assert set(data["regularity_inputs"][DAY]["big_column_supervisors"]) == {"am", "pm"}   # الفاضي اتملى


def test_user_reset_of_regularity_inputs_blocks_the_import(tmp_path):
    data = _data()
    data["change_log"].append({"id": "CHG-1", "ts": "2026-10-25T10:00:00", "entity": "regularity_inputs",
                               "entity_id": DAY, "action": "reset", "day": DAY, "text": "", "edited_by": "u"})
    plan = gaps.plan_regularity(data, DAY, gaps.extract_regularity(_regularity(tmp_path, OFFICERS, NOTES)), apply=True)
    assert {op.status for op in plan.ops} == {"conflict"}
    assert "regularity_inputs" not in data or DAY not in data["regularity_inputs"]


def test_regularity_limits_are_unresolved_not_truncated(tmp_path):
    data = _data()
    too_many = [f"ملاحظة رقم {i}" for i in range(13)]
    plan = gaps.plan_regularity(data, DAY, gaps.extract_regularity(_regularity(tmp_path, [], too_many)), apply=True)
    assert [u["reason"] for u in plan.unresolved] == ["13 ملاحظة (الحد 12)"]
    assert "regularity_inputs" not in data


# ───────────── 2) عمل بالإدارة ─────────────

ADMIN = [("عقيد / احمد سمير", "مدير الادارة"),
         ("مقدم / طاهر سلطان", "متابعة أعمال الشئون   + متابعة التعيينات حتى 4م"),
         ("مقدم / السعيد عبده", "ضابط امن الادارة + \"ت\""),
         ("رائد / ناس مجهولين", "متابعة غامضة")]


def test_admin_work_extraction_and_plan(tmp_path):
    data = _data()
    data["day_officers"][DAY] = {"OFF-004": {"note": "متابعة أعمال الشئون + متابعة التعيينات حتى 4م"},
                                 "OFF-003": {"taqseera": True}}
    path = _board(tmp_path, admin=ADMIN)
    sections = gaps.extract_board_sections(path)
    assert [row["service"] for row in sections["admin_work"]][:2] == ["عقيد / احمد سمير", "مقدم / طاهر سلطان"]
    plan = gaps.plan_admin_work(data, DAY, sections, str(path), apply=True)
    states = data["day_officers"][DAY]
    assert states["OFF-004"] == {"admin_work_text": "متابعة أعمال الشئون + متابعة التعيينات حتى 4م"}   # note اتنقلت
    assert states["OFF-003"] == {"taqseera": True, "admin_work_text": 'ضابط امن الادارة + "ت"'}      # التقصيرة محفوظة
    assert "OFF-001" not in states and "admin_work" not in states["OFF-004"]                       # المدير: نص المنصب تلقائي
    assert [u["row"] for u in plan.unresolved] == ["رائد / ناس مجهولين"]
    assert [op.detail for op in plan.ops if op.key == "OFF-004"] == ["نقل من note (نفس النص)"]


def test_admin_work_never_overwrites_a_different_stored_text_and_director_text_replaces_label(tmp_path):
    data = _data()
    data["day_officers"][DAY] = {"OFF-004": {"admin_work_text": "نص المستخدم", "note": "حاجة تانية"}}
    admin = [("عقيد / احمد سمير", "مدير الادارة + المرور التعقبي 5م حتى 10م"),
             ("مقدم / طاهر سلطان", "نص الأرشيف")]
    path = _board(tmp_path, admin=admin)
    plan = gaps.plan_admin_work(data, DAY, gaps.extract_board_sections(path), str(path), apply=True)
    assert data["day_officers"][DAY]["OFF-004"] == {"admin_work_text": "نص المستخدم", "note": "حاجة تانية"}
    assert data["day_officers"][DAY]["OFF-001"] == {"admin_work_text": "مدير الادارة + المرور التعقبي 5م حتى 10م"}
    assert [op.status for op in plan.ops if op.key == "OFF-004"] == ["conflict"]


def test_admin_work_text_deliberately_cleared_by_user_is_a_conflict(tmp_path):
    data = _data()
    data["change_log"].append({"id": "CHG-2", "ts": "2026-10-25T10:00:00", "entity": "officer_state",
                               "entity_id": "OFF-004", "action": "update", "day": None, "reason": f"يوم {DAY}",
                               "text": f"في يومية {DAY}", "before": {"admin_work_text": "قديم"},
                               "after": {"admin_work_text": ""}, "edited_by": "u"})
    path = _board(tmp_path, admin=[("مقدم / طاهر سلطان", "نص الأرشيف")])
    plan = gaps.plan_admin_work(data, DAY, gaps.extract_board_sections(path), str(path), apply=True)
    assert [op.status for op in plan.ops] == ["conflict"] and not data["day_officers"]


def test_admin_text_longer_than_limit_is_unresolved(tmp_path):
    data = _data()
    path = _board(tmp_path, admin=[("مقدم / طاهر سلطان", "ا" * 501)])
    plan = gaps.plan_admin_work(data, DAY, gaps.extract_board_sections(path), str(path), apply=True)
    assert plan.ops == [] and "أطول من 500" in plan.unresolved[0]["reason"]


def test_ambiguous_names_are_broken_by_rank_never_guessed(tmp_path):
    data = _data()
    index = gaps.OfficerIndex(data, DAY)
    assert index.resolve("محمد مصطفى", "")["id"] == "OFF-006"               # الاسم المختصر المسجّل بيكسب
    data["officers"].append(_officer({"id": "OFF-008", "name": "محمد مصطفي علي حسن", "short_name": "محمد مصطفي",
                                      "role": "عقيد"}))
    index = gaps.OfficerIndex(data, DAY)
    assert index.resolve("محمد مصطفى", "مقدم")["id"] == "OFF-006"          # نفس الاسم: الرتبة بتفك
    assert index.resolve("محمد مصطفى", "")["status"] == "ambiguous"        # من غير رتبة مانخمّنش
    assert index.resolve("محمد", "")["status"] == "none"


# ───────────── 3) تقسيم الهدف ─────────────

SPLIT_ROWS = [("مشرف الأهداف", "مقدم / السعيد عبده حتى 4م + ت"), ("", "مقدم / طاهر سلطان من 4م"),
              ("سوميد", "رائد / رامي سامي")]


def test_split_target_adds_periods_and_officer_ids(tmp_path):
    data = _data()
    path = _board(tmp_path, targets=SPLIT_ROWS)
    sections = gaps.extract_board_sections(path)
    assert [(r["service"], r["value"]) for r in sections["targets"]][:2] == [
        ("مشرف الأهداف", "مقدم / السعيد عبده حتى 4م + ت"), ("مشرف الأهداف", "مقدم / طاهر سلطان من 4م")]
    plan = gaps.plan_split_targets(data, DAY, sections, str(path), apply=True)
    row = data["day_assignments"][DAY][0]
    assert row["periods"] == [{"officer_id": "OFF-003", "until": "16:00"}, {"officer_id": "OFF-004", "from": "16:00"}]
    assert row["officer_ids"] == ["OFF-003", "OFF-004"]
    from backend.assignments import valid_periods
    assert valid_periods(row) == row["periods"]                         # نفس اللي apply_assignment بيقبله
    assert [op.status for op in plan.ops] == ["add"] and plan.unresolved == []
    assert "periods" not in data["day_assignments"][DAY][1]              # سوميد ضابط واحد: ولا تقسيم


def test_split_target_conflicts_when_another_officer_is_stored(tmp_path):
    data = _data()
    data["day_assignments"][DAY][0]["officer_ids"] = ["OFF-007"]
    path = _board(tmp_path, targets=SPLIT_ROWS)
    plan = gaps.plan_split_targets(data, DAY, gaps.extract_board_sections(path), str(path), apply=True)
    assert [op.status for op in plan.ops] == ["conflict"]
    assert data["day_assignments"][DAY][0]["officer_ids"] == ["OFF-007"] and "periods" not in data["day_assignments"][DAY][0]


def test_split_target_requires_matching_until_and_from_times(tmp_path):
    data = _data()
    rows = [("مشرف الأهداف", "مقدم / السعيد عبده حتى 4م"), ("", "مقدم / طاهر سلطان من 5م")]
    path = _board(tmp_path, targets=rows)
    plan = gaps.plan_split_targets(data, DAY, gaps.extract_board_sections(path), str(path), apply=True)
    assert plan.ops == [] and plan.unresolved[0]["reason"].startswith("الفترتان مش")
    assert data["day_assignments"][DAY][0]["officer_ids"] == ["OFF-003"]


def test_split_target_same_values_are_idempotent(tmp_path):
    data = _data()
    path = _board(tmp_path, targets=SPLIT_ROWS)
    sections = gaps.extract_board_sections(path)
    gaps.plan_split_targets(data, DAY, sections, str(path), apply=True)
    snapshot = deepcopy(data)
    again = gaps.plan_split_targets(data, DAY, sections, str(path), apply=True)
    assert [op.status for op in again.ops] == ["same"] and data == snapshot


# ───────────── اليوم الكامل ─────────────

def test_day_files_ignore_files_dated_for_another_day(tmp_path):
    root = tmp_path / "2026"
    _board(root, day="2026-10-22")
    _regularity(root, OFFICERS[:1], day="2026-10-21", folder_day=22)
    folder = root / "10" / "22"
    for name in ("21-10-2026 افراد.docx", "22-10-2026.docx", "22-10-2026 افراد.docx", "يومية الحملة عن يوم 22-10-2026.docx"):
        Document().save(folder / name)
    names = sorted(p.name for p in single_day.day_files(root, "2026-10-22"))
    assert names == ["22-10-2026 افراد.docx", "22-10-2026.docx", "22.docx"]


def test_missing_day_detection_and_skips(tmp_path):
    root = tmp_path / "2026"
    _board(root, day="2026-10-22")
    data = _data()
    assert single_day.missing_days(data, root, "2026-08-01", None) == ["2026-10-22"]
    assert single_day.missing_days(data, root, "2026-10-23", None) == []
    result = single_day.import_day(data, DAY, root, "b")
    assert result["status"] == "skip" and "مخزّنة" in result["reason"]
    assert single_day.import_day(data, "2026-10-23", root, "b")["reason"].startswith("مفيش لوحة")
    assert "2026-10-22" not in data["day_assignments"]                  # الـskip ما لمسش حاجة


def test_strength_only_label_is_moved_back_to_emergency_and_reported():
    records = [{"record_type": "header", "half_labels": ["الخدمات الاساسية", "الخدمات الطارئة"]},
               {"record_type": "board_row", "half": 1, "label": "حفل 5م", "manning": "مقدم / طاهر سلطان"},
               {"record_type": "board_label", "half": 1, "label": "فرد + وحدة", "manning": ""},
               {"record_type": "board_row", "half": 1, "label": "تبة ضرب النار", "manning": "مقدم / محمد مصطفى"}]
    sections = single_day.strength_label_sections(records)
    assert sections == {"فرد + وحدة": "emergency"}
    blob = {"assignments": [{"section": "فرد + وحدة", "name": "تبة ضرب النار"},
                            {"section": "الأهداف", "name": "سوميد"}]}
    notes = single_day.repair_strength_sections(blob, sections)
    assert blob["assignments"][0]["section"] == "الخدمات الطارئة" and "مانسبتش القوة" in notes[0]
    with pytest.raises(ValueError):
        single_day.repair_strength_sections({"assignments": [{"section": "فرد + وحدة", "name": "س"}]},
                                            {"فرد + وحدة": "targets"})


# ───────────── الهجرة 023: dry-run / --write / idempotency ─────────────

def _tree_digest():
    digest = hashlib.sha1()
    for path in sorted(store.DATA_DIR.rglob("*.json")):
        digest.update(path.name.encode() + path.read_bytes())
    return digest.hexdigest()


def test_migration_dry_run_writes_nothing_then_write_is_idempotent(data_file, monkeypatch, capsys, tmp_path):
    migration = import_module("migrations.023_import_archive_gaps")
    root = tmp_path / "2026"
    _regularity(root, OFFICERS, NOTES)
    _board(root, targets=SPLIT_ROWS, admin=ADMIN[1:3])
    data = store.load_data(store.ALL_DAYS)
    seeded = _data()
    for key in ("officers", "command", "day_assignments", "day_officers", "day_import"):
        data[key] = seeded[key]
    store.save_data(data)
    report_path = tmp_path / "preview.md"
    base = ["023", "--archive-root", str(root), "--from", "2026-07-01", "--no-day-import"]
    run = lambda: migration.run(from_schema=6, to_schema=6, migrate=migration.migrate, title="t")   # noqa: E731
    before = _tree_digest()

    monkeypatch.setattr(sys, "argv", base + ["--report", str(report_path)])
    run()
    assert _tree_digest() == before                                       # dry-run ما كتبش حاجة
    assert "عرض بس" in capsys.readouterr().out
    text = report_path.read_text(encoding="utf-8")
    assert DAY in text and "مدخلات الانتظام" in text and "تقسيم الهدف بالوقت" in text

    monkeypatch.setattr(sys, "argv", base + ["--write"])
    run()
    written = _tree_digest()
    assert written != before
    data = store.load_data(store.ALL_DAYS)
    assert data["regularity_inputs"][DAY]["officer_extra"]["OFF-002"].startswith("المرور التعقبي")
    assert data["day_officers"][DAY]["OFF-004"]["admin_work_text"].startswith("متابعة أعمال الشئون")
    assert data["day_assignments"][DAY][0]["officer_ids"] == ["OFF-003", "OFF-004"]
    entry = next(c for c in data["change_log"] if c["entity"] == "import" and c["entity_id"] == "archive-gaps-023")
    assert entry["after"]["fields_added"] == {"regularity": 4, "admin_work": 2, "split_target": 1}

    capsys.readouterr()
    run()                                                                 # --write تاني
    assert _tree_digest() == written                                      # مفيش تغيير ولا إعادة كتابة
    assert "متملية بالفعل" in capsys.readouterr().out
