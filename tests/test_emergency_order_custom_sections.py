"""الطوارئ مترتبة بمعاد الانتظام، وأقسام المشغّل (ماتش…) ظاهرة في التفصيلية وتصديرها."""
import io

from docx import Document
from docx.table import _Cell

from backend.afraad import _time_key
from backend.board_export import LAST_FIT, build_docx

DAY = "2026-10-21"


def test_time_key_is_chronological_with_minutes_noon_and_night():
    times = ["12ليل", "2:30م", "8م", "2م", "12ظ", "9ص", "", "8ص", "7ص", "11 ص حتى الانتهاء", "—"]
    assert sorted(times, key=_time_key) == [
        "7ص", "8ص", "9ص", "11 ص حتى الانتهاء", "12ظ", "2م", "2:30م", "8م", "12ليل", "", "—"]
    assert _time_key("8 لاسلكي") == _time_key("8")    # «ل» جوه كلمة مش ليل


def _service(name, time, section="الخدمات الطارئة"):
    return {"id": f"S-{name}", "label": name, "name": name, "time": time, "section": section,
            "officers": [], "personnel": [], "conscripts": [], "conscript_count": 1}


def _cells(document):
    table = document.tables[0]
    return table, ["|".join(_Cell(tc, table).text for tc in row._tr.tc_lst) for row in table.rows]


def test_afraad_page_orders_emergency_by_schedule(client):
    for name, time in (("ثالثة", "2:30م"), ("أولى", "8ص"), ("تانية", "12ظ"), ("بدون", "")):
        assert client.post(f"/api/assignments/{DAY}", json={
            "name": name, "kind": "خارجية", "section": "الخدمات الطارئة", "time": time}).status_code == 201
    rows = client.get(f"/api/afraad/{DAY}").get_json()["occasional"]
    assert [row["name"] for row in rows] == ["أولى", "تانية", "ثالثة", "بدون"]


def test_board_export_separates_special_block_and_keeps_grid_coverage():
    emergency = [_service(f"طارئة {i}", f"{(i % 9) + 1}م") for i in range(12)]
    custom = [_service("تأمين مباراة", "2:30م", "مباراة ش"), _service("الباب الرئيسي", "8ص", "مباراة ش")]
    board = {"date": DAY, "sections": [
        {"name": "الخدمات الطارئة", "type": "services", "rows": emergency},
        {"name": "مباراة ش", "type": "services", "rows": custom}]}
    document = Document(io.BytesIO(build_docx(board).getvalue()))
    table, lines = _cells(document)
    text = "\n".join(lines)
    for i in range(12):
        assert f"طارئة {i}" in text
    assert "تأمين مباراة" in text and "الباب الرئيسي" in text
    assert text.index("خدمات مباراة ش") < text.index("الباب الرئيسي") < text.index("تأمين مباراة")
    grid = [int(c.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"))
            for c in table._tbl.tblGrid]
    for tr in table._tbl.tr_lst:
        spans = []
        for tc in tr.tc_lst:
            node = tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan")
            spans.append(int(node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")) if node is not None else 1)
        assert sum(spans) == 5
    assert any("خدمات مباراة ش" in "".join(t.text or "" for t in tc.iter())
               and tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan").get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "3"
               for tr in table._tbl.tr_lst for tc in tr.tc_lst
               if tc.tcPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan") is not None)
    assert LAST_FIT["ratio"] <= 1
    # لا صفوف حشو: عنوان + 11 صف طوارئ/مخصص + «الاهداف» + عنوان الخوارج + صفّا المعسكر الفرعي
    assert len(table.rows) >= 15


def test_prison_rows_beyond_seven_are_kept():
    prison = [_service(f"سجن {i}", "", "خدمات سجن قوات الأمن") for i in range(9)]
    document = Document(io.BytesIO(build_docx({"date": DAY, "sections": [
        {"name": "خدمات سجن قوات الأمن", "type": "services", "rows": prison}]}).getvalue()))
    text = "\n".join(_cells(document)[1])
    for i in range(9):
        assert f"سجن {i}" in text


def test_ordinary_emergency_rows_stay_above_special_blocks_with_targets_row():
    # مع وجود كتلة أهداف، صف الطوارئ العادي مايتنقلش لصف «الاهداف» تحت كتلة المباراة
    from backend.board_export import _latest_content
    emergency = [_service(f"طارئة {i}", f"{i + 1}م") for i in range(4)]
    custom = [_service("مدرجات", "2م", "مباراة ش")]
    board = {"date": DAY, "sections": [
        {"name": "الخدمات الطارئة", "type": "services", "rows": emergency},
        {"name": "مباراة ش", "type": "services", "rows": custom},
        {"name": "الأهداف", "type": "targets", "rows": [{"name": "سوميد", "officers": []}]}]}
    texts = [" ".join(c["text"] for c in row["cells"]) for row in _latest_content(board)]
    title = next(i for i, t in enumerate(texts) if "خدمات مباراة ش" in t)
    assert all(any(f"طارئة {i}" in t for t in texts[:title]) for i in range(4))
    targets_row = next(t for t in texts if t.startswith("الاهداف"))
    assert "طارئة" not in targets_row
