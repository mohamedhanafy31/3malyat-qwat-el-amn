#!/usr/bin/env python3
"""مقارنة يومية انتظام مولدة مع نسخة الأرشيف، مع فصل فجوات البيانات عن عيوب الكود."""
import argparse
import os
import re
import tempfile
from datetime import date
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from zipfile import ZipFile
from lxml import etree

from backend.regularity import build
from backend.regularity_export import build_docx
from backend.store import load_data
from backend.text import core_service_name


def _day(value):
    if re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", value):
        return date.fromisoformat(value).isoformat()
    parts = value.split("/")
    if len(parts) == 3:
        y, m, d = map(int, parts)
        return date(y, m, d).isoformat()
    raise SystemExit("اكتب اليوم بصيغة YYYY-MM-DD أو YYYY/M/D")


def _archive(root, day):
    d = date.fromisoformat(day)
    wanted = f"{d.day}..docx"
    root = Path(root)
    base = root / str(d.year) if root.name != str(d.year) else root
    direct = base / str(d.month) / str(d.day) / wanted
    if direct.exists():
        return direct
    # بعض الأيام محفوظة في مجلد لاحق (مثل 21 أكتوبر داخل 22/21..docx).
    matches = sorted(base.glob(f"{d.month}/*/{wanted}"))
    return matches[0] if matches else None


def _norm(value):
    value = str(value or "")
    value = re.sub(r"\s+", "", value).replace("ـ", "")
    return value.translate(str.maketrans({"ى": "ي", "ة": "ه", "أ": "ا", "إ": "ا", "آ": "ا"}))


def _archive_values(path):
    d = Document(path)
    header = " ".join(p.text for s in d.sections for p in s.header.paragraphs).strip()
    table1 = [[c.text for c in row.cells] for row in d.tables[0].rows]
    table2 = [[c.text for c in row.cells] for row in d.tables[1].rows]
    notes = [row[1] for row in table2[1:] if len(set(_norm(c) for c in row[1:])) == 1]
    return header, table1, table2, notes


def _boxes(path):
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    return [" ".join("".join(node.itertext()).split()) for node in root.xpath(".//w:txbxContent", namespaces=ns)]


def _generated_values(payload):
    payload = dict(payload)
    payload["day"] = payload.get("day") or payload.get("date")
    buf = build_docx(payload)
    with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
        handle.write(buf.read()); handle.flush()
        d = Document(handle.name)
        header = " ".join(p.text for s in d.sections for p in s.header.paragraphs).strip()
        table1 = [[c.text for c in row.cells] for row in d.tables[0].rows]
        table2 = [[c.text for c in row.cells] for row in d.tables[1].rows]
        boxes = _boxes(handle.name)
    return header, table1, table2, boxes


def compare_day(day, archive_root="../../2026", data=None):
    """يرجع نتيجة مقارنة قابلة للاختبار، وتستخدمها واجهة CLI نفسها."""
    day = _day(day) if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(day)) else str(day)
    archive = _archive(archive_root, day)
    if not archive:
        raise FileNotFoundError(f"لم أجد نسخة أرشيف لليوم {day}")
    data = data if data is not None else load_data([day])
    payload = build(data, day)
    payload["day"] = day
    generated = _generated_values(payload)
    archived = _archive_values(archive)
    generated_rows = generated[1][1:]
    archive_rows = archived[1][1:]
    generated_by_core = {}
    archive_by_core = {}
    for row in generated_rows:
        if len(row) > 6:
            generated_by_core.setdefault(_norm(core_service_name(row[1])), []).append(row)
    for row in archive_rows:
        if len(row) > 6:
            archive_by_core.setdefault(_norm(core_service_name(row[1])), []).append(row)
    generated_service_names = sorted(generated_by_core)
    archive_service_names = sorted(archive_by_core)
    shared = sorted(set(generated_by_core) & set(archive_by_core))
    row_checks = []
    for core in shared:
        generated_keys = {}
        for row in generated_by_core[core]:
            generated_keys.setdefault(_norm(row[6]), []).append(row)
        for archive_row in archive_by_core[core]:
            key = _norm(archive_row[6])
            candidates = generated_keys.get(key, [])
            row_checks.append({
                "service": core, "time": key,
                "generated": candidates,
                "archive": archive_row,
                "fields": {
                    field: any(_norm(candidate[index]) == _norm(archive_row[index])
                               for candidate in candidates)
                    for field, index in (("time", 6), ("count", 7),
                                         ("place", 5), ("people", 2))
                },
            })
    return {"day": day, "archive": archive, "payload": payload,
            "generated": generated, "archived": archived,
            "generated_service_names": generated_service_names,
            "archive_service_names": archive_service_names,
            "shared_service_names": shared, "row_checks": row_checks}


def _field(label, left, right):
    ok = _norm(left) == _norm(right)
    print(f"{label:<18} {'MATCH' if ok else 'MISMATCH'}")
    if not ok:
        print(f"  generated: {left!r}\n  archive:   {right!r}\n  class: data gap (archive is historical) unless the generated rule is wrong")
    return ok


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("day")
    parser.add_argument("--archive-root", default="../../2026")
    args = parser.parse_args()
    day = _day(args.day)
    archive = _archive(args.archive_root, day)
    if not archive:
        raise SystemExit(f"لم أجد نسخة أرشيف لليوم {day}")
    result = compare_day(day, args.archive_root)
    payload = result["payload"]
    gh, gt1, gt2, generated_boxes = result["generated"]
    ah, at1, at2, notes = result["archived"]
    print(f"day={day} archive={archive}")
    counts = {"matched": 0, "mismatched": 0}
    def check(label, left, right):
        if _field(label, left, right): counts["matched"] += 1
        else: counts["mismatched"] += 1
    check("header", gh, ah)
    for i, name in enumerate(("م", "الخدمة", "قائم للخدمة", "رقم الهاتف", "المركبة", "مكان الانتظام", "ساعة الانتظام", "عدد المجندين")):
        if i < len(gt1[0]) and i < len(at1[0]): check(f"table1.{name}", gt1[0][i], at1[0][i])
    generated_services = {_norm(row[1]) for row in gt1[1:] if len(row) > 1}
    archive_services = {_norm(row[1]) for row in at1[1:] if len(row) > 1}
    check("service names set", sorted(generated_services), sorted(archive_services))
    # الصفوف الرسمية تُطابق بالخدمة + الساعة، لا بترتيب الصف، لأن المركبات
    # أو صفوف الشخص الواحد قد تقسم التكليف إلى أكثر من صف.
    by_key = {}
    for row in gt1[1:]:
        if len(row) > 6:
            by_key.setdefault((_norm(row[1]), _norm(row[6])), []).append(row)
    for field, idx in (("people", 2), ("phone", 3), ("vehicle", 4),
                       ("place", 5), ("counts", 7)):
        matched = mismatched = missing = 0
        for archive_row in at1[1:]:
            if len(archive_row) <= idx:
                continue
            candidates = by_key.get((_norm(archive_row[1]), _norm(archive_row[6])), [])
            if not candidates:
                missing += 1
                continue
            if any(_norm(candidate[idx]) == _norm(archive_row[idx]) for candidate in candidates):
                matched += 1
            else:
                mismatched += 1
        print(f"table1.{field:<10} MATCH={matched} MISMATCH={mismatched} MISSING={missing}")
        counts["matched"] += matched
        counts["mismatched"] += mismatched + missing
    check("table2.headers", gt2[0], at2[0])
    check("notes", [_norm(row[1]) for row in gt2[-3:] if len(row) > 1],
          [_norm(x) for x in notes[-3:]])
    # التوقيعات داخل text boxes؛ يكفي هنا عدّ ظهور كل اسم في النسختين، لأن
    # Choice وFallback يحملان نفس النص بصريًا.
    archive_boxes = _boxes(archive)
    for key, item in ("director", payload.get("signatories", {}).get("director")), ("deputy", payload.get("signatories", {}).get("deputy")), ("operations", payload.get("signatories", {}).get("operations")):
        name = _norm((item or {}).get("name"))
        count = sum(name and name in _norm(box) for box in archive_boxes)
        generated_count = sum(name and name in _norm(box) for box in generated_boxes)
        print(f"signature.{key:<10} generated={generated_count} archive={count} "
              f"{'MATCH' if generated_count == count else 'MISMATCH (data gap)'}")
    print(f"SUMMARY matched={counts['matched']} mismatched={counts['mismatched']}")


if __name__ == "__main__":
    main()
