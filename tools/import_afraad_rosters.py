#!/usr/bin/env python3
"""Fill «يومية الأفراد» (الخدمات الأساسية فقط) from the daily «افراد» .docx files.

For every roster file from --start (default 2026-08-01) up to the latest one
found, the basic-services table is read (everything between the header row and
the «الخدمات الطارئة» row — the emergency part is never captured) and written to
`day_afraad[day]` with: القائم بها صباحًا/ليلًا (+ التليفون وربطه بالفرد إن
تعرّفنا عليه)، قوام الخدمة، التسليح، الانتظام.

Dry-run by default; --write replaces the afraad section of the captured days
only (every other section of those days is left alone) after a snapshot and a
pre-image copy of the day files.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import docx  # noqa: E402

from backend import store  # noqa: E402
from backend.afraad_ocr import _person_match, _service_score  # noqa: E402
from backend.dated import afraad_basic_on  # noqa: E402
from backend.repo import PeopleRepo  # noqa: E402
from backend.text import norm  # noqa: E402

DEFAULT_ROOT = ROOT.parent.parent / "2026"
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
TITLE_DATE = re.compile(r"(\d{1,2})\s*[/\-]\s*(\d{1,2})\s*[/\-]\s*(20\d\d)")
FILE_DATE = re.compile(r"^(\d{1,2})-(\d{1,2})-(20\d\d)\s+افراد\.docx$")
PLACEHOLDER = re.compile(r"^[\sـ\-–—_.]*$")
PHONE = re.compile(r"(?:\+?20)?0?1[0125][\d\s\-]{7,12}")


def clean(text: str) -> str:
    return " ".join(str(text or "").replace("‏", " ").replace("‎", " ")
                    .replace('"', " ").split())


def field(text: str) -> str:
    text = clean(text)
    return "" if PLACEHOLDER.match(text) else text


def iso(d: int, m: int, y: int) -> str:
    return f"{y:04d}-{m:02d}-{d:02d}"


# ───────────────────────── discovery ─────────────────────────
def roster_day(path: Path, document) -> tuple[str, str]:
    """(day, source) — the date printed in the sheet title wins over the filename,
    because a few files were copied forward without being renamed."""
    for paragraph in document.paragraphs:
        if "يوم" in paragraph.text:
            hit = TITLE_DATE.search(paragraph.text.translate(DIGITS))
            if hit:
                return iso(int(hit[1]), int(hit[2]), int(hit[3])), "title"
    hit = FILE_DATE.match(path.name)
    return (iso(int(hit[1]), int(hit[2]), int(hit[3])), "filename") if hit else ("", "")


def discover(root: Path, start: str, end: str | None):
    """-> ({day: path}, notes). Duplicates for one day: prefer the file whose
    folder and filename agree with the day, then the most recently saved."""
    candidates: dict[str, list[tuple[tuple, Path]]] = {}
    notes = []
    for path in sorted(root.rglob("*افراد.docx")):
        if path.name.startswith("~$") or not FILE_DATE.match(path.name):
            continue
        try:
            document = docx.Document(path)
        except Exception as exc:  # corrupt / locked file
            notes.append(f"تعذر فتح {path}: {exc}")
            continue
        day, source = roster_day(path, document)
        if not day or day < start or (end and day > end):
            continue
        name_day = FILE_DATE.match(path.name)
        agrees = iso(int(name_day[1]), int(name_day[2]), int(name_day[3])) == day
        folder_ok = path.parent.name == str(int(day[8:]))
        candidates.setdefault(day, []).append(((agrees, folder_ok, path.stat().st_mtime), path))
    chosen = {}
    for day, items in sorted(candidates.items()):
        items.sort(key=lambda item: item[0], reverse=True)
        chosen[day] = items[0][1]
        if len(items) > 1:
            notes.append(f"{day}: أكثر من ملف — اخترت {items[0][1]} وتجاهلت "
                         + "، ".join(str(p) for _k, p in items[1:]))
    return chosen, notes


# ───────────────────────── parsing ─────────────────────────
def unique_cells(row):
    seen, cells = set(), []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell.text)
    return cells


def basic_rows(document):
    """Rows of the basic table between the header and «الخدمات الطارئة»."""
    for table in document.tables:
        header = next((i for i, row in enumerate(table.rows)
                       if "الخدمه الصباحيه" in norm(" ".join(unique_cells(row)))
                       and "الخدمه الليليه" in norm(" ".join(unique_cells(row)))), None)
        if header is None:
            continue
        rows = []
        for row in table.rows[header + 1:]:
            cells = unique_cells(row)
            if "الخدمات الطاريه" in norm(" ".join(cells)):
                break
            if len([c for c in cells if c.strip()]) < 3:
                continue
            # القوام/التسليح/الانتظام/الليلية/الصباحية دائمًا آخر خمس خلايا،
            # حتى لو ظهرت خلية إضافية («بيك اب»/«مدرعة») بعد اسم الخدمة.
            rows.append({"service": clean(cells[1] if len(cells) > 1 else ""),
                         "morning": cells[-5], "night": cells[-4],
                         "count": cells[-3], "weapon": cells[-2], "schedule": cells[-1]})
        return rows
    return []


def person_slot(text: str, personnel):
    raw = "\n".join(line for line in str(text or "").translate(DIGITS).splitlines())
    if PLACEHOLDER.match(clean(raw)):
        return {"name": "", "phone": "", "person_id": ""}, "empty", None
    phone = re.sub(r"\D", "", (PHONE.search(raw) or [""])[0])
    by_phone = [p for p in personnel if phone and len(phone) >= 10
                and re.sub(r"\D", "", p.get("phone") or "").endswith(phone[-10:])]
    # رقم التليفون المطبوع هو أقوى دليل؛ الاسم وحده يقع في تعادل مع من يشبهه.
    match = ({"matched": True, "id": by_phone[0]["id"]} if len(by_phone) == 1
             else _person_match(raw, personnel))
    name = clean(PHONE.sub(" ", raw))
    if match.get("matched"):
        person = next(p for p in personnel if p["id"] == match["id"])
        label = "/ ".join(part for part in (person.get("role"), person.get("name")) if part)
        return {"name": label, "phone": person.get("phone") or phone,
                "person_id": person["id"]}, "linked", match
    return {"name": name, "phone": phone, "person_id": ""}, "unlinked", match


def align(rows, services):
    """Row -> service. The sheet has a fixed section, so position is trusted when
    the counts agree (name similarity is only a sanity check); otherwise fall
    back to unique best name matches."""
    result, warnings = {}, []
    if len(rows) == len(services):
        for row, service in zip(rows, services):
            score = _service_score(row["service"], service["name"])
            if score < 0.5:
                warnings.append(f"«{row['service']}» ≠ «{service['name']}» ({score:.2f})")
            result[service["id"]] = row
        return result, warnings
    warnings.append(f"عدد الصفوف {len(rows)} بدل {len(services)} — محاذاة بالترتيب والاسم")
    # The sheet keeps the service order, so find the best order-preserving
    # matching (a row may be dropped, a service may be missing from the sheet).
    n, m = len(rows), len(services)
    score = [[_service_score(r["service"], s["name"]) for s in services] for r in rows]
    best = [[0.0] * (m + 1) for _ in range(n + 1)]
    take = [[False] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            best[i][j] = max(best[i - 1][j], best[i][j - 1])
            gain = score[i - 1][j - 1]
            if gain >= 0.45 and best[i - 1][j - 1] + gain > best[i][j]:
                best[i][j] = best[i - 1][j - 1] + gain
                take[i][j] = True
    i, j = n, m
    while i > 0 and j > 0:
        if take[i][j]:
            result[services[j - 1]["id"]] = rows[i - 1]
            i, j = i - 1, j - 1
        elif best[i][j] == best[i - 1][j]:
            i -= 1
        else:
            j -= 1
    used = {id(r) for r in result.values()}
    warnings += [f"تعذر تحديد الخدمة «r['service']»".replace("r['service']", r["service"])
                 for r in rows if id(r) not in used]
    return result, warnings


def build_day(data, day, path):
    document = docx.Document(path)
    rows = basic_rows(document)
    services = afraad_basic_on(data, day)
    personnel = [p.as_dict() for p in PeopleRepo(data).on_force(day, "personnel")]
    placed, warnings = align(rows, services)
    entries, linked, unlinked = {}, 0, []
    for service in services:
        row = placed.get(service["id"])
        if row is None:
            warnings.append(f"لا يوجد صف للخدمة «{service['name']}»")
            continue
        entry = {"count": field(row["count"]), "weapon": field(row["weapon"]),
                 "schedule": field(row["schedule"])}
        for shift, key in (("morning", "morning"), ("night", "night")):
            slot, status, match = person_slot(row[key], personnel)
            entry[f"{shift}_name"], entry[f"{shift}_phone"] = slot["name"], slot["phone"]
            entry[f"{shift}_person_id"] = slot["person_id"]
            linked += status == "linked"
            if status == "unlinked":
                unlinked.append({"service": service["name"], "shift": shift,
                                 "text": clean(row[key]), "candidates": (match or {}).get("candidates", [])[:2]})
        entries[service["id"]] = entry
    return entries, {"rows": len(rows), "entries": len(entries), "linked": linked,
                     "unlinked": unlinked, "warnings": warnings}


# ───────────────────────── main ─────────────────────────
def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT, help="مجلد 2026 (فيه 8/ 9/ ...)")
    parser.add_argument("--start", default="2026-08-01")
    parser.add_argument("--end", default=None, help="افتراضيًا آخر يومية موجودة")
    parser.add_argument("--write", action="store_true", help="تنفيذ الكتابة (الافتراضي: معاينة فقط)")
    parser.add_argument("--report", type=Path, default=None, help="حفظ تقرير JSON")
    args = parser.parse_args()

    files, notes = discover(args.root, args.start, args.end)
    if not files:
        print("لم أجد أي يومية أفراد في النطاق.")
        return 1
    data = store.load_data(store.ALL_DAYS)
    plan, report = {}, {}
    for day, path in files.items():
        entries, info = build_day(data, day, path)
        report[day] = {"file": str(path), **info}
        if entries:
            plan[day] = entries
    first, last = min(files), max(files)
    all_days = {datetime.fromisoformat(first).toordinal() + i for i in
                range(datetime.fromisoformat(last).toordinal() - datetime.fromisoformat(first).toordinal() + 1)}
    missing = sorted(datetime.fromordinal(o).date().isoformat() for o in all_days
                     if datetime.fromordinal(o).date().isoformat() not in files)

    print(f"range={first}..{last} files={len(files)} days_to_write={len(plan)} missing_files={missing}")
    for note in notes:
        print("note:", note)
    for day, info in report.items():
        flags = []
        if info["warnings"]:
            flags.append(f"warnings={len(info['warnings'])}")
        if info["unlinked"]:
            flags.append(f"unlinked={len(info['unlinked'])}")
        print(f"{day} rows={info['rows']} entries={info['entries']} linked={info['linked']} {' '.join(flags)}")
        for warning in info["warnings"]:
            print("    !", warning)
    if args.report:
        args.report.write_text(json.dumps({"missing": missing, "notes": notes, "days": report},
                                          ensure_ascii=False, indent=1), encoding="utf-8")
        print("report:", args.report)
    if not args.write:
        print("dry-run — أضف --write للتنفيذ.")
        return 0

    store.acquire_process_lock()
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    preimage = store.DATA_DIR / "import" / "preimage" / f"afraad-rosters-{stamp}"
    for day in plan:
        source = store.day_path(day)
        if source.exists():
            target = preimage / "files" / source.relative_to(store.DATA_DIR)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
    snapshot = store.snapshot_now()
    for day, entries in plan.items():
        data.setdefault("day_afraad", {})[day] = entries
    store.save_data(data)
    preimage.mkdir(parents=True, exist_ok=True)
    (preimage / "report.json").write_text(json.dumps(
        {"at": datetime.now().isoformat(timespec="seconds"), "snapshot": str(snapshot or ""),
         "missing": missing, "notes": notes, "days": report}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"written={len(plan)} preimage={preimage}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
