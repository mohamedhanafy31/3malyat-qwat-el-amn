#!/usr/bin/env python3
"""مقارنة اليوميات الستة المولدة بملفات الأرشيف اليدوية.

الأداة للقياس فقط: لا تكتب الملفات المولدة ولا تعدل ``data/``. ويتم استدعاء
مولدات ``backend.day_exports`` نفسها عشان يبقى القياس مطابق لمسار التطبيق.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from functools import lru_cache
from collections import Counter, defaultdict
from datetime import date, timedelta
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

KINDS = ("board", "duty", "afraad", "campaign", "counts", "regularity")
ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
DIACRITICS = re.compile(r"[\u0610-\u061a\u064b-\u065f\u0670\u06d6-\u06ed]")
PLACEHOLDER = re.compile(r"^[\s\-_ـ‐‑‒–—―.\/،؛:]*$")
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(20\d{2})(?!\d)")
_COMPARE_DATA: dict[str, Any] = {}
_COMPARE_DAY = ""
_COMPARE_EXACT: dict[str, str | None] = {}
_COMPARE_PREFIX: dict[tuple[str, ...], str | None] = {}
_COMPARE_FIRST: set[str] = set()


def normalize_arabic(value: object) -> str:
    """تطبيع مقارنة متعمد من غير ما نغير النص الذي يظهر في التقرير."""
    text = unicodedata.normalize("NFKC", str(value or "")).translate(ARABIC_DIGITS)
    text = DIACRITICS.sub("", text).replace("ـ", "")
    text = text.translate(str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا",
                                        "ى": "ي", "ة": "ه", "ؤ": "و", "ئ": "ي"}))
    if PLACEHOLDER.fullmatch(text):
        return ""
    text = re.sub(r"[\-_‐‑‒–—―،؛,;:|\\/()【】\[\]{}\"'«»]+", " ", text)
    return " ".join(text.lower().split())


def _compact(value: object) -> str:
    return normalize_arabic(value).replace(" ", "")


def _unique_cells(row) -> list[str]:
    """يرجع الخلايا المنطقية مرة واحدة بدل تكرار خلية Word المدمجة."""
    seen, values = set(), []
    for cell in row.cells:
        marker = id(cell._tc)
        if marker not in seen:
            seen.add(marker)
            values.append(cell.text)
    return values


def _document(source: Path | bytes | bytearray | BytesIO) -> Document:
    if isinstance(source, Path):
        return Document(source)
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    return Document(source)


def _rows(source: Path | bytes | bytearray | BytesIO, table: int = 0) -> list[list[str]]:
    document = _document(source)
    if table >= len(document.tables):
        return []
    return [_unique_cells(row) for row in document.tables[table].rows]


def _grid_rows(source: Path | bytes | bytearray | BytesIO, table: int = 0) -> list[list[str]]:
    """شبكة Word بعد توسيع الخلايا المدمجة لأعمدتها الفعلية."""
    document = _document(source)
    if table >= len(document.tables):
        return []
    return [[cell.text for cell in row.cells] for row in document.tables[table].rows]


def _result(day: str, kind: str) -> dict[str, Any]:
    return {"day": day, "kind": kind, "status": "compared", "mismatches": [],
            "layout_differences": [], "fields": {}}


def _compare_value(result: dict[str, Any], section: str, row_key: str, field: str,
                   archive: object, system: object) -> None:
    stats = result["fields"].setdefault(field, {"compared": 0, "matched": 0, "mismatched": 0})
    stats["compared"] += 1
    if normalize_arabic(archive) == normalize_arabic(system):
        stats["matched"] += 1
        return
    stats["mismatched"] += 1
    result["mismatches"].append({
        "day": result["day"], "kind": result["kind"], "section": section,
        "row_key": row_key, "field": field,
        "archive_value": "" if archive is None else str(archive),
        "system_value": "" if system is None else str(system),
    })


@lru_cache(maxsize=8192)
def _resolved_identity_cached(day: str, text: str) -> str | None:
    value = text
    if not normalize_arabic(value) or not _COMPARE_DATA or not _COMPARE_DAY:
        return None
    from importer.person_match import _rank_and_name, resolve_person
    # الرتبة والتليفون بيتشالوا قبل المفتاح (ا.ش/ اسم)، وإلا الاسم مايتلقاش.
    exact_key = _person_identity(_rank_and_name(value)[1] or value)
    if exact_key in _COMPARE_EXACT:
        return _COMPARE_EXACT[exact_key]
    tokens = exact_key.split()
    prefix = tuple(tokens[:2])
    if len(prefix) == 2 and prefix in _COMPARE_PREFIX:
        return _COMPARE_PREFIX[prefix]
    if not tokens or tokens[0] not in _COMPARE_FIRST:
        return None
    matches = [resolve_person(_COMPARE_DATA, str(value), day, role=role)
               for role in ("officer", "personnel", "conscript-driver")]
    matches = [match for match in matches if match.person_id]
    if not matches:
        return None
    matches.sort(key=lambda match: (-match.score, match.person_id or ""))
    if len(matches) > 1 and matches[0].person_id != matches[1].person_id and abs(matches[0].score - matches[1].score) < .015:
        return None
    return matches[0].person_id


def _resolved_identity(value: object) -> str | None:
    return _resolved_identity_cached(_COMPARE_DAY, str(value or ""))


def _compare_person_identity(result: dict[str, Any], field: str, archive: object, system: object) -> None:
    """مقياس منفصل: الهوية لو النص اتحل، والنص كاحتياطي لو ماتحلش."""
    left, right = _resolved_identity(archive), _resolved_identity(system)
    key = field + "_identity"
    stats = result["fields"].setdefault(key, {"compared": 0, "matched": 0, "mismatched": 0,
                                                "resolved": 0, "text_fallback": 0})
    stats["compared"] += 1
    if left and right:
        stats["resolved"] += 1
        matched = left == right
    else:
        stats["text_fallback"] += 1
        matched = normalize_arabic(archive) == normalize_arabic(system)
    stats["matched" if matched else "mismatched"] += 1


def _key(row: dict[str, Any], fields: tuple[str, ...]) -> str:
    # نحافظ على حدود الكلمات عشان نعرف نطابق الاسم المختصر
    # «خالد خلف» بالاسم الكامل «خالد خلف الله».
    return "|".join(normalize_arabic(row.get(field, "")) for field in fields)


def _key_similarity(left: str, right: str) -> float:
    if left == right:
        return 1.0
    # الاسم المختصر (كلمتين) هو نفس هوية الاسم الثلاثي/الرباعي،
    # وفرق الطول نفسه يتقاس كفرق label بعد مطابقة الصفين.
    left_words, right_words = left.split(), right.split()
    if min(len(left_words), len(right_words)) >= 2:
        short, long = sorted((left_words, right_words), key=len)
        if long[:len(short)] == short:
            return 1.0
    sequence = SequenceMatcher(None, left, right).ratio() if left and right else float(left == right)
    a, b = set(left_words), set(right_words)
    # الاحتواء مفيد لـ«مأمورية علاقات»/«مأمورية علاقات تصوير»، لكن
    # مجرد نفس الكلمات بترتيب مختلف لا يجعل اسمي ضابطين هوية واحدة.
    containment = min(0.99, min(len(a), len(b)) / max(len(a), len(b)) + 0.35) \
        if a and b and a != b and (a < b or b < a) else 0.0
    return max(sequence, containment)


def _service_identity(value: object) -> str:
    """مفتاح الخدمة بعد فصل الزمن/الجهة اللي لهم حقول مستقلة."""
    from backend.service_equivalence import service_parts

    raw = re.sub(r"\([^)]*\)\s*$", " ", str(value or ""))
    parts = list(service_parts(raw))
    if len(parts) > 1 and "افطار" in parts:
        parts.remove("افطار")
    key = " + ".join(parts)
    return key.replace("ترحيه نيابه", "ترحيله نيابه")


def _person_identity(value: object) -> str:
    key = normalize_arabic(value)
    key = re.sub(r"\b01\d{9}\b", " ", key)
    key = re.sub(r"^(?:عقيد|مقدم|رايد|نقيب|ملازم(?: اول)?|م\s*(?:اول|ش)|ا\s*ش|رقيب|مساعد)\s+", "", key)
    return " ".join(key.split())


def match_rows(archive_rows: list[dict[str, Any]], system_rows: list[dict[str, Any]],
               key_fields: tuple[str, ...], *, value_fields: tuple[str, ...] = (),
               fuzzy: bool = True) -> list[tuple[dict[str, Any] | None, dict[str, Any] | None]]:
    """يطابق المحتوى، مش رقم الصف أو ترتيب التكرار.

    لما نفس المفتاح يتكرر (زي «راحة»)، الاقتران العالمي بيختار أقرب
    قيمة بدل أول ظهور. ده يمنع تبديل صفين من إنه يتحسب فرقين.
    """
    edges: list[tuple[float, int, int]] = []
    for left_index, archive_row in enumerate(archive_rows):
        archive_key = _key(archive_row, key_fields)
        for right_index, system_row in enumerate(system_rows):
            system_key = _key(system_row, key_fields)
            key_score = _key_similarity(archive_key, system_key)
            if archive_key != system_key and (not fuzzy or key_score < 0.84):
                continue
            value_scores = [SequenceMatcher(None, _compact(archive_row.get(field, "")),
                                            _compact(system_row.get(field, ""))).ratio()
                            for field in value_fields]
            value_score = sum(value_scores) / len(value_scores) if value_scores else 0.0
            # تطابق المفتاح أساسي؛ القيمة تفك تعادل المفتاح المكرر.
            identity_score = 100.0 if archive_key == system_key else key_score * 10
            edges.append((identity_score + value_score, left_index, right_index))
    left_used: set[int] = set()
    right_used: set[int] = set()
    paired: dict[int, int] = {}
    for _score, left_index, right_index in sorted(edges, reverse=True):
        if left_index in left_used or right_index in right_used:
            continue
        left_used.add(left_index)
        right_used.add(right_index)
        paired[left_index] = right_index
    pairs = [(row, system_rows[paired[index]] if index in paired else None)
             for index, row in enumerate(archive_rows)]
    pairs.extend((None, row) for index, row in enumerate(system_rows) if index not in right_used)
    return pairs


def _compare_rows(result: dict[str, Any], section: str, archive_rows: list[dict[str, Any]],
                  system_rows: list[dict[str, Any]], key_fields: tuple[str, ...],
                  fields: tuple[str, ...], *, fuzzy: bool = True,
                  match_value_fields: tuple[str, ...] | None = None,
                  person_fields: tuple[str, ...] = ()) -> None:
    occurrences: Counter[str] = Counter()
    value_fields = match_value_fields or tuple(field for field in fields if field not in key_fields)
    for archive, system in match_rows(archive_rows, system_rows, key_fields,
                                      value_fields=value_fields, fuzzy=fuzzy):
        probe = archive or system or {}
        base = " | ".join(str(probe.get(field, "")) for field in key_fields) or "(empty)"
        occurrences[normalize_arabic(base)] += 1
        row_key = f"{base}#{occurrences[normalize_arabic(base)]}"
        if archive is None or system is None:
            difference = "extra_in_system" if archive is None else "missing_in_system"
            result["mismatches"].append({
                "day": result["day"], "kind": result["kind"], "section": section,
                "row_key": row_key, "field": "row", "difference": difference,
                "archive_value": "" if archive is None else json.dumps(archive, ensure_ascii=False, default=str),
                "system_value": "" if system is None else json.dumps(system, ensure_ascii=False, default=str),
            })
            present = archive or system or {}
            for field in fields:
                if field not in present:
                    continue
                stats = result["fields"].setdefault(field, {"compared": 0, "matched": 0, "mismatched": 0})
                stats["compared"] += 1
                stats["mismatched"] += 1
            continue
        for field in fields:
            # بعض القوالب التاريخية ماكانش فيها الحقل أصلًا (مثل هاتف/مركبة
            # الانتظام القديم). غياب العمود مش قيمة فاضية وماينفعش نقيسه.
            if archive is not None and system is not None and (field not in archive or field not in system):
                continue
            _compare_value(result, section, row_key, field,
                           archive.get(field, "") if archive else "",
                           system.get(field, "") if system else "")
            if field in person_fields:
                _compare_person_identity(result, field, archive.get(field, ""), system.get(field, ""))


_BOARD_SECTIONS = {
    "الخدمات الاساسيه": "basic", "خدمات الاساسيه": "basic",
    "الخدمات الطاريه": "emergency", "الطواري": "emergency",
    "الاهداف": "targets", "معسكر فرعي": "subcamp",
    "المعسكر الفرعي": "subcamp", "ضابط عظيم وامن المعسكر الفرعي": "subcamp",
    "ضابط عظيم الاداره": "great", "ضابط الامن بالاداره": "security",
    "ضابط الامن": "security", "خدمات سجن قوات الامن": "prison",
    "خدمات السجن": "prison", "عمل بالاداره": "admin_work",
    "الخوارج": "outsiders", "الراحات": "rests", "تقصيره": "rests",
}
_STATUS_LABELS = {"راحه", "راحه اسبوعيه", "راحه شهريه", "ماموريه", "فرقه", "مرضي", "اجازه", "اجازات"}


def _board_section(label: object) -> str | None:
    key = normalize_arabic(label)
    for heading, section in _BOARD_SECTIONS.items():
        if normalize_arabic(heading) in key:
            return section
    return None


def _looks_person(value: object) -> bool:
    key = normalize_arabic(value)
    return bool(re.search(r"(?:^| )(?:عقيد|مقدم|رايد|نقيب|م\s*(?:اول|ش)|ا\s*ش|رقيب|مساعد|مندوب)(?: |$)", key))


def _board_rows(source) -> dict[str, Any]:
    """يستخرج نصفي اللوحة كل واحد لوحده باستخدام مستخرج المستورد.

    ده مهم لأن صف Word ممكن يحمل خدمة طوارئ في اليمين وعنوان قسم
    في الشمال؛ الاتنين مش صف منطقي واحد.
    """
    from importer.extract import _extract_board, read_docx_raw

    raw = read_docx_raw(source if isinstance(source, Path) else BytesIO(source.getvalue() if isinstance(source, BytesIO) else source))
    records = _extract_board({"date": "", "role": "board", "path": "memory"}, raw)
    sections: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current = {0: "basic", 1: "emergency"}
    previous_service: dict[int, str] = {}
    previous_slot_value: dict[int, str] = {}
    half_labels = next((record.get("half_labels") or [] for record in records
                        if record.get("record_type") == "header"), [])
    for half, label in enumerate(half_labels):
        current[half] = _board_section(label) or current.get(half, "emergency")
    for record in records:
        record_type = record.get("record_type")
        if record_type not in {"board_row", "board_label"}:
            continue
        half = int(record.get("half") or 0)
        label = str(record.get("label") or "").strip()
        value = str(record.get("manning") or "").strip()
        heading = _board_section(label)
        if heading and not value:
            current[half] = heading
            previous_service.pop(half, None)
            continue
        section = current.get(half, "basic" if half == 0 else "emergency")
        label_key = normalize_arabic(label)
        if section in {"subcamp", "great", "security"} and ("صباح" in label_key or "ليل" in label_key):
            service = "ليلية" if "ليل" in label_key else "صباحية"
            if not value:
                value = previous_slot_value.get(half, "")
            if value:
                previous_slot_value[half] = value
            row = {"identity": service, "service": service, "value": value,
                   "row_index": record.get("row_index"), "half": half}
            sections[section].append(row)
            continue
        # العناوين الفرعية بعد «الخوارج» لازم تفضل قسم منطقي، مش صف بيانات.
        if record_type == "board_label" and section in {"outsiders", "rests"} and label_key in _STATUS_LABELS:
            continue
        # السطر المنفرد باسم ضابط تحت هدف هو تغطية زمنية تانية لنفس الهدف.
        if record_type == "board_label" and section == "targets" and _looks_person(label) and previous_service.get(half):
            service, value = previous_service[half], label
        elif section.startswith("outsiders") or section in {"rests", "admin_work"}:
            if label_key in _STATUS_LABELS and _looks_person(value):
                service, value = value, label
            elif _looks_person(label):
                service = label
            else:
                continue
        elif record_type == "board_label":
            # خانة فاضية/عنوان عرض، مفيش محتوى يتقارن.
            continue
        else:
            service = label
        if not normalize_arabic(service) and not normalize_arabic(value):
            continue
        identity = _service_identity(service)
        if section == "targets" and normalize_arabic(service) == normalize_arabic("السخنة"):
            identity = "كهرباء السخنة"
        if section.startswith("outsiders") or section in {"rests", "admin_work"}:
            identity = _person_identity(service)
        row = {"identity": identity, "service": service, "value": value,
               "row_index": record.get("row_index"), "half": half}
        sections[section].append(row)
        if section == "targets" and normalize_arabic(service):
            previous_service[half] = service
    return {"sections": dict(sections)}


def _board_packing(rows: dict[str, Any]) -> set[tuple[str, str]]:
    by_physical: dict[int, dict[int, str]] = defaultdict(dict)
    for section, entries in rows["sections"].items():
        for entry in entries:
            by_physical[int(entry.get("row_index") or 0)][int(entry.get("half") or 0)] = (
                f"{section}|{_key(entry, ('service',))}")
    return {(halves[0], halves[1]) for halves in by_physical.values() if 0 in halves and 1 in halves}


def compare_board(day: str, archive, system_bytes: bytes) -> dict[str, Any]:
    result = _result(day, "board")
    left, right = _board_rows(archive), _board_rows(system_bytes)
    for section in sorted(set(left["sections"]) | set(right["sections"])):
        _compare_rows(result, section, left["sections"].get(section, []),
                      right["sections"].get(section, []), ("identity",), ("service", "value"),
                      person_fields=("value",))
    # تختيم نصفي اللوحة على نفس صف Word اختلاف layout فقط. نقيسه
    # بين العناصر الموجودة في النسختين عشان الناقص مايتحسبش layout كمان.
    left_items = {f"{section}|{_key(row, ('service',))}" for section, entries in left["sections"].items() for row in entries}
    right_items = {f"{section}|{_key(row, ('service',))}" for section, entries in right["sections"].items() for row in entries}
    common = left_items & right_items
    left_packing = {pair for pair in _board_packing(left) if pair[0] in common and pair[1] in common}
    right_packing = {pair for pair in _board_packing(right) if pair[0] in common and pair[1] in common}
    if left_packing != right_packing:
        result["layout_differences"].append({
            "day": day, "kind": "board", "section": "five_column_packing",
            "archive_value": sorted(left_packing - right_packing),
            "system_value": sorted(right_packing - left_packing),
        })
    # ترتيب الأهداف مقفول في المنتج؛ باقي الأقسام ترتيبها لا يغير المعنى.
    left_order = [_key(row, ("identity",)) for row in left["sections"].get("targets", [])]
    right_order = [_key(row, ("identity",)) for row in right["sections"].get("targets", [])]
    shared = set(left_order) & set(right_order)
    if [value for value in left_order if value in shared] != [value for value in right_order if value in shared]:
        result["layout_differences"].append({
            "day": day, "kind": "board", "section": "targets_order",
            "archive_value": [value for value in left_order if value in shared],
            "system_value": [value for value in right_order if value in shared],
        })
    return result


def _duty_officers(source) -> list[dict[str, str]]:
    out = []
    for cells in _rows(source, 0)[1:]:
        if len(cells) < 3 or not normalize_arabic(cells[2]):
            continue
        out.append({"rank": cells[1] if len(cells) > 1 else "", "name": cells[2],
                    "assigned": cells[3] if len(cells) > 3 else "",
                    "daily": cells[4] if len(cells) > 4 else "",
                    "rest": cells[5] if len(cells) > 5 else ""})
    return out


def _duty_summary(source) -> list[dict[str, str]]:
    # الرأس مدمج لكن القيم موزعة على 17 عمود، فالشبكة أدق من الخلايا المنطقية هنا.
    rows = _grid_rows(source, 1)
    if len(rows) < 3:
        return []
    width = max(map(len, rows))
    result = []
    for index in range(width):
        result.append({"index": str(index),
                       "metric": f"{rows[0][index] if index < len(rows[0]) else ''} / {rows[1][index] if index < len(rows[1]) else ''}",
                       "counted": rows[2][index] if index < len(rows[2]) else "",
                       "displayed": rows[3][index] if len(rows) > 3 and index < len(rows[3]) else ""})
    return result


def compare_duty(day: str, archive, system_bytes: bytes) -> dict[str, Any]:
    result = _result(day, "duty")
    left_officers, right_officers = _duty_officers(archive), _duty_officers(system_bytes)
    for row in left_officers + right_officers:
        row["identity"] = _person_identity(row.get("name"))
    _compare_rows(result, "officers", left_officers, right_officers,
                  ("identity",), ("rank", "name", "assigned", "daily", "rest"), person_fields=("name",))
    _compare_rows(result, "summary", _duty_summary(archive), _duty_summary(system_bytes),
                  ("index",), ("metric", "counted", "displayed"), fuzzy=False)
    return result


def _afraad_rows(source) -> dict[str, list[dict[str, str]]]:
    rows = _rows(source)
    grid = _grid_rows(source)
    heading = next((i for i, row in enumerate(rows) if "طاري" in normalize_arabic(" ".join(row))), len(rows))
    header = [normalize_arabic(value) for value in (grid[0] if grid else [])]
    def column(needle: str, fallback: int) -> int:
        return next((i for i, value in enumerate(header) if needle in value), fallback)
    columns = {"service": column("خدمه", 1), "morning": column("صباح", 3),
               "night": column("ليل", 5), "strength": column("قوام", 8),
               "weapon": column("تسليح", 11), "time": column("انتظام", 12)}
    basic = []
    for cells in grid[1:heading]:
        if len(cells) > max(columns.values()):
            basic.append({field: cells[index] for field, index in columns.items()})
    emergency = []
    for cells in rows[heading + 1:]:
        if any(normalize_arabic(cell) for cell in cells):
            emergency.append({"service": cells[0] if cells else "", "person": cells[1] if len(cells) > 1 else "",
                              "strength": cells[2] if len(cells) > 2 else "", "weapon": cells[3] if len(cells) > 3 else "",
                              "time": cells[4] if len(cells) > 4 else ""})
    return {"basic": basic, "emergency": emergency}


def compare_afraad(day: str, archive, system_bytes: bytes) -> dict[str, Any]:
    result = _result(day, "afraad")
    left, right = _afraad_rows(archive), _afraad_rows(system_bytes)
    for side in (left, right):
        for entries in side.values():
            for row in entries:
                row["identity"] = _service_identity(row.get("service"))
    _compare_rows(result, "basic", left["basic"], right["basic"], ("identity",),
                  ("service", "morning", "night", "strength", "weapon", "time"))
    _compare_rows(result, "emergency", left["emergency"], right["emergency"], ("identity",),
                  ("service", "person", "strength", "weapon", "time"), person_fields=("person",))
    return result


def _campaign_parts(source) -> dict[str, list[dict[str, str]]]:
    rows = _rows(source)
    divider = next((i for i, row in enumerate(rows) if "خط" in normalize_arabic(" ".join(row))
                    and "توص" in normalize_arabic(" ".join(row))), len(rows))
    def plate_row(row: list[str]) -> bool:
        values = [normalize_arabic(x) for x in row if normalize_arabic(x)]
        return bool(values) and all("/" in raw and any(c.isdigit() for c in raw) for raw in row if normalize_arabic(raw))
    plate_start = next((i for i, row in enumerate(rows[divider + 1:], divider + 1) if plate_row(row)), len(rows))
    fields = ("ordinal", "service", "driver", "plate", "vehicle")
    def convert(source_rows):
        return [{field: (row[index] if index < len(row) else "") for index, field in enumerate(fields)}
                for row in source_rows if any(normalize_arabic(x) for x in row)]
    plates = [{"plate": cell} for row in rows[plate_start:] for cell in row
              if normalize_arabic(cell) and "/" in cell and any(char.isdigit() for char in cell)]
    return {"checkpoints": convert(rows[2:divider]),
            "delivery": convert(rows[divider + 2:plate_start]),
            "plates": plates}


def compare_campaign(day: str, archive, system_bytes: bytes) -> dict[str, Any]:
    result = _result(day, "campaign")
    left, right = _campaign_parts(archive), _campaign_parts(system_bytes)
    for side in (left, right):
        for section, entries in side.items():
            for row in entries:
                row["identity"] = _service_identity(row.get("service"))
                row["plate_key"] = "".join(re.findall(r"\d+", str(row.get("plate") or "")))
    for section in ("checkpoints", "delivery"):
        _compare_rows(result, section, left[section], right[section], ("identity",),
                      ("service", "driver", "plate", "vehicle"),
                      match_value_fields=("plate_key", "plate_key", "plate_key", "driver", "vehicle"),
                      person_fields=("driver",))
    _compare_rows(result, "plates", left["plates"], right["plates"], ("plate_key",), ("plate",))
    return result


def _workbook(source):
    return load_workbook(source if isinstance(source, Path) else BytesIO(source), data_only=True)


def _counts_rows(source) -> dict[str, list[dict[str, Any]]]:
    workbook = _workbook(source)
    sheet = workbook[workbook.sheetnames[0]]
    def pairs(numbers: Iterable[int]) -> list[dict[str, Any]]:
        result = []
        for row in numbers:
            for name_col, count_col in (("B", "C"), ("D", "E"), ("F", "G")):
                name, count = sheet[f"{name_col}{row}"].value, sheet[f"{count_col}{row}"].value
                if name not in (None, "") or count not in (None, ""):
                    result.append({"service": "" if name is None else str(name), "count": "" if count is None else str(count)})
        return result
    result = {"basic_am": pairs(range(4, 13)), "basic_pm": pairs(range(14, 24)),
              "emergency": pairs(range(27, 36))}
    workbook.close()
    return result


def compare_counts(day: str, archive, system_bytes: bytes) -> dict[str, Any]:
    result = _result(day, "counts")
    left, right = _counts_rows(archive), _counts_rows(system_bytes)
    for side in (left, right):
        for entries in side.values():
            for row in entries:
                row["identity"] = _service_identity(row.get("service"))
    for section in ("basic_am", "basic_pm", "emergency"):
        _compare_rows(result, section, left[section], right[section], ("identity",), ("service", "count"))
    return result


def _regularity_rows(source) -> dict[str, list[dict[str, str]]]:
    grid = _grid_rows(source, 0)
    header = [normalize_arabic(value) for value in (grid[0] if grid else [])]
    def find(needles: tuple[str, ...]) -> int | None:
        return next((index for index, value in enumerate(header)
                     if all(needle in value for needle in needles)), None)
    columns = {"service": find(("خدمه",)), "person": find(("قايم", "خدمه")),
               "phone": find(("هاتف",)) if find(("هاتف",)) is not None else find(("تلفون",)),
               "vehicle": find(("مركبه",)), "place": find(("مكان", "انتظام")),
               "time": find(("ساعه", "انتظام")), "count": find(("عدد", "مجند"))}
    columns = {field: index for field, index in columns.items() if index is not None}
    services = []
    for cells in grid[1:]:
        if "service" in columns and len(cells) > columns["service"] and normalize_arabic(cells[columns["service"]]):
            services.append({field: cells[index] if index < len(cells) else "" for field, index in columns.items()})
    officer_grid = _grid_rows(source, 1)
    officer_header = [normalize_arabic(value) for value in (officer_grid[0] if officer_grid else [])]
    def officer_col(needle: str, fallback: int) -> int:
        return next((i for i, value in enumerate(officer_header) if needle in value), fallback)
    officer_columns = {"rank": officer_col("رتبه", 1), "name": officer_col("اسم", 2),
                       "service": officer_col("خدمه", 3)}
    officers = []
    for cells in officer_grid[1:]:
        name_index = officer_columns["name"]
        if len(cells) > name_index and normalize_arabic(cells[name_index]):
            officers.append({field: cells[index] if index < len(cells) else ""
                             for field, index in officer_columns.items()})
    return {"services": services, "officers": officers}


def compare_regularity(day: str, archive, system_bytes: bytes) -> dict[str, Any]:
    result = _result(day, "regularity")
    left, right = _regularity_rows(archive), _regularity_rows(system_bytes)
    for side in (left, right):
        for row in side["services"]:
            row["identity"] = _service_identity(row.get("service"))
        for row in side["officers"]:
            row["identity"] = _person_identity(row.get("name"))
    _compare_rows(result, "services", left["services"], right["services"], ("identity", "time"),
                  ("service", "person", "phone", "vehicle", "place", "time", "count"),
                  person_fields=("person",))
    _compare_rows(result, "officers", left["officers"], right["officers"], ("identity",),
                  ("rank", "name", "service"), person_fields=("name",))
    return result


COMPARATORS: dict[str, Callable[[str, Any, bytes], dict[str, Any]]] = {
    "board": compare_board, "duty": compare_duty, "afraad": compare_afraad,
    "campaign": compare_campaign, "counts": compare_counts, "regularity": compare_regularity,
}


def _kind_for_name(name: str) -> str | None:
    low = name.lower()
    if name.startswith("~$") or low.endswith(".pdf") or low.endswith(".tmp"):
        return None
    if re.fullmatch(r"\d+\.\.docx", name):
        return "regularity"
    if re.fullmatch(r"\d+-\d+-20\d\d\s+افراد\.docx", name):
        return "afraad"
    if re.fullmatch(r"\d+-\d+-20\d\d\.docx", name):
        return "duty"
    if name.startswith("يومية الحملة") and low.endswith(".docx"):
        return "campaign"
    if name.startswith("اعداد الخدمات") and Path(name).suffix.lower() in {".xlsx", ".docx", ".doc"}:
        return "counts"
    if re.fullmatch(r"\d+\.docx", name):
        return "board"
    return None


def _named_day(path: Path, kind: str, archive_root: Path) -> str | None:
    translated = path.name.translate(ARABIC_DIGITS)
    hit = DATE_RE.search(translated)
    if hit:
        try:
            return date(int(hit.group(3)), int(hit.group(2)), int(hit.group(1))).isoformat()
        except ValueError:
            return None
    if kind in {"board", "regularity"}:
        day_hit = re.match(r"(\d+)", translated)
        try:
            month = int(path.parent.parent.name)
            year = int(archive_root.name) if archive_root.name.isdigit() else 2026
            return date(year, month, int(day_hit.group(1))).isoformat() if day_hit else None
        except ValueError:
            return None
    return None


def discover_archive(archive_root: Path) -> tuple[dict[tuple[str, str], Path], list[dict[str, Any]]]:
    """جرد باليوم المطبوع في الاسم، مش فولدر اليوم فقط."""
    candidates: dict[tuple[str, str], list[Path]] = defaultdict(list)
    ignored = []
    for path in archive_root.glob("*/*/*"):
        if not path.is_file():
            continue
        kind = _kind_for_name(path.name)
        if not kind:
            continue
        day = _named_day(path, kind, archive_root)
        if day:
            candidates[(day, kind)].append(path)
        else:
            ignored.append({"path": str(path), "reason": "date_not_resolved"})
    chosen = {}
    for key, paths in candidates.items():
        day = date.fromisoformat(key[0])
        # النسخة الموضوعة في فولدر يومها لها الأولوية، مع الاحتفاظ بالنسخة المزاحة لو مفيش غيرها.
        chosen[key] = sorted(paths, key=lambda p: (p.parent.name == str(day.day), p.suffix.lower() == ".xlsx",
                                                   p.stat().st_mtime_ns), reverse=True)[0]
        for duplicate in paths:
            if duplicate != chosen[key]:
                ignored.append({"path": str(duplicate), "reason": "duplicate", "chosen": str(chosen[key])})
    return chosen, ignored


def _days(start: str, end: str) -> list[str]:
    current, stop = date.fromisoformat(start), date.fromisoformat(end)
    output = []
    while current <= stop:
        output.append(current.isoformat())
        current += timedelta(days=1)
    return output


def summarize(results: list[dict[str, Any]], start: str, end: str, kinds: list[str]) -> dict[str, Any]:
    summary: dict[str, Any] = {"range": {"from": start, "to": end}, "kinds": {}}
    for kind in kinds:
        selected = [r for r in results if r["kind"] == kind]
        fields: dict[str, Counter] = defaultdict(Counter)
        mismatch_count = 0
        layout_count = 0
        for result in selected:
            mismatch_count += len(result.get("mismatches", []))
            layout_count += len(result.get("layout_differences", []))
            for field, values in result.get("fields", {}).items():
                fields[field].update(values)
        totals = Counter()
        for values in fields.values():
            totals.update(values)
        summary["kinds"][kind] = {
            "found": sum(r["status"] != "archive_missing" for r in selected),
            "compared": sum(r["status"] == "compared" for r in selected),
            "skipped": sum(r["status"] != "compared" for r in selected),
            "skip_reasons": dict(Counter(r["status"] for r in selected if r["status"] != "compared")),
            "mismatches": mismatch_count,
            "layout_differences": layout_count,
            "fields": {field: dict(values) for field, values in sorted(fields.items())},
            "totals": dict(totals),
            "match_rate": round(totals["matched"] / totals["compared"], 6) if totals["compared"] else None,
        }
    return summary


def run(start: str, end: str | None, kinds: list[str], archive_root: Path) -> dict[str, Any]:
    global _COMPARE_DATA, _COMPARE_DAY, _COMPARE_EXACT, _COMPARE_PREFIX, _COMPARE_FIRST
    inventory, ignored = discover_archive(archive_root)
    available_days = sorted(day for day, _kind in inventory if day >= start)
    if end is None:
        if not available_days:
            raise RuntimeError("مفيش ملفات أرشيف في النطاق")
        end = available_days[-1]
    from backend import day_exports
    from backend import store
    _COMPARE_DATA = store.assemble()
    exact: dict[str, set[str]] = defaultdict(set)
    prefixes: dict[tuple[str, ...], set[str]] = defaultdict(set)
    for group in ("officers", "personnel", "conscripts"):
        for person in _COMPARE_DATA.get(group) or []:
            for label in (person.get("name"), person.get("short_name"), person.get("nickname")):
                if label:
                    key = _person_identity(label)
                    exact[key].add(person.get("id"))
                    words = key.split()
                    if len(words) >= 2:
                        prefixes[tuple(words[:2])].add(person.get("id"))
    _COMPARE_EXACT = {key: next(iter(ids)) if len(ids) == 1 else None for key, ids in exact.items()}
    _COMPARE_PREFIX = {key: next(iter(ids)) if len(ids) == 1 else None for key, ids in prefixes.items()}
    _COMPARE_FIRST = {key[0] for key in prefixes}
    _resolved_identity_cached.cache_clear()
    results = []
    for day in _days(start, end):
        _COMPARE_DAY = day
        for kind in kinds:
            archive = inventory.get((day, kind))
            if archive is None:
                results.append({"day": day, "kind": kind, "status": "archive_missing", "archive": None,
                                "mismatches": [], "layout_differences": [], "fields": {}})
                continue
            if archive.suffix.lower() in {".doc", ".xls"}:
                results.append({"day": day, "kind": kind, "status": "unsupported_legacy_format",
                                "archive": str(archive), "mismatches": [], "fields": {}})
                continue
            try:
                generated = getattr(day_exports, f"build_{kind}_bytes")(day)
                compared = COMPARATORS[kind](day, archive, generated)
                compared["archive"] = str(archive)
                results.append(compared)
            except Exception as exc:  # نحتفظ بباقي الأيام للتشخيص المجمع.
                results.append({"day": day, "kind": kind, "status": "comparison_error", "archive": str(archive),
                                "error": f"{type(exc).__name__}: {exc}", "mismatches": [],
                                "layout_differences": [], "fields": {}})
    return {"summary": summarize(results, start, end, kinds), "results": results,
            "inventory": {f"{day}:{kind}": str(path) for (day, kind), path in sorted(inventory.items())
                          if start <= day <= end and kind in kinds},
            "ignored": ignored}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--from", dest="start", default="2026-08-01")
    parser.add_argument("--to", dest="end")
    parser.add_argument("--kinds", default=",".join(KINDS))
    parser.add_argument("--archive-root", type=Path, default=ROOT.parent.parent / "2026")
    parser.add_argument("--json", type=Path)
    args = parser.parse_args(argv)
    kinds = [value.strip() for value in args.kinds.split(",") if value.strip()]
    unknown = sorted(set(kinds) - set(KINDS))
    if unknown:
        parser.error("أنواع غير معروفة: " + ", ".join(unknown))
    report = run(args.start, args.end, kinds, args.archive_root.resolve())
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
