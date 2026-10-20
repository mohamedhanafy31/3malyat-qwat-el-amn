#!/usr/bin/env python3
"""Build the seed data for the «الحملة» module from the archive.

Sources (read-only):
  * 2026/9/*/اضافات/خدمات الحمله سيارات ومجندين من داخل قوات الامن.docx
      -> the fleet inventory, grouped by type heading (69 plates);
  * 2026/9/*/اضافات/دفوع الحملة المحدث جدا.docx
      -> the conscripts marked «سواق» (names only — no batch/ranking/role data);
  * 2026/{8,9}/*/يومية الحملة عن يوم *.docx
      -> vehicle model text (شفورلية، بانسير …) and plates that are in use but
         missing from the inventory (reported, never auto-added).

Output: migrations/017_campaign_seed.json, consumed by migrations/017_seed_campaign.py.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
ARCHIVE = ROOT.parent.parent / "2026"
OUT = ROOT / "migrations" / "017_campaign_seed.json"

DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")
FLEET_DOC = "خدمات الحمله سيارات ومجندين من داخل قوات الامن.docx"
BATCH_DOC = "دفوع الحملة المحدث جدا.docx"

TYPE_BY_HEADING = {"مايكات": "ميك", "اللواري": "لوري", "بيك اب": "بيك اب",
                   "المدرعات": "مدرعة", "دبل كبينة": "بيك اب دوبل كبينة"}
TYPE_BY_LABEL = {"جرار": "جرار", "اسعاف": "اسعاف", "إسعاف": "اسعاف", "أتوبيس": "اوتوبيس",
                 "اتوبيس": "اوتوبيس", "نقل مياه": "نقل مياه", "ميكروباص": "ميكروباص",
                 "لوجان": "لوجان (ملاكي)"}
# التسمية الأساسية للنوع في اليوميات، لقياس تعارض النوع مع الإضافات
BUCKET = [("مدرعه", "مدرعة"), ("مدرعة", "مدرعة"), ("مايك", "ميك"), ("ميك", "ميك"),
          ("بيك", "بيك اب"), ("لوري", "لوري"), ("لورى", "لوري"), ("اتوبيس", "اوتوبيس"),
          ("اوتوبيس", "اوتوبيس")]
MODEL_STRIP = ("المدرعه", "المدرعة", "مدرعه", "مدرعة", "مايك", "ميك", "بيك اب", "لوري", "لورى")
DAILY_FIXED = ["التدخل السريع", "التدخل السريع", "خط الغاز", "القول المكبر",
               "ارتكاز الاسعاف", "سائق بورتوفيق"]

PLATE = re.compile(r"^(?:(ب\s*\d{1,2})\s*/\s*)?(\d{3,5})$")
DAILY_PLATE = re.compile(r"^(\d{4,5})\s*(?:ب\s*\d{1,2})?$")


def cells(row):
    seen, out = set(), []
    for cell in row.cells:
        if cell._tc not in seen:
            seen.add(cell._tc)
            out.append(" ".join(cell.text.split()))
    return out


def latest(pattern: str) -> Path:
    files = sorted(ARCHIVE.glob(pattern), key=lambda p: int(p.parts[-3]))
    if not files:
        raise SystemExit(f"لم أجد {pattern} تحت {ARCHIVE}")
    return files[-1]


def parse_plate(text: str):
    text = text.translate(DIGITS).strip()
    hit = PLATE.match(text)
    if not hit:
        return None
    series = re.sub(r"\s+", "", hit[1] or "")
    return series, hit[2]


def read_fleet(path: Path):
    document = docx.Document(path)
    fleet, seen = [], set()
    for table in document.tables[2:]:
        rows = [cells(r) for r in table.rows]
        heading = rows[0][0] if rows and len(rows[0]) == 1 else ""
        pairs = rows if not heading else rows[1:]
        for row in pairs:
            plate = parse_plate(row[0]) if row else None
            if not plate:
                continue
            if heading:
                vtype = TYPE_BY_HEADING[heading]
            else:
                label = row[1].strip() if len(row) > 1 else ""
                vtype = TYPE_BY_LABEL.get(label)
                if not vtype:
                    raise SystemExit(f"نوع غير معروف في جدول المركبات الخاصة: {label!r}")
            if plate in seen:
                continue
            seen.add(plate)
            entry = {"plate_series": plate[0], "plate_number": plate[1], "type": vtype,
                     "model": "", "status": "operational", "fault": "", "notes": "",
                     "archived": False}
            if not plate[0]:
                entry["notes"] = "اللوحة بدون سلسلة المحافظة في ملف الإضافات — راجعها"
            fleet.append(entry)
    return fleet


def read_conscripts(path: Path):
    document = docx.Document(path)
    names = []
    for table in document.tables[:3]:
        for row in table.rows[2:]:
            row_cells = cells(row)
            if len(row_cells) < 4 or "سواق" not in row_cells[-1]:
                continue
            name = re.sub(r"\(.*?\)", "", row_cells[1]).strip()
            if name and name not in names:
                names.append(name)
    return names


def daily_usage():
    """number -> Counter of full type strings; number -> days seen."""
    types, days = defaultdict(Counter), defaultdict(set)
    for month in (8, 9):
        for path in sorted(ARCHIVE.glob(f"{month}/*/يومية الحملة عن يوم *.docx")):
            if path.name.startswith("~$"):
                continue
            try:
                table = docx.Document(path).tables[0]
            except Exception:
                continue
            for row in table.rows:
                row_cells = cells(row)
                if len(row_cells) < 5:
                    continue
                plate = DAILY_PLATE.match(row_cells[3].translate(DIGITS))
                if plate:
                    types[plate[1]][row_cells[4]] += 1
                    days[plate[1]].add(path.parent.name + f"-{month}")
    return types, days


def bucket(text: str) -> str:
    for needle, name in BUCKET:
        if needle in text:
            return name
    return ""


def model_of(text: str) -> str:
    for word in MODEL_STRIP:
        text = text.replace(word, " ")
    return " ".join(text.split())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()

    fleet = read_fleet(latest(f"9/*/اضافات/{FLEET_DOC}"))
    conscripts = read_conscripts(latest(f"9/*/اضافات/{BATCH_DOC}"))
    types, days = daily_usage()

    inventory = {v["plate_number"]: v for v in fleet}
    conflicts = []
    for number, counter in types.items():
        vehicle = inventory.get(number)
        if not vehicle:
            continue
        text, count = counter.most_common(1)[0]
        vehicle["model"] = model_of(text) if bucket(text) else ""
        seen_bucket = bucket(text)
        shown = {"بيك اب": {"بيك اب", "بيك اب دوبل كبينة"}}.get(seen_bucket, {seen_bucket})
        if seen_bucket and vehicle["type"] not in shown and count >= 5:
            vehicle["notes"] = f"نوعها في اليوميات «{text}» ({count} يوم) — راجع النوع"
            conflicts.append({"plate": f"{vehicle['plate_series']}/{number}", "addendum": vehicle["type"],
                              "daily": text, "days": count})

    unlisted = sorted(({"plate_number": n, "type": types[n].most_common(1)[0][0], "days": len(days[n])}
                       for n in types if n not in inventory), key=lambda r: -r["days"])
    payload = {
        "vehicles": fleet,
        "conscripts": [{"name": name, "phone": "", "is_driver": True} for name in conscripts],
        "campaign_daily": DAILY_FIXED,
        "report": {"conflicts": conflicts, "unlisted_plates_seen_in_daily": unlisted,
                   "counts": Counter(v["type"] for v in fleet)},
    }
    args.out.write_text(json.dumps(payload, ensure_ascii=False, indent=1, default=dict), encoding="utf-8")
    print(f"vehicles={len(fleet)} conscripts={len(conscripts)} conflicts={len(conflicts)} "
          f"unlisted={len(unlisted)} -> {args.out}")
    for kind, count in sorted(Counter(v["type"] for v in fleet).items(), key=lambda i: -i[1]):
        print(f"  {kind}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
