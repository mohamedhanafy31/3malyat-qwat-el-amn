"""Generate the final two-column personnel contact sheet from live roster data."""
from copy import deepcopy
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend import store  # noqa: E402
from backend.roster_import import name_key  # noqa: E402


SOURCE_ROOT = ROOT.parents[1]
TEMPLATE = SOURCE_ROOT / "2026" / "10" / "22" / "داتا افراد.docx"
ORDER_SOURCE = SOURCE_ROOT / "كشف_الأفراد_2026-10-21.doc"
OUTPUT = SOURCE_ROOT / "داتا افراد النهائية 22-10-2026.docx"


def official_order():
    text = ORDER_SOURCE.read_text(encoding="utf-8-sig")
    rows = re.findall(r"<tr[^>]*>(.*?)</tr>", text, flags=re.I | re.S)
    names = []
    for row in rows[1:]:
        cells = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", cell)).strip()
                 for cell in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", row, flags=re.I | re.S)]
        if len(cells) >= 3 and cells[2]:
            names.append(cells[2])
    return names


def abbreviation(role):
    role = str(role or "").strip()
    if role.startswith("رقيب"):
        return "رقيب"
    if role.startswith(("معاون", "جندي")):
        return "م.ش"
    return "ا.ش"


def set_bidi(paragraph):
    props = paragraph._p.get_or_add_pPr()
    if props.find(qn("w:bidi")) is None:
        # w:bidi precedes spacing/justification in the OOXML pPr sequence.
        props.insert(0, OxmlElement("w:bidi"))


def write_cell(cell, person):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraphs = [cell.paragraphs[0], cell.add_paragraph()]
    values = [f"{abbreviation(person.get('role'))}/ {person['name']}",
              person.get("phone") or "—"]
    for paragraph, value in zip(paragraphs, values):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        set_bidi(paragraph)
        run = paragraph.add_run(value)
        run.bold = True
        run.font.size = Pt(7)
        run.font.rtl = True


def generate():
    data = store.load_data(())
    active = [row for row in data.get("personnel", [])
              if row.get("status", "active") == "active"]
    by_name = {name_key(row.get("name")): row for row in active}
    ordered = []
    for name in official_order():
        person = by_name.pop(name_key(name), None)
        if person:
            ordered.append(person)
    ordered.extend(sorted(by_name.values(), key=lambda row: name_key(row.get("name"))))
    if len(ordered) != 163:
        raise RuntimeError(f"Expected 163 active personnel, found {len(ordered)}")

    document = Document(TEMPLATE)
    document.paragraphs[0].text = "بيانات الأفراد النهائية بتاريخ 22/10/2026"
    title = document.paragraphs[0]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_bidi(title)
    for run in title.runs:
        run.bold = True
        run.font.size = Pt(14)
        run.font.rtl = True

    table = document.tables[0]
    needed = (len(ordered) + 1) // 2
    while len(table.rows) < needed:
        table._tbl.append(deepcopy(table.rows[-1]._tr))
    while len(table.rows) > needed:
        table._tbl.remove(table.rows[-1]._tr)
    for index, row in enumerate(table.rows):
        for column in range(2):
            person_index = index * 2 + column
            if person_index < len(ordered):
                write_cell(row.cells[column], ordered[person_index])
            else:
                row.cells[column].text = ""

    document.core_properties.title = "بيانات الأفراد النهائية"
    document.core_properties.subject = "قائمة اتصال الأفراد على القوة بتاريخ 22/10/2026"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    generate()
