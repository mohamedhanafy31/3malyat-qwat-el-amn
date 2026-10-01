"""مرحلة الاستخراج الحرفي للوثائق التي اختارها DISCOVER."""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from .ledger import Ledger, atomic_write_jsonl
from .textnorm import dates_in_text, norm, strip_format_marks, weekday_in_text


VERSION = "2"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEADER_ALIASES = {
    "م": "م", "الرتبه": "الرتبة", "رقم الاقدميه": "رقم الأقدمية",
    "الاقدميه": "رقم الأقدمية", "الاسم": "الاسم",
    "العمل المسند اليه": "العمل المسند إليه", "التشغيل اليومي": "التشغيل اليومي",
    "الراحات": "الراحات", "الحاله": "الحالة", "رقم التلفون": "رقم التلفون",
    "رقم الهاتف": "رقم الهاتف", "ملاحظات": "ملاحظات",
    "الامضاء": "الإمضاء",
    "العمل المسند مجددا": "العمل المسند مجددًا",
}
ROSTER_SECTIONS = ("الحراسات المشدده", "الخوارج")
BOARD_SECTIONS = (
    "الخدمات الاساسيه", "الخدمات اساسيه", "الخدمات الطاريه", "الطواري", "الاهداف", "معسكر فرعي",
    "المعسكر الفرعي", "ضابط الامن", "ضابط عظيم", "تقصيره", "الخوارج", "الراحات",
    "عمل بالاداره", "فتره صباحيه", "فتره ليليه",
)
WEAPONS = ("الي", "خرطوش", "فيدرال", "دونك", "كلبش", "مايك", "بريتا")
TIME_RE = re.compile(r"(?:\b\d{1,2}\s*(?::|\.)\s*\d{1,2}\b|\b\d{1,2}\s*[صمظ]\b|الساعه|انتظام)")
PERSON_RE = re.compile(r"(?:\b01\d{9}\b|م\s*[.\/]\s*ش|ا\s*[.\/]\s*ش|رقيب|مساعد|مندوب|معاون)")
STRENGTH_RE = re.compile(r"(?:\d+\s*(?:مج|مجند|فرد)|سائق)")


@dataclass(frozen=True)
class Cell:
    text: str
    columns: tuple[int, ...]


@dataclass(frozen=True)
class Table:
    rows: tuple[tuple[Cell, ...], ...]


@dataclass(frozen=True)
class DocxData:
    paragraphs: tuple[str, ...]
    textboxes: tuple[str, ...]
    tables: tuple[Table, ...]
    body_sha1: str
    headings: tuple[str, ...]


def _raw_text(element: ET.Element, *, skip_textboxes: bool = False) -> str:
    parts: list[str] = []
    def visit(node: ET.Element) -> None:
        if skip_textboxes and node.tag == W + "txbxContent":
            return
        if node.tag == W + "t":
            parts.append(node.text or "")
        elif node.tag == W + "tab":
            parts.append("\t")
        elif node.tag in {W + "br", W + "cr"}:
            parts.append("\n")
        for child in node:
            visit(child)
    visit(element)
    return strip_format_marks("".join(parts))


def _cell_text(cell: ET.Element) -> str:
    paragraphs = [_raw_text(p) for p in cell.findall(".//" + W + "p")]
    return "\n".join(paragraphs)


def _grid_value(node: ET.Element | None, default: int = 0) -> int:
    if node is None:
        return default
    try:
        return int(node.get(W + "val", str(default)))
    except ValueError:
        return default


def read_docx_raw(path: Path) -> DocxData:
    with zipfile.ZipFile(path) as package:
        xml = package.read("word/document.xml")
    root = ET.fromstring(xml)
    body = root.find(W + "body")
    paragraphs: list[str] = []
    textboxes: list[str] = []
    tables: list[Table] = []
    headings: list[str] = []
    preceding = ""
    if body is not None:
        for element in body:
            if element.tag == W + "p":
                value = _raw_text(element, skip_textboxes=True)
                paragraphs.append(value)
                if value:
                    preceding = value
                for box in element.findall(".//" + W + "txbxContent"):
                    box_text = "\n".join(_raw_text(p) for p in box.findall(".//" + W + "p"))
                    if box_text and box_text not in textboxes:
                        textboxes.append(box_text)
            elif element.tag == W + "tbl":
                headings.append(preceding)
                parsed_rows: list[tuple[Cell, ...]] = []
                for row in element.findall(W + "tr"):
                    grid = _grid_value(row.find("./" + W + "trPr/" + W + "gridBefore"))
                    parsed_cells: list[Cell] = []
                    for cell in row.findall(W + "tc"):
                        span = max(1, _grid_value(cell.find("./" + W + "tcPr/" + W + "gridSpan"), 1))
                        columns = tuple(range(grid, grid + span))
                        parsed_cells.append(Cell(_cell_text(cell), columns))
                        grid += span
                    parsed_rows.append(tuple(parsed_cells))
                tables.append(Table(tuple(parsed_rows)))
    return DocxData(tuple(paragraphs), tuple(textboxes), tuple(tables), hashlib.sha1(xml).hexdigest(), tuple(headings))


def _selected(discover_path: Path, start: dt.date, end: dt.date) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    with discover_path.open(encoding="utf-8") as stream:
        for line in stream:
            item = json.loads(line)
            for value in item.get("chosen_for", []):
                day = dt.date.fromisoformat(value)
                if start <= day <= end:
                    selected.append({**item, "date": value, "stale": value in item.get("stale_for", [])})
    return sorted(selected, key=lambda item: (item["date"], item["role"], item["path"]))


def _provenance(doc: dict[str, Any], *, record_type: str, raw_text: str | None = None,
                raw_cells: list[str] | None = None, table: int | None = None,
                row: int | None = None, columns: list[list[int]] | list[int] | None = None,
                **extra: Any) -> dict[str, Any]:
    return {
        "record_type": record_type,
        "date": doc["date"],
        "role": doc["role"],
        "path": doc["path"],
        "file_sha1": doc.get("sha1"),
        "body_hash": doc.get("body_sha1"),
        "layout_version": doc.get("layout_version"),
        "stale": bool(doc.get("stale")),
        "table_index": table,
        "row_index": row,
        "column_indexes": columns if columns is not None else [],
        "raw_text": raw_text,
        "raw_cells": raw_cells if raw_cells is not None else [],
        **extra,
    }


def _row_record(doc: dict[str, Any], kind: str, table_index: int, row_index: int,
                row: tuple[Cell, ...], **extra: Any) -> dict[str, Any]:
    return _provenance(doc, record_type=kind, raw_cells=[cell.text for cell in row],
                       table=table_index, row=row_index,
                       columns=[list(cell.columns) for cell in row], **extra)


def _row_by_grid(row: tuple[Cell, ...]) -> dict[int, str]:
    result: dict[int, str] = {}
    for cell in row:
        for column in cell.columns:
            result[column] = cell.text
    return result


def _header_map(row: tuple[Cell, ...]) -> dict[int, str]:
    result: dict[int, str] = {}
    for cell in row:
        canonical = HEADER_ALIASES.get(norm(cell.text), strip_format_marks(cell.text).strip())
        for column in cell.columns:
            result[column] = canonical
    return result


def _mapped(row: tuple[Cell, ...], headers: dict[int, str]) -> dict[str, str]:
    values = _row_by_grid(row)
    return {heading: values.get(column, "") for column, heading in headers.items() if heading}


def _title_records(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for index, value in enumerate(data.paragraphs):
        if not value.strip():
            continue
        dates = dates_in_text(value)
        records.append(_provenance(
            doc, record_type="title", raw_text=value, row=index,
            parsed_dates=[day.isoformat() for day in dates], weekday=weekday_in_text(value),
        ))
    for index, value in enumerate(data.textboxes):
        records.append(_provenance(doc, record_type="signature", raw_text=value, row=index))
    return records


def _extract_roster(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    main_index = None
    main_header = None
    for ti, table in enumerate(data.tables):
        for ri, row in enumerate(table.rows[:4]):
            joined = norm(" ".join(cell.text for cell in row))
            if "الرتبه" in joined and "الاسم" in joined:
                main_index, main_header = ti, ri
                break
        if main_index is not None:
            break
    for ti, table in enumerate(data.tables):
        if ti == main_index and main_header is not None:
            headers = _header_map(table.rows[main_header])
            records.append(_row_record(doc, "header", ti, main_header, table.rows[main_header], headers=headers))
            for ri, row in enumerate(table.rows[main_header + 1:], main_header + 1):
                values = _mapped(row, headers)
                joined = norm(" ".join(cell.text for cell in row))
                section = next((name for name in ROSTER_SECTIONS if name in joined), None)
                if section:
                    records.append(_row_record(doc, "section", ti, ri, row, section=section))
                else:
                    records.append(_row_record(doc, "roster_row", ti, ri, row, fields=values))
        else:
            top = _row_by_grid(table.rows[0]) if table.rows else {}
            for ri, row in enumerate(table.rows):
                row_values = _row_by_grid(row)
                for cell in row:
                    for column in cell.columns:
                        path = [value for value in (top.get(column, ""), row_values.get(0, "")) if value]
                        records.append(_provenance(
                            doc, record_type="summary_cell", raw_text=cell.text, table=ti, row=ri,
                            columns=list(cell.columns), header_path=path,
                        ))
    return records


def _unknown_roster_layout(data: DocxData) -> str | None:
    for table in data.tables:
        for row in table.rows[:4]:
            header = norm(" ".join(cell.text for cell in row))
            if "الرتبه" not in header or "الاسم" not in header:
                continue
            has_signature = "الامضاء" in header
            has_notes = "ملاحظات" in header
            if has_signature and has_notes:
                return "roster_v3_signature_notes"
            if has_signature:
                return "roster_v3_signature"
            if has_notes:
                return "roster_v3_notes"
    return None


def _extract_old_officer(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    for ti, table in enumerate(data.tables):
        heading = data.headings[ti]
        for ri, row in enumerate(table.rows):
            values = [cell.text for cell in row if cell.text.strip()]
            records.append(_row_record(doc, "old_officer_row", ti, ri, row, heading=heading,
                                       rank_name=values[0] if values else "",
                                       service_status=values[1] if len(values) > 1 else ""))
    return records


# «خدمات أساسية» من غير «ال» في لوحات 2024-2026
_BASIC_HEADERS = ("خدمات الاساسيه", "خدمات اساسيه")


def _extract_board(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    for ti, table in enumerate(data.tables):
        if not table.rows:
            continue
        header_index = next((ri for ri, row in enumerate(table.rows) if
                             any(value in norm(" ".join(c.text for c in row)) for value in _BASIC_HEADERS) and
                             "الاهداف" in norm(" ".join(c.text for c in row))), 0)
        header = table.rows[header_index]
        starts = sorted({cell.columns[0] for cell in header
                         if any(signature in norm(cell.text) for signature in (*_BASIC_HEADERS, "الاهداف"))})
        known_header = len(starts) >= 2
        if len(starts) < 2:
            max_column = max((column for row in table.rows for cell in row for column in cell.columns), default=3)
            starts = [0, (max_column + 1) // 2]
        starts = starts[:2]
        records.append(_row_record(doc, "header", ti, header_index, header, half_starts=starts,
                                   known_section=known_header))
        for half_index, start in enumerate(starts):
            end = starts[half_index + 1] if half_index + 1 < len(starts) else 10**6
            for ri, row in enumerate(table.rows[header_index + 1:], header_index + 1):
                # الخلية العابرة للحد (gridSpan) بتتبع النص اللي فيه مركزها — مش النصين
                cells = tuple(cell for cell in row if start <= (cell.columns[0] + cell.columns[-1]) / 2 < end)
                nonempty = [cell.text for cell in cells if cell.text.strip()]
                if not nonempty:
                    continue
                kind = "board_label" if len(nonempty) == 1 else "board_row"
                records.append(_row_record(doc, kind, ti, ri, cells, half=half_index,
                                           label=nonempty[0], manning=nonempty[1] if len(nonempty) > 1 else "",
                                           known_section=any(section in norm(nonempty[0]) for section in BOARD_SECTIONS)))
    return records


def _align_afraad(values: list[str]) -> dict[str, str]:
    aligned: dict[str, str] = {}
    people: list[str] = []
    for value in values[1:]:
        simplified = norm(value)
        if not value.strip():
            continue
        if PERSON_RE.search(simplified):
            people.append(value)
        elif STRENGTH_RE.search(simplified):
            aligned.setdefault("قوام الخدمة", value)
        elif any(word in simplified for word in WEAPONS):
            aligned.setdefault("التسليح", value)
        elif TIME_RE.search(simplified):
            aligned.setdefault("الانتظام", value)
    if people:
        aligned["الخدمة الصباحية"] = people[0]
    if len(people) > 1:
        aligned["الخدمة الليلية"] = people[1]
    return aligned


def _extract_afraad(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    emergency = False
    for ti, table in enumerate(data.tables):
        header_index = next((ri for ri, row in enumerate(table.rows) if
                             "الخدمه الصباحيه" in norm(" ".join(c.text for c in row)) and
                             "الخدمه الليليه" in norm(" ".join(c.text for c in row))), None)
        for ri, row in enumerate(table.rows):
            joined = norm(" ".join(cell.text for cell in row))
            if ri == header_index:
                records.append(_row_record(doc, "afraad_header", ti, ri, row))
                continue
            if "الخدمات الطاريه" in joined:
                emergency = True
                records.append(_row_record(doc, "afraad_subheading", ti, ri, row, section="الخدمات الطارئة"))
                continue
            values = [cell.text for cell in row]
            nonempty = [value for value in values if value.strip()]
            if not nonempty:
                continue
            if len(nonempty) == 1:
                records.append(_row_record(doc, "afraad_subheading", ti, ri, row,
                                           section=nonempty[0], emergency=emergency))
            elif emergency:
                keys = ("service", "leader", "count", "weapon", "time_party")
                records.append(_row_record(doc, "afraad_emergency_row", ti, ri, row,
                                           fields=dict(zip(keys, values)), emergency=True))
            else:
                records.append(_row_record(doc, "afraad_basic_row", ti, ri, row,
                                           section="الخدمات الأساسية", aligned=_align_afraad(values)))
    return records


def _extract_duty(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    structured = any("قايم للخدمه" in norm(" ".join(c.text for c in row)) and
                     "ساعه الانتظام" in norm(" ".join(c.text for c in row))
                     for table in data.tables for row in table.rows)
    if not structured:
        for index, value in enumerate(data.paragraphs):
            if not value.strip():
                continue
            simplified = norm(value)
            kind = ("duty_supervisor" if "مشرف" in simplified else
                    "duty_time_heading" if TIME_RE.search(simplified) else "duty_service_line")
            records.append(_provenance(doc, record_type=kind, raw_text=value, row=index, era="A"))
    for ti, table in enumerate(data.tables):
        headers: dict[int, str] = {}
        for ri, row in enumerate(table.rows):
            joined = norm(" ".join(cell.text for cell in row))
            if structured and ("قايم للخدمه" in joined or ("الرتبه" in joined and "الخدمه" in joined)):
                headers = _header_map(row)
                records.append(_row_record(doc, "duty_header", ti, ri, row, era="B", headers=headers))
            elif not structured and len([c for c in row if c.text.strip()]) == 1 and TIME_RE.search(joined):
                records.append(_row_record(doc, "duty_time_heading", ti, ri, row, era="A"))
            else:
                kind = ("duty_officer_row" if any(norm(value) == "الرتبه" for value in headers.values())
                        else "duty_row")
                records.append(_row_record(doc, kind, ti, ri, row, era="B" if structured else "A",
                                           fields=_mapped(row, headers) if headers else {}))
    return records


def _extract_counts_docx(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    block = ""
    for ti, table in enumerate(data.tables):
        for ri, row in enumerate(table.rows):
            joined = norm(" ".join(cell.text for cell in row))
            for name in ("صباحية", "ليلية", "طوارئ"):
                if norm(name) in joined:
                    block = name
            kind = "counts_total" if "المجموع" in joined or "الاجمالي" in joined else "counts_row"
            records.append(_row_record(doc, kind, ti, ri, row, block=block))
    return records


def _extract_counts_xlsx(doc: dict[str, Any], path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    values_book = load_workbook(path, read_only=False, data_only=True)
    formula_book = load_workbook(path, read_only=False, data_only=False)
    records: list[dict[str, Any]] = []
    try:
        for sheet_index, sheet_name in enumerate(values_book.sheetnames):
            values_sheet = values_book[sheet_name]
            formula_sheet = formula_book[sheet_name]
            block = ""
            for ri in range(1, values_sheet.max_row + 1):
                values: list[str] = []
                formulas: dict[int, str] = {}
                columns: list[list[int]] = []
                for ci in range(1, values_sheet.max_column + 1):
                    value = values_sheet.cell(ri, ci).value
                    formula = formula_sheet.cell(ri, ci).value
                    raw = "" if value is None else strip_format_marks(str(value))
                    values.append(raw)
                    columns.append([ci - 1])
                    if isinstance(formula, str) and formula.startswith("=") and "SUM" in formula.upper():
                        formulas[ci - 1] = formula
                joined = norm(" ".join(values))
                for name in ("صباحية", "ليلية", "طوارئ"):
                    if norm(name) in joined:
                        block = name
                if any(value != "" for value in values) or formulas:
                    kind = "counts_total" if formulas or "المجموع" in joined or "الاجمالي" in joined else "counts_row"
                    records.append(_provenance(doc, record_type=kind, raw_cells=values,
                                               table=sheet_index, row=ri - 1, columns=columns,
                                               sheet=sheet_name, block=block, formulas=formulas))
    finally:
        values_book.close()
        formula_book.close()
    return records


def _xlsx_body_hash(path: Path) -> str:
    digest = hashlib.sha1()
    with zipfile.ZipFile(path) as package:
        names = sorted(name for name in package.namelist()
                       if name.startswith("xl/worksheets/") or name == "xl/sharedStrings.xml")
        for name in names:
            digest.update(name.encode("utf-8"))
            digest.update(package.read(name))
    return digest.hexdigest()


def _extract_command(doc: dict[str, Any], data: DocxData) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for index, value in enumerate(data.paragraphs):
        if not value.strip():
            continue
        parts = re.split(r"\s*[:：]\s*", value, maxsplit=1)
        records.append(_provenance(doc, record_type="command_row", raw_text=value, row=index,
                                   service=parts[0], detail=parts[1] if len(parts) > 1 else ""))
    return records


def _extract_generic_tables(doc: dict[str, Any], data: DocxData, prefix: str) -> list[dict[str, Any]]:
    records = _title_records(doc, data)
    for ti, table in enumerate(data.tables):
        headers = _header_map(table.rows[0]) if table.rows else {}
        for ri, row in enumerate(table.rows):
            records.append(_row_record(doc, f"{prefix}_{'header' if ri == 0 else 'row'}", ti, ri, row,
                                       fields=_mapped(row, headers) if ri else headers))
    return records


def extract_document(archive: Path, doc: dict[str, Any]) -> list[dict[str, Any]]:
    path = archive / doc["path"]
    document = _provenance(doc, record_type="document", raw_text="",
                           source_title_date=doc.get("title_date"))
    try:
        if path.suffix.lower() == ".xlsx":
            document["body_hash"] = _xlsx_body_hash(path)
            records = _extract_counts_xlsx(doc, path)
        elif path.suffix.lower() == ".docx":
            data = read_docx_raw(path)
            document["body_hash"] = data.body_sha1
            role = doc["role"]
            if role == "roster":
                if doc.get("layout_version") == "unknown":
                    inferred = _unknown_roster_layout(data)
                    if inferred:
                        doc = {**doc, "layout_version": inferred}
                        document["layout_version"] = inferred
                records = _extract_roster(doc, data)
            elif role == "old_officer_sheet":
                records = _extract_old_officer(doc, data)
            elif role == "board":
                records = _extract_board(doc, data)
            elif role == "afraad":
                records = _extract_afraad(doc, data)
            elif role == "duty_list":
                records = _extract_duty(doc, data)
            elif role in {"counts", "counts_leader"}:
                records = _extract_counts_docx(doc, data)
            elif role == "command_order":
                records = _extract_command(doc, data)
            elif role in {"match", "deployment", "ref_officers", "ref_personnel"}:
                records = _extract_generic_tables(doc, data, role)
            else:
                records = _extract_generic_tables(doc, data, role)
        else:
            raise ValueError(f"unsupported chosen extension: {path.suffix}")
        for record in records:
            record["body_hash"] = document["body_hash"]
        return [document, *records]
    except Exception as exc:  # فشل الملف يعزله ولا يوقف الدفعة.
        return [_provenance(doc, record_type="extract_error", raw_text="",
                            error_type=type(exc).__name__, error=str(exc))]


def run_extract(archive: Path, ledger: Ledger, start: dt.date, end: dt.date,
                state: dict[str, Any], *, resume: bool = False) -> dict[str, Any]:
    output = ledger.staging_path("extract")
    prior = state.get("stages", {}).get("extract", {})
    if resume and prior.get("status") == "complete" and output.exists():
        return prior.get("checkpoint", {})
    discover = ledger.staging_path("discover")
    if not discover.exists():
        raise FileNotFoundError("يجب تشغيل discover للدفعة أولًا")
    ledger.mark_stage(state, "extract", "running")
    documents = _selected(discover, start, end)
    records: list[dict[str, Any]] = []
    errors = 0
    for doc in documents:
        extracted = extract_document(archive, doc)
        records.extend(extracted)
        errors += int(extracted[0]["record_type"] == "extract_error")
    atomic_write_jsonl(output, records)
    checkpoint = {"documents": len(documents), "records": len(records), "errors": errors,
                  "version": VERSION}
    ledger.mark_stage(state, "extract", "complete", checkpoint)
    return checkpoint


__all__ = ["VERSION", "extract_document", "read_docx_raw", "run_extract"]
